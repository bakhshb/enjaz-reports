#!/usr/bin/env python3
# -*- coding: utf-8 -*-
# /// script
# requires-python = ">=3.12"
# dependencies = [
#   "pandas>=2.2,<3",
#   "openpyxl>=3.1,<4",
# ]
# ///
"""
Build the two-sheet "ملخص المهام" / "تفاصيل المهام" workbook from a raw
strategic-tasks export.

Usage:
    uv run build_report.py --input /path/to/source.xlsx --output "/mnt/user-data/outputs/ملخص وتفاصيل المهام الاستراتيجية.xlsx"

See ../SKILL.md for the full spec this implements. This script is deliberately
column-name-tolerant (aliases below) because the sheet/column names in the
source file drift slightly between exports; it never guesses *values* (status
mapping, missing fields) — only *which column* holds a given field.
"""
import argparse
import sys

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

COL_ALIASES = {
    "task": ["المهمة"],
    "sector": ["القطاع"],
    "source": ["المصدر"],
    "planned_date": ["تاريخ الإنجاز المخطط", "تاريخ الانجاز المخطط", "تاريخ الإنجاز المخطط له", "تاريخ الانجاز المخطط له"],
    "status": ["الحالة"],
    "notes": ["الملاحظات"],
    "related_parties": ["جهات ذات علاقة", "الجهات ذات العلاقة"],
    "support_request": ["طلب الدعم"],
}

STATUS_MAP = {
    "مكتمل": "مكتملة", "مكتملة": "مكتملة", "منجز": "مكتملة", "منجزة": "مكتملة",
    "على المخطط": "على المخطط", "علي المخطط": "على المخطط",
    "متأخر": "متأخر", "متاخر": "متأخر", "المتأخر": "متأخر", "متأخرة": "متأخر",
    "معلق": "معلق", "معلقة": "معلق",
}
STATUSES = ["مكتملة", "على المخطط", "متأخر", "معلق"]
SOURCE_ORDER = ["اجتماع القيادات", "اجتماع التجربة الرقمية", "لجنة المتابعة"]
DETAIL_STATUS_ORDER = ["متأخر", "معلق", "مكتملة", "على المخطط"]

def find_col(df, key):
    for alias in COL_ALIASES[key]:
        for c in df.columns:
            if str(c).strip() == alias:
                return c
    return None

def pick_sheets(xl):
    main_sheet = support_sheet = None
    for name in xl.sheet_names:
        df = xl.parse(name, nrows=1)
        cols = [str(c).strip() for c in df.columns]
        if any(a in cols for a in COL_ALIASES["support_request"]):
            support_sheet = name
        elif any(a in cols for a in COL_ALIASES["status"]):
            main_sheet = name
    return main_sheet, support_sheet

def normalize_status(raw, unmapped):
    s = str(raw).strip()
    if s in STATUS_MAP:
        return STATUS_MAP[s]
    unmapped.add(s)
    return s

def build(input_path, output_path):
    xl = pd.ExcelFile(input_path)
    main_name, support_name = pick_sheets(xl)
    if main_name is None:
        sys.exit("لم يتم العثور على شيت المهام الرئيسي (يحتاج عمود 'الحالة').")
    df = xl.parse(main_name)

    c_task = find_col(df, "task")
    c_sector = find_col(df, "sector")
    c_source = find_col(df, "source")
    c_date = find_col(df, "planned_date")
    c_status = find_col(df, "status")
    c_notes = find_col(df, "notes")
    c_related = find_col(df, "related_parties")
    missing = [k for k, v in {
        "المهمة": c_task, "القطاع": c_sector, "المصدر": c_source,
        "تاريخ الإنجاز المخطط": c_date, "الحالة": c_status,
    }.items() if v is None]
    if missing:
        sys.exit(f"أعمدة مفقودة في شيت المهام: {', '.join(missing)}")

    df[c_status] = df[c_status].astype(str).str.strip()
    df[c_sector] = df[c_sector].astype(str).str.strip()
    df[c_source] = df[c_source].astype(str).str.strip()

    unmapped = set()
    df["_status_norm"] = df[c_status].apply(lambda v: normalize_status(v, unmapped))
    if unmapped:
        sys.exit(
            "توجد حالات غير معروفة يجب مراجعتها يدويا قبل المتابعة (لم يتم تخمين تصنيفها): "
            + ", ".join(sorted(unmapped))
        )

    if df[c_task].duplicated().any():
        dups = df[c_task][df[c_task].duplicated()].tolist()
        print("تحذير: مهام مكررة بنفس النص:", dups, file=sys.stderr)

    total = len(df)
    status_counts = {s: int((df["_status_norm"] == s).sum()) for s in STATUSES}
    assert sum(status_counts.values()) == total

    sectors = list(dict.fromkeys(df[c_sector].tolist()))
    sector_status = {}
    for sec in sectors:
        sub = df[df[c_sector] == sec]
        sector_status[sec] = {s: int((sub["_status_norm"] == s).sum()) for s in STATUSES}
    for s in STATUSES:
        assert sum(v[s] for v in sector_status.values()) == status_counts[s]

    completed_tasks = df[df["_status_norm"] == "مكتملة"][[c_task, c_sector]].rename(
        columns={c_task: "task", c_sector: "sector"}).to_dict("records")
    delayed_tasks = df[df["_status_norm"] == "متأخر"][[c_task, c_sector]].rename(
        columns={c_task: "task", c_sector: "sector"}).to_dict("records")

    present_sources = list(dict.fromkeys(df[c_source].tolist()))
    sources = [s for s in SOURCE_ORDER if s in present_sources] + \
              [s for s in present_sources if s not in SOURCE_ORDER]

    sector_rank = {sec: i for i, sec in enumerate(sectors)}
    status_rank = {st: i for i, st in enumerate(DETAIL_STATUS_ORDER)}
    df["_sector_rank"] = df[c_sector].map(sector_rank)
    df["_status_rank"] = df["_status_norm"].map(status_rank)
    source_groups = {
        src: df[df[c_source] == src]
        .sort_values(["_status_rank", "_sector_rank"], kind="stable")
        .to_dict("records")
        for src in sources
    }
    assert sum(len(v) for v in source_groups.values()) == total

    sup_records = []
    if support_name is not None:
        sup = xl.parse(support_name)
        c_stask = find_col(sup, "task")
        c_ssector = find_col(sup, "sector")
        c_sreq = find_col(sup, "support_request")
        for rec in sup.to_dict("records"):
            sup_records.append({
                "task": None if c_stask is None or pd.isna(rec.get(c_stask)) else rec.get(c_stask),
                "sector": None if c_ssector is None or pd.isna(rec.get(c_ssector)) else rec.get(c_ssector),
                "support_request": None if c_sreq is None or pd.isna(rec.get(c_sreq)) else rec.get(c_sreq),
            })

    FONT_NAME = "Abar Mid"
    FONT_SIZE = 10
    NAVY, LIGHT_GRAY, WHITE = "1F4E78", "D9D9D9", "FFFFFF"
    title_font = Font(name=FONT_NAME, size=FONT_SIZE, bold=True, color=WHITE)
    title_fill = PatternFill("solid", fgColor=NAVY)
    header_font = Font(name=FONT_NAME, size=FONT_SIZE, bold=True)
    header_fill = PatternFill("solid", fgColor=LIGHT_GRAY)
    label_font = Font(name=FONT_NAME, size=FONT_SIZE, bold=True)
    value_font = Font(name=FONT_NAME, size=FONT_SIZE)
    normal_font = Font(name=FONT_NAME, size=FONT_SIZE)
    thin = Side(style="thin", color="A6A6A6")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    right_wrap = Alignment(horizontal="right", vertical="center", wrap_text=True)
    right_nowrap = Alignment(horizontal="right", vertical="center", wrap_text=False)
    center = Alignment(horizontal="center", vertical="center", wrap_text=False)

    def style_title(ws, row, col_span, text):
        ws.cell(row=row, column=1, value=text)
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=col_span)
        c = ws.cell(row=row, column=1)
        c.font, c.fill, c.alignment = title_font, title_fill, right_nowrap
        ws.row_dimensions[row].height = 22

    def style_header_row(ws, row, headers):
        for i, h in enumerate(headers, start=1):
            c = ws.cell(row=row, column=i, value=h)
            c.font, c.fill, c.border, c.alignment = header_font, header_fill, border, center

    def write_row(ws, row, values, wrap_cols=None, center_cols=None):
        wrap_cols, center_cols = wrap_cols or set(), center_cols or set()
        for i, v in enumerate(values, start=1):
            c = ws.cell(row=row, column=i, value=v)
            c.font, c.border = normal_font, border
            c.alignment = center if i in center_cols else (right_wrap if i in wrap_cols else right_nowrap)

    wb = Workbook()
    ws1 = wb.active
    ws1.title = "ملخص المهام"
    ws1.sheet_view.rightToLeft = True
    for col, w in {1: 40, 2: 46, 3: 30, 4: 42, 5: 14}.items():
        ws1.column_dimensions[get_column_letter(col)].width = w

    r = 1
    style_title(ws1, r, 5, "الملخص العام للمهام")
    r += 1
    for label, val in [
        ("إجمالي عدد المهام", total),
        ("المهام المكتملة", status_counts["مكتملة"]),
        ("المهام على المخطط", status_counts["على المخطط"]),
        ("المهام المتأخرة", status_counts["متأخر"]),
        ("المهام المعلقة", status_counts["معلق"]),
    ]:
        lc = ws1.cell(row=r, column=1, value=label)
        lc.font, lc.border, lc.alignment = label_font, border, right_nowrap
        vc = ws1.cell(row=r, column=2, value=val)
        vc.font, vc.border, vc.alignment = value_font, border, center
        r += 1
    r += 1

    style_title(ws1, r, 5, "تفصيل حالات المهام حسب القطاع")
    r += 1
    style_header_row(ws1, r, ["القطاع", "مكتملة", "على المخطط", "متأخر", "معلق"])
    r += 1
    for sec in sectors:
        counts = sector_status[sec]
        write_row(ws1, r, [sec] + [counts[s] or None for s in STATUSES], wrap_cols={1}, center_cols={2, 3, 4, 5})
        r += 1
    r += 1

    style_title(ws1, r, 5, "المهام المكتملة")
    r += 1
    style_header_row(ws1, r, ["#", "المهمة", "القطاع"])
    r += 1
    for i, t in enumerate(completed_tasks, start=1):
        write_row(ws1, r, [i, t["task"], t["sector"]], wrap_cols={2, 3}, center_cols={1})
        r += 1
    r += 1

    style_title(ws1, r, 5, "المهام المتأخرة")
    r += 1
    style_header_row(ws1, r, ["#", "المهمة", "القطاع"])
    r += 1
    for i, t in enumerate(delayed_tasks, start=1):
        write_row(ws1, r, [i, t["task"], t["sector"]], wrap_cols={2, 3}, center_cols={1})
        r += 1
    r += 1

    style_title(ws1, r, 5, "طلبات الدعم")
    r += 1
    style_header_row(ws1, r, ["#", "المهمة", "القطاع", "طلب الدعم"])
    r += 1
    for i, rec in enumerate(sup_records, start=1):
        write_row(ws1, r, [i, rec["task"], rec["sector"], rec["support_request"]], wrap_cols={2, 3, 4}, center_cols={1})
        r += 1

    ws2 = wb.create_sheet("تفاصيل المهام")
    ws2.sheet_view.rightToLeft = True
    for col, w in {1: 6, 2: 55, 3: 30, 4: 16, 5: 14, 6: 30, 7: 30}.items():
        ws2.column_dimensions[get_column_letter(col)].width = w

    headers2 = ["#", "المهمة", "القطاع", "تاريخ الانجاز المخطط", "الحالة", "ملاحظات", "الجهات ذات العلاقة"]
    r = 1
    for src in sources:
        style_title(ws2, r, 7, f"مهام مصدر ({src})")
        r += 1
        style_header_row(ws2, r, headers2)
        r += 1
        for i, rec in enumerate(source_groups[src], start=1):
            notes = rec.get(c_notes) if c_notes else None
            related = rec.get(c_related) if c_related else None
            notes = None if pd.isna(notes) else notes
            related = None if pd.isna(related) else related
            write_row(ws2, r, [
                i, rec.get(c_task), rec.get(c_sector), rec.get(c_date), rec.get("_status_norm"), notes, related,
            ], wrap_cols={2, 3, 6, 7}, center_cols={1, 4, 5})
            r += 1
        r += 1

    wb.save(output_path)

    print(f"تم الحفظ: {output_path}")
    print(f"إجمالي المهام: {total}")
    print(f"الحالات: {status_counts}")
    print(f"القطاعات ({len(sectors)}): {sectors}")
    print(f"المصادر ({len(sources)}): {sources}")
    print("تحقق: مجموع الحالات عبر القطاعات يطابق الإجماليات — OK")
    print("تحقق: مجموع المهام عبر المصادر يطابق الإجمالي — OK")

if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--input", required=True)
    ap.add_argument("--output", required=True)
    args = ap.parse_args()
    build(args.input, args.output)

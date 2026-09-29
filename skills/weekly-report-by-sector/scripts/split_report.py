# -*- coding: utf-8 -*-
"""Split the ministry-wide weekly report (تقرير المهام) into one deck per جهة.

Usage:
    python split_report.py MASTER.pptx --outdir out [--transactions FINAL.xlsx]
                           [--aliases aliases.json] [--no-refine]

The master deck supplies data and design. The final transactions Excel supplies
sector membership when the displayed transactions table omits that column. Nothing here
re-designs anything: every table, row, KPI card and title bar is cloned from a
shape that already exists in the master, and only the text, the fills that encode
status, and the geometry are changed.
"""
import argparse, json, math, os, posixpath, re, shutil, subprocess, sys, zipfile, pathlib, tempfile
from collections import defaultdict
from copy import deepcopy

from pptx import Presentation
from pptx.util import Emu
from lxml import etree

A = 'http://schemas.openxmlformats.org/drawingml/2006/main'
P = 'http://schemas.openxmlformats.org/presentationml/2006/main'
R = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'
def q(t):  return '{%s}%s' % (A, t)
def qp(t): return '{%s}%s' % (P, t)
EMU = 914400
SOFFICE = shutil.which('soffice')

# ---- layout constants taken from the master deck -------------------------
CX, CW   = 474245, 6623050      # content column: x and width of the KPI card
TOP_Y    = 1780255              # first table starts where the comparison chart was
GAP      = 200000               # vertical gap between stacked tables
KPI_ZONE = 1700000              # everything below this y on a section slide is body
FOOTER_Y = 12687759
NAV_XS   = (5741129, 5741130, 3970297, 2199462)   # nav tabs that stay
NAV_DROP_X = 428628                                # the التفاصيل tab
NAV_SHIFT  = 1012345            # shift left so the 3 remaining tabs are page-centred
NAV_TOPS   = (-1, -2, -3, 137425, 194575)

# ---- palette (read off the master's KPI cards) ---------------------------
STATUS_FILL = {'مكتملة':'DCE6F2', 'على المخطط':'EBF1DE', 'متأخرة':'F2DCDB',
               'لم تبدأ':'F2F2F2', 'معلقة':'F2F2F2'}
ACCENT = {'متأخرة':'87452B', 'مكتملة':'174766', 'على المخطط':'0A8062',
          'معلقة':'A6A6A6', 'لم تبدأ':'A6A6A6', 'تحديثات':'D9A86A'}

TASK_ORDER = ['متأخرة', 'معلقة', 'مكتملة', 'على المخطط']
TASK_TITLE = {'متأخرة':'المهام المتأخرة', 'مكتملة':'المهام المكتملة',
              'على المخطط':'المهام على المخطط', 'معلقة':'المهام المعلقة'}
PROJ_ORDER = ['متأخرة', 'على المخطط', 'لم تبدأ', 'مكتملة']
PROJ_TITLE = {'متأخرة':'المشاريع المتأخرة', 'على المخطط':'المشاريع على المخطط',
              'لم تبدأ':'المشاريع لم تبدأ', 'مكتملة':'المشاريع المكتملة'}

# =========================== small helpers ================================
def norm(s):
    return ' '.join(str(s).split())

def norm_status(s):
    s = norm(s)
    if s.startswith('مكتمل'): return 'مكتملة'
    if s.startswith('على المخطط'): return 'على المخطط'
    if s.startswith('متأخر'): return 'متأخرة'
    if s.startswith('لم تبدأ') or s.startswith('لم يبدأ'): return 'لم تبدأ'
    if s.startswith('معلق'): return 'معلقة'
    return s

def cell_text(c):
    """Cell text with <a:br/> preserved as newlines."""
    out = []
    for para in c.text_frame.paragraphs:
        s = ''
        for ch in para._p:
            if ch.tag == q('r'):   s += ch.find(q('t')).text or ''
            elif ch.tag == q('br'): s += '\n'
        out.append(s)
    return '\n'.join(out).strip()

def cell_text_el(tc):
    out = []
    for p in tc.iter(q('p')):
        sfrag = ''
        for ch in p:
            if ch.tag == q('r'): sfrag += ch.find(q('t')).text or ''
            elif ch.tag == q('br'): sfrag += '\n'
        out.append(sfrag)
    return '\n'.join(out).strip()

def table_rows(shape):
    return [[cell_text(c) for c in r.cells] for r in shape.table.rows]

def first_rpr(tc):
    for r in tc.iter(q('r')):
        rpr = r.find(q('rPr'))
        if rpr is not None: return deepcopy(rpr)
    for e in tc.iter(q('endParaRPr')):
        rpr = deepcopy(e); rpr.tag = q('rPr'); return rpr
    return None

def set_cell(tc, text, algn=None, fallback_rpr=None, bold_date_line=False):
    txb = tc.find(q('txBody'))
    rpr = first_rpr(tc)
    if rpr is None and fallback_rpr is not None: rpr = deepcopy(fallback_rpr)
    ps = txb.findall(q('p'))
    ppr = None
    if ps:
        e = ps[0].find(q('pPr'))
        if e is not None: ppr = deepcopy(e)
    for p in ps: txb.remove(p)
    p = etree.SubElement(txb, q('p'))
    if ppr is None:
        ppr = etree.Element(q('pPr')); ppr.set('rtl', '1')
    if algn:
        ppr.set('algn', algn); ppr.set('rtl', '1')
    p.append(ppr)
    text = re.sub('[\u064B-\u0652]', '', str(text))
    for i, ln in enumerate(str(text).split('\n')):
        if i:
            br = etree.SubElement(p, q('br'))
            if rpr is not None: br.append(deepcopy(rpr))
        r = etree.SubElement(p, q('r'))
        run_rpr = deepcopy(rpr) if rpr is not None else etree.Element(q('rPr'))
        if bold_date_line:
            is_date_line = (i == 0 and re.fullmatch(
                r'\s*تاريخ التحديث\s+\d{1,2}\s+\S+\s+\d{4}\s*', ln) is not None)
            run_rpr.set('b', '1' if is_date_line else '0')
        r.append(run_rpr)
        etree.SubElement(r, q('t')).text = ln

def set_fill(tc, hexcol):
    tcpr = tc.find(q('tcPr'))
    if tcpr is None: return
    for tag in ('solidFill', 'noFill', 'gradFill'):
        for e in tcpr.findall(q(tag)): tcpr.remove(e)
    sf = etree.SubElement(tcpr, q('solidFill'))
    etree.SubElement(sf, q('srgbClr')).set('val', hexcol)

def clean_title_row(tr, accent):
    """Borderless title row: only the small coloured accent bar survives.

    PowerPoint silently drops a whole <a:tcPr> whose children are out of schema
    order and falls back to a black default border, so rebuild it in order:
    lnL, lnR, lnT, lnB, then the fill.
    """
    for i, tc in enumerate(tr.findall(q('tc'))):
        tcpr = tc.find(q('tcPr'))
        if tcpr is None: tcpr = etree.SubElement(tc, q('tcPr'))
        attrs = dict(tcpr.attrib)
        for ch in list(tcpr): tcpr.remove(ch)
        tcpr.attrib.clear(); tcpr.attrib.update(attrs)
        for side in ('lnL', 'lnR', 'lnT', 'lnB'):
            ln = etree.SubElement(tcpr, q(side))
            ln.set('cap', 'flat'); ln.set('cmpd', 'sng'); ln.set('algn', 'ctr')
            if side == 'lnL' and i == 0 and accent:
                ln.set('w', '57150')
                sf = etree.SubElement(ln, q('solidFill'))
                etree.SubElement(sf, q('srgbClr')).set('val', accent)
            else:
                ln.set('w', '12700')
                etree.SubElement(ln, q('noFill'))
            etree.SubElement(ln, q('prstDash')).set('val', 'solid')
        etree.SubElement(tcpr, q('noFill'))

def expand_title_row(tr, ncols):
    """Re-merge a cloned title row across ncols columns.

    A <a:tc> added after <a:extLst> breaks the schema order and PowerPoint then
    ignores hMerge, leaving unmerged cells with black borders — so insert before it.
    """
    tcs = tr.findall(q('tc'))
    tcs[0].set('gridSpan', str(ncols))
    for extra in ('hMerge', 'vMerge', 'rowSpan'):
        tcs[0].attrib.pop(extra, None)
    filler = deepcopy(tcs[1]) if len(tcs) > 1 else deepcopy(tcs[0])
    filler.attrib.pop('gridSpan', None); filler.set('hMerge', '1')
    ext = tr.find(q('extLst'))
    while len(tr.findall(q('tc'))) > ncols:
        tr.remove(tr.findall(q('tc'))[-1])
    while len(tr.findall(q('tc'))) < ncols:
        f = deepcopy(filler)
        ext.addprevious(f) if ext is not None else tr.append(f)
    for tc in tr.findall(q('tc'))[1:]:
        tc.attrib.pop('gridSpan', None); tc.set('hMerge', '1')
    return tr

def gf_tbl(gf): return gf.find('.//' + q('tbl'))

def tbl_colwidths(tbl):
    return [int(g.get('w')) for g in tbl.find(q('tblGrid')).findall(q('gridCol'))]

def rescale(tbl, new_w):
    gcs = tbl.find(q('tblGrid')).findall(q('gridCol'))
    old = [int(g.get('w')) for g in gcs]; tot = sum(old)
    new = [int(round(w * new_w / tot)) for w in old]
    new[-1] += new_w - sum(new)
    for g, w in zip(gcs, new): g.set('w', str(w))

def gf_set_pos(gf, x, y, cx, cy):
    xf = gf.find(qp('xfrm'))
    xf.find(q('off')).set('x', str(int(x))); xf.find(q('off')).set('y', str(int(y)))
    xf.find(q('ext')).set('cx', str(int(cx))); xf.find(q('ext')).set('cy', str(int(cy)))

def next_id(spTree):
    ids = [int(e.get('id')) for e in spTree.iter(qp('cNvPr'))]
    return max(ids) + 1 if ids else 2

def add_frame(spTree, gf, name):
    gf = deepcopy(gf)
    nv = gf.find(qp('nvGraphicFramePr') + '/' + qp('cNvPr'))
    nv.set('id', str(next_id(spTree))); nv.set('name', name)
    spTree.append(gf)
    return gf

def make_table(tmpl_gf, title, rows, algns, status_col=None, title_tr=None, accent=None,
               bold_update_col=None, headers=None):
    gf = deepcopy(tmpl_gf)
    tbl = gf_tbl(gf)
    trs = tbl.findall(q('tr'))
    ncols = len(tbl_colwidths(tbl))
    body = deepcopy(trs[2])
    fb = first_rpr(body.findall(q('tc'))[0])
    if title_tr is not None:
        newt = expand_title_row(deepcopy(title_tr), ncols)
        tbl.replace(trs[0], newt); ttr = newt
    else:
        ttr = trs[0]
    tcs_t = ttr.findall(q('tc'))
    tcell = next((tc for tc in tcs_t if tc.get('gridSpan')), tcs_t[0])
    clean_title_row(ttr, accent)
    set_cell(tcell, title, algn='r', fallback_rpr=fb)
    if headers is not None:
        if len(headers) != ncols:
            raise ValueError('Header width does not match table: ' + title)
        for tc, value in zip(tbl.findall(q('tr'))[1].findall(q('tc')), headers):
            set_cell(tc, value, fallback_rpr=fb)
    for tr in tbl.findall(q('tr'))[2:]: tbl.remove(tr)
    for row in rows:
        if isinstance(row, dict):
            tr = deepcopy(row['tr']); tcs = tr.findall(q('tc'))
            set_cell(tcs[0], row['num'], algn='ctr', fallback_rpr=first_rpr(tcs[0]))
            if status_col is not None:
                set_fill(tcs[status_col],
                         STATUS_FILL.get(norm_status(cell_text_el(tcs[status_col])), 'F2F2F2'))
            tbl.append(tr)
            continue
        tr = deepcopy(body); tcs = tr.findall(q('tc'))
        if len(row) != ncols:
            raise ValueError('Row width does not match table: ' + title)
        for ci, val in enumerate(row):
            set_cell(tcs[ci], val, algn=algns[ci], fallback_rpr=fb,
                     bold_date_line=(ci == bold_update_col))
        if status_col is not None:
            set_fill(tcs[status_col], STATUS_FILL.get(norm_status(row[status_col]), 'F2F2F2'))
        tr.set('h', '152400')
        tbl.append(tr)
    return gf

# ---- first-pass height estimate (refine_positions corrects it later) -----
CPI, LINE, PAD = 15.5, 0.145, 0.12
def est_rows(rows, widths):
    if rows and isinstance(rows[0], dict):
        return 0.58 * EMU + sum(int(r['tr'].get('h', 152400)) for r in rows)
    return est_table(rows, widths)
def est_table(rows, widths):
    h = 0.28 + 0.30
    for r in rows:
        lines = 1
        for txt, w in zip(r, widths):
            cpl = max(6, (w / EMU - 0.12) * CPI)
            n = sum(max(1, math.ceil(len(seg) / cpl)) for seg in str(txt).split('\n'))
            lines = max(lines, n)
        h += PAD + lines * LINE
    return h * EMU

def set_estimated_row_heights(gf, rows, widths, frame_height):
    """Keep editable PowerPoint row geometry consistent before PDF refinement."""
    heights = [int(0.28 * EMU), int(0.30 * EMU)]
    if rows and isinstance(rows[0], dict):
        heights.extend(int(row['tr'].get('h', 152400)) for row in rows)
    else:
        for row in rows:
            lines = 1
            for value, width in zip(row, widths):
                cpl = max(6, (width / EMU - 0.12) * CPI)
                lines = max(lines, sum(max(1, math.ceil(len(part) / cpl))
                                       for part in str(value).split('\n')))
            heights.append(int((PAD + lines * LINE) * EMU))
    heights[-1] += int(frame_height) - sum(heights)
    for tr, height in zip(gf_tbl(gf).findall(q('tr')), heights):
        tr.set('h', str(height))

# ======================= reading the master deck ==========================
def find_tables(slide, must_contain):
    hits = []
    for sh in slide.shapes:
        if not sh.has_table: continue
        rows = table_rows(sh)
        head = ' '.join(rows[1]) if len(rows) > 1 else ''
        if all(k in head for k in must_contain): hits.append((sh, rows))
    return hits

def table_title(rows):
    return norm(next((x for x in rows[0] if norm(x)), '')) if rows else ''

def title_tables(slides, titles):
    for slide in slides:
        for sh in slide.shapes:
            if sh.has_table:
                rows = table_rows(sh)
                if rows and table_title(rows) in titles:
                    yield sh, rows

def summary_records(slides, title, text_key):
    out = []
    for _, rows in title_tables(slides, {title}):
        head = [norm(x) for x in rows[1]]
        detail_key = 'طلب الدعم' if title == 'طلبات الدعم' else 'التحدي'
        for required in (text_key, 'القطاع', detail_key):
            if required not in head:
                raise ValueError('%s is missing %s column' % (title, required))
        for row in rows[2:]:
            if any(norm(x) for x in row):
                out.append({'sector': row[head.index('القطاع')],
                            'name': row[head.index(text_key)],
                            'text': row[head.index(detail_key)]})
    return out

def locate(prs):
    """Identify slide roles by content, not by fixed slide numbers."""
    slides = list(prs.slides)
    visible = [i for i, s in enumerate(slides) if s._element.get('show') != '0']
    thanks = next((i for i in reversed(visible)
                   if any(sh.has_text_frame and 'شكرا' in sh.text_frame.text for sh in slides[i].shapes)),
                  visible[-1])
    role = {'cover': visible[0], 'thanks': thanks,
            'tasks': None, 'projects': None, 'trans': None,
            'task_details': [], 'proj_details': [], 'trans_slides': [],
            'task_summaries': [], 'project_summaries': []}
    for i in visible:
        s = slides[i]
        has_chart = any(sh.has_chart for sh in s.shapes)
        labels = ' '.join(sh.text_frame.text for sh in s.shapes if sh.has_text_frame)
        tables = [(sh, table_rows(sh)) for sh in s.shapes if sh.has_table]
        if any('رقم المعاملة' in [norm(x) for x in r[1]] for _, r in tables if len(r) > 1):
            role['trans_slides'].append(i)
            if role['trans'] is None: role['trans'] = i
        task_summary = any(table_title(r) in {'طلبات الدعم', 'المهام المعلقة',
            'المهام المكتملة', 'المهام المتأخرة', 'المهام على المخطط'} for _, r in tables)
        project_summary = any(table_title(r) in {'أبرز التحديثات', 'التحديات',
            'المشاريع المتأخرة'} for _, r in tables)
        if task_summary:
            role['task_summaries'].append(i)
        if project_summary:
            role['project_summaries'].append(i)
        if has_chart and 'إجمالي المهام' in labels and role['tasks'] is None:
            role['tasks'] = i
        if has_chart and 'إجمالي مشاريع سهيل' in labels and role['projects'] is None:
            role['projects'] = i
        if any(len(r[1]) == 6 and 'اسم المشروع' in r[1] for _, r in tables if len(r) > 1):
            role['proj_details'].append(i)
        if any(len(r[1]) == 6 and 'المهمة' in r[1] and 'الحالة' in r[1]
               for _, r in tables if len(r) > 1):
            role['task_details'].append(i)
    missing = [k for k in ('tasks', 'projects') if role[k] is None]
    if missing:
        raise SystemExit('Could not locate section slide(s): %s — check the master layout' % missing)
    if role['trans'] is None:
        print('  note: no delayed-transactions table (رقم المعاملة) in the master this week')
    return role

def read_data(prs, role):
    slides = list(prs.slides)
    tasks, projs, trans, ups, support, challenges = [], [], [], [], [], []
    for i in role['task_details']:
        for sh, rows in find_tables(slides[i], ['المهمة', 'الحالة']):
            if len(rows[1]) != 6: continue
            for r in rows[2:]:
                if r[0].strip(): tasks.append({'text': r[1], 'sector': r[2],
                                               'status': norm_status(r[4]),
                                               'note': r[5] if len(r) > 5 else ''})
    for i in role['proj_details']:
        for sh, rows in find_tables(slides[i], ['اسم المشروع']):
            if len(rows[1]) != 6: continue
            sector = rows[0][0]
            trs = sh._element.find('.//' + q('tbl')).findall(q('tr'))
            for ri, r in enumerate(rows[2:], start=2):
                if r[0].strip():
                    projs.append({'sector': sector, 'name': r[1], 'start': r[2],
                                  'end': r[3], 'status': norm_status(r[4]), 'note': r[5],
                                  'tr': deepcopy(trs[ri])})
    for i in role['trans_slides']:
        for sh, rows in find_tables(slides[i], ['رقم المعاملة']):
            head = [norm(x) for x in rows[1]]
            id_col = head.index('رقم المعاملة')
            for r in rows[2:]:
                if norm(r[id_col]): trans.append({'row': r, 'head': head})
    for i in role['project_summaries']:
        for sh, rows in title_tables([slides[i]], {'أبرز التحديثات'}):
            head = [norm(x) for x in rows[1]]
            for r in rows[2:]:
                if norm(r[head.index('المشروع')]):
                    ups.append({'project': r[head.index('المشروع')],
                                'sector': r[head.index('القطاع')],
                                'text': r[head.index('التحديث')]})
    for i in role['task_summaries']:
        support.extend(summary_records([slides[i]], 'طلبات الدعم', 'المهمة'))
    for i in role['project_summaries']:
        challenges.extend(summary_records([slides[i]], 'التحديات', 'المشروع'))
    return tasks, projs, trans, ups, support, challenges

TRANS_FIELDS = ('رقم المعاملة', 'موضوع المعاملة', 'الجهة الوارد منها المعاملة',
                'تاريخ انشاء المعاملة', 'تاريخ الإنجاز المخطط له')

def transaction_value(value):
    if value is None: return ''
    if hasattr(value, 'strftime'): return value.strftime('%d/%m/%Y')
    return str(value).strip()

def field_key(name):
    return norm(name).replace('إنشاء', 'انشاء').replace('الانجاز', 'الإنجاز').replace('المخطط له', 'المخطط')

def transaction_key(values):
    return tuple(norm(transaction_value(x)) for x in values)

def attach_transaction_sectors(trans, workbook_path):
    """Join the five displayed PPTX fields to the final Excel's sector column."""
    if not trans: return trans
    if all('القطاع' in item['head'] for item in trans):
        for item in trans:
            item['sector'] = item['row'][item['head'].index('القطاع')]
        return trans
    if not workbook_path:
        raise ValueError('The five-column transactions table needs --transactions FINAL.xlsx '
                         'to identify sectors.')
    from openpyxl import load_workbook
    wb = load_workbook(workbook_path, read_only=True, data_only=True)
    ws = wb['المعاملات'] if 'المعاملات' in wb.sheetnames else wb.active
    excel_rows = list(ws.values)
    header_i = next((i for i, r in enumerate(excel_rows)
                     if {'رقم المعاملة', 'موضوع المعاملة', 'القطاع'}.issubset(
                         {field_key(x) for x in r if x is not None})), None)
    if header_i is None:
        raise ValueError('Could not find the delayed-transactions header in ' + workbook_path)
    head = [field_key(x) for x in excel_rows[header_i]]
    col = {name: head.index(field_key(name)) for name in TRANS_FIELDS}
    sector_col = head.index('القطاع')
    lookup = defaultdict(list)
    for row in excel_rows[header_i + 1:]:
        if not norm(transaction_value(row[col['رقم المعاملة']])): continue
        values = [row[col[name]] for name in TRANS_FIELDS]
        lookup[transaction_key(values)].append(transaction_value(row[sector_col]))
    for item in trans:
        head = [field_key(x) for x in item['head']]
        try:
            values = [item['row'][head.index(field_key(name))] for name in TRANS_FIELDS]
        except ValueError as exc:
            raise ValueError('Unexpected five-column transactions header') from exc
        key = transaction_key(values)
        sectors = lookup.get(key, [])
        if not sectors:
            raise ValueError('Transaction %s does not match the final Excel workbook' % values[0])
        if len(set(sectors)) != 1:
            raise ValueError('Transaction %s matches multiple sectors in Excel' % values[0])
        item['sector'] = sectors.pop()
    wb.close()
    return trans

def kpi_shapes(slide):
    """Pair each KPI number box with its label box by horizontal proximity."""
    nums, labels = [], []
    for sh in slide.shapes:
        if not sh.has_text_frame or sh.top is None or sh.top >= KPI_ZONE: continue
        t = norm(sh.text_frame.text)
        if not t: continue
        cx = sh.left + sh.width / 2
        (nums if t.replace(',', '').isdigit() else labels).append((cx, sh, t))
    out = {}
    for cx, sh, t in nums:
        if not labels: continue
        _, lsh, lt = min(labels, key=lambda L: abs(L[0] - cx))
        key = 'إجمالي' if 'إجمالي' in lt else norm_status(lt)
        out[key] = {'num': sh, 'label': lsh}
    return out

def set_text(shape, value):
    done = False
    for p in shape.text_frame.paragraphs:
        for r in p.runs:
            if not done: r.text = str(value); done = True
            else: r.text = ''
    if not done: shape.text_frame.text = str(value)

# ============================ deck assembly ===============================
DY = 700000   # room opened under the cover title for the sector line

def fix_cover(prs, sector):
    slide = prs.slides[0]
    grp = slide.shapes[0]._element
    xf = grp.find(qp('grpSpPr')).find(q('xfrm'))
    off, ext, chext = xf.find(q('off')), xf.find(q('ext')), xf.find(q('chExt'))
    sps = [ch for ch in grp if ch.tag in (qp('sp'), qp('cxnSp'), qp('pic'), qp('graphicFrame'))]
    title_el = max(sps, key=lambda e: _font_size(e))
    date_el = next((e for e in sps if e is not title_el and e.find('.//' + q('t')) is not None), None)
    txf = _xfrm(title_el)
    base_y = int(txf.find(q('off')).get('y')) + int(txf.find(q('ext')).get('cy'))
    for ch in sps:
        if ch is title_el: continue
        x = _xfrm(ch)
        if x is None: continue
        o = x.find(q('off'))
        if o is not None: o.set('y', str(int(o.get('y')) + DY))
    ext.set('cy', str(int(ext.get('cy')) + DY))
    chext.set('cy', str(int(chext.get('cy')) + DY))
    off.set('y', str(int(off.get('y')) - DY // 2))

    new = deepcopy(date_el)
    new.find('.//' + qp('cNvPr')).set('id', '900')
    new.find('.//' + qp('cNvPr')).set('name', 'SectorName')
    xf2 = _xfrm(new)
    xf2.find(q('off')).set('x', str(int(xf.find(q('chOff')).get('x'))))
    xf2.find(q('off')).set('y', str(base_y + 30000))
    xf2.find(q('ext')).set('cx', str(int(chext.get('cx'))))
    xf2.find(q('ext')).set('cy', '620000')
    bp = new.find('.//' + q('bodyPr'))
    for af in bp.findall(q('spAutoFit')): bp.remove(af)
    bp.set('anchor', 'ctr'); bp.set('wrap', 'square')
    first = True
    for p in new.findall('.//' + q('p')):
        for r in p.findall(q('r')):
            if first:
                r.find(q('t')).text = sector
                rpr = r.find(q('rPr')); rpr.set('sz', '1600')
                for f in rpr.findall(q('solidFill')): rpr.remove(f)
                sf = etree.Element(q('solidFill'))
                etree.SubElement(sf, q('srgbClr')).set('val', 'FFFFFF')
                rpr.insert(0, sf); first = False
            else:
                p.remove(r)
    grp.append(new)

def _xfrm(el):
    for holder in (qp('spPr'), qp('grpSpPr')):
        h = el.find(holder)
        if h is not None:
            x = h.find(q('xfrm'))
            if x is not None: return x
    return None

def _font_size(sp):
    return max([int(r.get('sz', 0)) for r in sp.iter(q('rPr'))] + [0])

def strip_and_centre_nav(prs):
    for layout in prs.slide_masters[0].slide_layouts:
        for sh in list(layout.shapes):
            if sh.left == NAV_DROP_X and sh.top in NAV_TOPS:
                layout.shapes._spTree.remove(sh._element)
        for sh in layout.shapes:
            if sh.left in NAV_XS and sh.top in NAV_TOPS:
                sh.left = Emu(sh.left - NAV_SHIFT)

def clear_body(slide):
    """Drop the cross-sector chart, its card and the template table."""
    spTree = slide.shapes._spTree
    for sh in list(slide.shapes):
        if sh.top is not None and sh.top >= KPI_ZONE:
            spTree.remove(sh._element)

def continuation_slide(prs, source_index, keep, base_count):
    """Clone the section heading/KPIs for a continuation without its data body."""
    source = prs.slides[source_index]
    slide = prs.slides.add_slide(source.slide_layout)
    for sh in list(slide.shapes): slide.shapes._spTree.remove(sh._element)
    for sh in source.shapes:
        if sh.top is not None and sh.top < KPI_ZONE:
            slide.shapes._spTree.append(deepcopy(sh._element))
    new_index = len(prs.slides) - 1
    at = keep.index(source_index) + 1
    while at < len(keep) and keep[at] >= base_count:
        at += 1
    keep.insert(at, new_index)
    return slide, new_index

def place_table(prs, slide, y, source_index, keep, base_count, tmpl, title, rows, algns,
                frame_name, estimate, x=CX, frame_width=CW, page_top=TOP_Y, **kwargs):
    """Place whole rows on as many same-design summary pages as required."""
    width = tbl_colwidths(gf_tbl(tmpl))
    remaining = list(rows)
    while remaining:
        n = 0
        for count in range(1, len(remaining) + 1):
            if estimate(remaining[:count], width) + y <= FOOTER_Y - 300000:
                n = count
            else:
                break
        if n == 0:
            if y == page_top:
                raise ValueError('A single %s row cannot fit on an empty section page' % title)
            slide, _ = continuation_slide(prs, source_index, keep, base_count)
            y = page_top
            continue
        chunk = remaining[:n]
        gf = make_table(tmpl, title, chunk, algns, **kwargs)
        h = estimate(chunk, width)
        set_estimated_row_heights(gf, chunk, width, h)
        gf_set_pos(add_frame(slide.shapes._spTree, gf, frame_name), x, y, frame_width, h)
        y += h + GAP
        remaining = remaining[n:]
        if remaining:
            slide, _ = continuation_slide(prs, source_index, keep, base_count)
            y = page_top
    return slide, y

def single_kpi(slide, count, label_text):
    """Transactions: keep only the gold total, recentred, relabelled."""
    kp = kpi_shapes(slide)
    keep = kp.get('إجمالي')
    if keep is None: return
    keep_ids = {keep['num'].shape_id, keep['label'].shape_id}
    spTree = slide.shapes._spTree
    for sh in list(slide.shapes):
        if sh.top is not None and sh.top < KPI_ZONE:
            if sh.width and sh.width > 5000000:
                continue
            if sh.shape_id in keep_ids: continue
            spTree.remove(sh._element)
    set_text(keep['num'], count)
    set_text(keep['label'], label_text)
    keep['num'].left = Emu(int(CX + (CW - keep['num'].width) / 2))
    keep['label'].width = Emu(2600000)
    keep['label'].left = Emu(int(CX + (CW - 2600000) / 2))
    for p in keep['label'].text_frame.paragraphs:
        pPr = p._p.find(q('pPr'))
        if pPr is None:
            pPr = etree.Element(q('pPr')); p._p.insert(0, pPr)
        pPr.set('algn', 'ctr'); pPr.set('rtl', '1')

def build(master, outdir, aliases, transactions=None, only=None):
    os.makedirs(outdir, exist_ok=True)
    full_dir = os.path.join(outdir, '_full'); os.makedirs(full_dir, exist_ok=True)

    src = Presentation(master)
    role = locate(src)
    tasks, projs, trans, ups, support, challenges = read_data(src, role)
    attach_transaction_sectors(trans, transactions)

    canon = {norm(k): v for k, v in aliases.items()}
    def C(name): return canon.get(norm(name), norm(name))

    sectors = {}
    def B(n): return sectors.setdefault(n, {'tasks': [], 'projs': [], 'trans': [],
                                           'ups': [], 'support': [], 'challenges': []})
    for t in tasks: B(C(t['sector']))['tasks'].append(t)
    for p in projs: B(C(p['sector']))['projs'].append(p)
    for r in trans: B(C(r['sector']))['trans'].append(r)
    for u in ups:   B(C(u['sector']))['ups'].append(u)
    for u in support: B(C(u['sector']))['support'].append(u)
    for u in challenges: B(C(u['sector']))['challenges'].append(u)

    sl = list(src.slides)
    asset_path = pathlib.Path(__file__).resolve().parents[2] / 'weekly-consolidated-report' / 'assets' / 'weekly-report-master.pptx'
    design = Presentation(str(asset_path)) if asset_path.exists() else None
    design_slides = list(design.slides) if design else []
    def titled_frame(indices, title):
        found = list(title_tables([sl[i] for i in indices], {title}))
        if not found and design_slides:
            found = list(title_tables(design_slides, {title}))
        return deepcopy(found[0][0]._element) if found else None
    task_templates = {st: titled_frame(role['task_summaries'], TASK_TITLE[st])
                      for st in TASK_ORDER}
    task_fallback = next((v for v in task_templates.values() if v is not None), None)
    TMPL_SUPPORT = titled_frame(role['task_summaries'], 'طلبات الدعم')
    TMPL_UPD = titled_frame(role['project_summaries'], 'أبرز التحديثات')
    TMPL_CHALLENGE = titled_frame(role['project_summaries'], 'التحديات')
    TMPL_TRANS = deepcopy(find_tables(sl[role['trans']], ['رقم المعاملة'])[0][0]._element) if role['trans'] is not None else None
    detail_shapes = (find_tables(sl[role['proj_details'][-1]], ['اسم المشروع'])
                     if role['proj_details'] else [])
    if not detail_shapes and design_slides:
        detail_shapes = next((find_tables(s, ['اسم المشروع']) for s in design_slides
                              if find_tables(s, ['اسم المشروع'])), [])
    if not detail_shapes:
        raise ValueError('No approved project detail table found for status formatting')
    TMPL_PROJ = deepcopy(detail_shapes[0][0]._element)
    if TMPL_UPD is None:
        raise ValueError('No approved Suhail summary table found for title formatting')
    GOLD_TITLE = deepcopy(gf_tbl(TMPL_UPD).findall(q('tr'))[0])
    rescale(gf_tbl(TMPL_PROJ), CW)

    manifest = []
    for sec in sorted(sectors, key=lambda s: (0 if s.startswith('وكالة') else 1, s)):
        if only and only not in sec: continue
        d = sectors[sec]
        if not any(d.values()): continue
        prs = Presentation(master)
        strip_and_centre_nav(prs)
        fix_cover(prs, sec)
        keep = [role['cover']]
        if d['tasks'] or d['support']: keep.append(role['tasks'])
        if d['projs'] or d['ups'] or d['challenges']: keep.append(role['projects'])
        if d['trans']: keep.append(role['trans'])
        keep.append(role['thanks'])

        if d['tasks'] or d['support']:
            s = prs.slides[role['tasks']]
            kp = kpi_shapes(s)
            cnt = {k: sum(1 for t in d['tasks'] if t['status'] == k) for k in TASK_ORDER}
            for key, val in [('إجمالي', len(d['tasks']))] + list(cnt.items()):
                if key in kp: set_text(kp[key]['num'], val)
            clear_body(s)
            y = TOP_Y
            for st in TASK_ORDER:
                sel = [x for x in d['tasks'] if x['status'] == st]
                if not sel: continue
                tmpl = task_templates[st] if task_templates[st] is not None else task_fallback
                if tmpl is None: raise ValueError('No task status table to clone')
                headers = [cell_text_el(tc) for tc in gf_tbl(tmpl).findall(q('tr'))[1].findall(q('tc'))]
                if task_templates[st] is None:
                    headers = ['#', 'المهمة', 'القطاع', 'ملاحظات'] if len(headers) == 4 else ['#', 'المهمة', 'القطاع']
                rows = []
                for i, t in enumerate(sel):
                    values = {'#': str(i+1), 'المهمة': t['text'], 'القطاع': sec,
                              'ملاحظات': t['note'].strip(), 'التحديث': t['note'].strip()}
                    rows.append([values.get(norm(h), '') for h in headers])
                s, y = place_table(prs, s, y, role['tasks'], keep, len(sl), tmpl,
                                   TASK_TITLE[st], rows, ['ctr' if h == '#' else 'r' for h in headers],
                                   'tbl_' + st, est_table, accent=ACCENT[st], headers=headers)
            if d['support']:
                if TMPL_SUPPORT is None: raise ValueError('No support table to clone')
                headers = [cell_text_el(tc) for tc in gf_tbl(TMPL_SUPPORT).findall(q('tr'))[1].findall(q('tc'))]
                rows = []
                for i, u in enumerate(d['support']):
                    values = {'#': str(i+1), 'المهمة': u['name'], 'القطاع': sec, 'طلب الدعم': u['text']}
                    rows.append([values.get(norm(h), '') for h in headers])
                s, y = place_table(prs, s, y, role['tasks'], keep, len(sl), TMPL_SUPPORT,
                                   'طلبات الدعم', rows, ['ctr' if h == '#' else 'r' for h in headers],
                                   'tbl_support', est_table, accent=ACCENT['تحديثات'])

        if d['projs'] or d['ups'] or d['challenges']:
            s = prs.slides[role['projects']]
            kp = kpi_shapes(s)
            cnt = {k: sum(1 for p in d['projs'] if p['status'] == k) for k in PROJ_ORDER}
            for key, val in [('إجمالي', len(d['projs']))] + list(cnt.items()):
                if key in kp: set_text(kp[key]['num'], val)
            clear_body(s)
            y = TOP_Y
            if d['ups']:
                if TMPL_UPD is None: raise ValueError('No updates table to clone')
                rows = [[str(i + 1), u['project'], sec, u['text']] for i, u in enumerate(d['ups'])]
                s, y = place_table(prs, s, y, role['projects'], keep, len(sl), TMPL_UPD,
                                   'أبرز التحديثات', rows, ['ctr', 'r', 'ctr', 'r'],
                                   'tbl_updates', est_table, accent=ACCENT['تحديثات'],
                                   bold_update_col=3)
            if d['challenges']:
                if TMPL_CHALLENGE is None: raise ValueError('No challenges table to clone')
                headers = [cell_text_el(tc) for tc in gf_tbl(TMPL_CHALLENGE).findall(q('tr'))[1].findall(q('tc'))]
                rows = []
                for i, u in enumerate(d['challenges']):
                    values = {'#': str(i+1), 'المشروع': u['name'], 'القطاع': sec, 'التحدي': u['text']}
                    rows.append([values.get(norm(h), '') for h in headers])
                s, y = place_table(prs, s, y, role['projects'], keep, len(sl), TMPL_CHALLENGE,
                                   'التحديات', rows, ['ctr' if h == '#' else 'r' for h in headers],
                                   'tbl_challenges', est_table, accent=ACCENT['متأخرة'])
            for st in PROJ_ORDER:
                rows = [{'num': str(i + 1), 'tr': p['tr']}
                        for i, p in enumerate([x for x in d['projs'] if x['status'] == st])]
                if not rows: continue
                s, y = place_table(prs, s, y, role['projects'], keep, len(sl), TMPL_PROJ,
                                   PROJ_TITLE[st], rows, ['ctr', 'r', 'ctr', 'ctr', 'ctr', 'r'],
                                   'proj_' + st, est_rows, status_col=4,
                                   title_tr=GOLD_TITLE, accent=ACCENT[st])

        if d['trans']:
            s = prs.slides[role['trans']]
            tframe = find_tables(s, ['رقم المعاملة'])[0][0]
            tx, ty, tw = tframe.left, tframe.top, tframe.width
            clear_body(s)
            single_kpi(s, len(d['trans']), 'إجمالي المعاملات المتأخرة')
            rows = []
            for item in d['trans']:
                rows.append([item['row'][item['head'].index(h)] for h in item['head']])
            place_table(prs, s, ty, role['trans'], keep, len(sl), TMPL_TRANS,
                        'المعاملات المتأخرة', rows,
                        ['ctr' if 'تاريخ' in h or 'رقم' in h else 'r' for h in d['trans'][0]['head']],
                        'tbl_trans', est_table, x=tx, frame_width=tw, page_top=ty,
                        accent=ACCENT['متأخرة'])

        path = os.path.join(full_dir, 'تقرير الإنجاز الأسبوعي - %s.pptx' % sec)
        prs.save(path)
        manifest.append({'sector': sec, 'file': path, 'keep': keep,
                         'edited': [i for i in keep if i in (role['tasks'], role['projects'], role['trans'])
                                    or i >= len(sl)],
                         'n': [len(d['tasks']), len(d['projs']), len(d['trans']), len(d['ups']),
                               len(d['support']), len(d['challenges'])],
                         'nslides': len(sl)})
    return manifest

# ================== pass 2: measure, then reflow ==========================
def _bands(page):
    rs = [r for r in page.rects if 20 < r['width'] < 560 and r['height'] < 200 and r['top'] > 60]
    rs.sort(key=lambda r: r['top'])
    gs = []
    for r in rs:
        if gs and r['top'] - gs[-1][1] < 10: gs[-1][1] = max(gs[-1][1], r['bottom'])
        else: gs.append([r['top'], r['bottom']])
    return gs

def to_pdf(pptx_path, work='_m'):
    shutil.copy(pptx_path, work + '.pptx')
    if os.path.exists(work + '.pdf'): os.remove(work + '.pdf')
    if not SOFFICE:
        raise RuntimeError('soffice is required for the refinement pass')
    pdf_dir = tempfile.mkdtemp(prefix='soffice-pdf-')
    profile_dir = tempfile.mkdtemp(prefix='soffice-profile-')
    try:
        subprocess.run([SOFFICE, '-env:UserInstallation=file://' + profile_dir,
                        '--headless', '--convert-to', 'pdf',
                        '--outdir', pdf_dir, work + '.pptx'],
                       capture_output=True, check=True)
        shutil.move(os.path.join(pdf_dir, os.path.basename(work) + '.pdf'),
                    work + '.pdf')
    finally:
        shutil.rmtree(pdf_dir, ignore_errors=True)
        shutil.rmtree(profile_dir, ignore_errors=True)
    os.remove(work + '.pptx')
    return work + '.pdf'

SAFETY = 1.12
OFFPAGE = 40_000_000

def _row_bands(page, x0, x1, y0):
    """Row rectangles of the one table rendered on this page, top to bottom (points)."""
    rs = [r for r in page.rects if r['x0'] > x0 - 3 and r['x1'] < x1 + 3
          and r['top'] > y0 - 3 and r['height'] > 2]
    bands = []
    for r in sorted(rs, key=lambda r: (r['top'], r['bottom'])):
        if bands and r['top'] < bands[-1][1] - 1:
            bands[-1][1] = max(bands[-1][1], r['bottom'])
        else:
            bands.append([r['top'], r['bottom']])
    return bands

def _frames(slide):
    return sorted([sh for sh in slide.shapes if sh.has_table], key=lambda s: s.top)

def refine_positions(manifest):
    """Measure every row and write the real height back, so a table's declared
    geometry matches what it occupies. Stacked tables overlap each other in the
    render, which would corrupt the measurement, so each pass sends every table
    but the one being measured off the page."""
    import pdfplumber
    for item in manifest:
        prs = Presentation(item['file'])
        nmax = max((len(_frames(prs.slides[si])) for si in item['edited']), default=0)
        heights = {}
        for k in range(nmax):
            probe = Presentation(item['file'])
            for si in item['edited']:
                for j, sh in enumerate(_frames(probe.slides[si])):
                    if j != k: sh.top = Emu(OFFPAGE + j * 100000)
            probe.save('_probe.pptx')
            pdf_path = to_pdf('_probe.pptx')
            with pdfplumber.open(pdf_path) as pdf:
                for si in item['edited']:
                    fr = _frames(prs.slides[si])
                    if k >= len(fr): continue
                    sh = fr[k]
                    ntr = len(sh._element.find('.//' + q('tbl')).findall(q('tr')))
                    bands = _row_bands(pdf.pages[si], sh.left / EMU * 72,
                                       (sh.left + sh.width) / EMU * 72, sh.top / EMU * 72)
                    if len(bands) != ntr - 1:
                        print('  ! row-band mismatch on slide %d of %s (%d bands, %d rows)'
                              % (si + 1, item['sector'], len(bands), ntr))
                        continue
                    title_h = max(0.0, bands[0][0] - sh.top / EMU * 72)
                    heights[(si, k)] = [int(h / 72 * EMU * SAFETY)
                                        for h in [title_h] + [b[1] - b[0] for b in bands]]
            os.remove(pdf_path); os.remove('_probe.pptx')
        for si in item['edited']:
            fr = _frames(prs.slides[si])
            y = fr[0].top if fr else 0
            for k, sh in enumerate(fr):
                hs = heights.get((si, k))
                if hs:
                    for tr, h in zip(sh._element.find('.//' + q('tbl')).findall(q('tr')), hs):
                        tr.set('h', str(h))
                    sh.height = Emu(sum(hs))
                sh.top = Emu(int(y))
                y += sh.height + GAP
        prs.save(item['file'])

# ============ pass 3: drop unused slides at the XML level =================
def shrink_media(wd):
    """The three full-bleed backgrounds are ~9 MB of the master. Re-encode them at
    the same pixel size — no visible change, and a sector deck lands around 2.7 MB."""
    from PIL import Image
    for f in pathlib.Path(wd, 'ppt', 'media').glob('*'):
        if f.suffix.lower() not in ('.jpg', '.jpeg') or f.stat().st_size < 500_000: continue
        im = Image.open(f)
        buf = __import__('io').BytesIO()
        im.save(buf, 'JPEG', quality=85, optimize=True)
        if buf.tell() < f.stat().st_size: f.write_bytes(buf.getvalue())

def reachable_parts(wd):
    """Remove links to discarded slides, then keep only reachable package parts."""
    root=pathlib.Path(wd)
    ns={'p':P,'r':'http://schemas.openxmlformats.org/package/2006/relationships'}
    pres=etree.parse(str(root/'ppt/presentation.xml'))
    ids={el.get('{'+R+'}id') for el in pres.findall('.//p:sldId',ns)}
    def target(name,rel):
        source=posixpath.join(posixpath.dirname(posixpath.dirname(name)),posixpath.basename(name)[:-5])
        return posixpath.normpath(posixpath.join(posixpath.dirname(source),rel.get('Target','').split('#',1)[0])).lstrip('/')
    prels='ppt/_rels/presentation.xml.rels'
    kept={target(prels,rel) for rel in etree.parse(str(root/prels)).getroot()
          if rel.get('Id') in ids}
    for path in root.rglob('*.rels'):
        name=path.relative_to(root).as_posix(); tree=etree.parse(str(path)); rels=tree.getroot(); removed=set()
        for rel in list(rels):
            if rel.get('Type','').endswith('/slide') and target(name,rel) not in kept:
                removed.add(rel.get('Id'));rels.remove(rel)
        if removed:
            tree.write(str(path),xml_declaration=True,encoding='UTF-8',standalone=True)
            source=path.parent.parent/path.name[:-5]
            if source.is_file() and source.suffix=='.xml':
                document=etree.parse(str(source))
                for link in document.xpath('//*[@r:id]',namespaces={'r':R}):
                    if link.get('{'+R+'}id') in removed and link.tag in {q('hlinkClick'),q('hlinkMouseOver')}:
                        link.getparent().remove(link)
                document.write(str(source),xml_declaration=True,encoding='UTF-8',standalone=True)
    reachable={'[Content_Types].xml'}; pending=['']
    while pending:
        source=pending.pop()
        relname=(posixpath.join(posixpath.dirname(source),'_rels',posixpath.basename(source)+'.rels') if source else '_rels/.rels')
        if relname in reachable or not (root/relname).exists():continue
        reachable.add(relname)
        for rel in etree.parse(str(root/relname)).getroot():
            if rel.get('TargetMode')=='External':continue
            part=target(relname,rel)
            if not (root/part).is_file():raise ValueError('Missing package target: '+part)
            if part not in reachable:reachable.add(part);pending.append(part)
    content_types=root/'[Content_Types].xml'; tree=etree.parse(str(content_types))
    for child in list(tree.getroot()):
        if child.tag.endswith('Override') and child.get('PartName','').lstrip('/') not in reachable:
            tree.getroot().remove(child)
    tree.write(str(content_types),xml_declaration=True,encoding='UTF-8',standalone=True)
    return reachable


def slim(manifest, outdir):
    """Remove unwanted slides and their unreachable content before packaging."""
    final = []
    for item in manifest:
        wd = os.path.join(outdir, '_work')
        shutil.rmtree(wd, ignore_errors=True)
        with zipfile.ZipFile(item['file']) as z: z.extractall(wd)
        pres = pathlib.Path(wd) / 'ppt' / 'presentation.xml'
        xml = pres.read_text(encoding='utf-8')
        m = re.search(r'(<p:sldIdLst>)(.*?)(</p:sldIdLst>)', xml, re.S)
        ids = re.findall(r'<p:sldId\b[^>]*/>', m.group(2))
        keep = item['keep']
        xml = xml[:m.start()] + m.group(1) + ''.join(ids[i] for i in keep) \
              + m.group(3) + xml[m.end():]
        pres.write_text(xml, encoding='utf-8')
        included=reachable_parts(wd)
        # back-references from layouts to deleted slides, and the nav hyperlinks
        for rels in pathlib.Path(wd).rglob('_rels/*.rels'):
            base = rels.parent.parent
            txt = rels.read_text(encoding='utf-8'); orig = txt
            for tag in re.findall(r'<Relationship\b[^>]*/>', txt):
                if 'TargetMode="External"' in tag: continue
                tm = re.search(r'Target="([^"]+)"', tag)
                if tm and not (base / tm.group(1)).resolve().exists():
                    txt = txt.replace(tag, '')
            if txt != orig: rels.write_text(txt, encoding='utf-8')
        for rels in pathlib.Path(wd).rglob('_rels/*.rels'):
            part = rels.parent.parent / rels.name[:-5]
            if not part.exists() or part.suffix != '.xml': continue
            valid = set(re.findall(r'Id="([^"]+)"', rels.read_text(encoding='utf-8')))
            txt = part.read_text(encoding='utf-8'); orig = txt
            def drop(mm):
                rid = re.search(r'r:id="([^"]*)"', mm.group(0))
                return '' if (rid and rid.group(1) and rid.group(1) not in valid) else mm.group(0)
            txt = re.sub(r'<a:hlinkClick\b[^>]*/>', drop, txt)
            txt = re.sub(r'<a:hlinkClick\b[^>]*>.*?</a:hlinkClick>', drop, txt, flags=re.S)
            if txt != orig: part.write_text(txt, encoding='utf-8')
        shrink_media(wd)
        out = os.path.abspath(os.path.join(outdir, os.path.basename(item['file'])))
        if os.path.exists(out): os.remove(out)
        with zipfile.ZipFile(out, 'w', compression=zipfile.ZIP_DEFLATED) as archive:
            for part in pathlib.Path(wd).rglob('*'):
                name=part.relative_to(wd).as_posix()
                if part.is_file() and name in included: archive.write(part,name)
        shutil.rmtree(wd, ignore_errors=True)
        item['out'] = out
        final.append(item)
    shutil.rmtree(os.path.join(outdir, '_full'), ignore_errors=True)
    return final

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('master')
    ap.add_argument('--outdir', default='out')
    ap.add_argument('--aliases', default=os.path.join(
        os.path.dirname(os.path.abspath(__file__)), '..', 'references', 'sector-aliases.json'))
    ap.add_argument('--no-refine', action='store_true')
    ap.add_argument('--only', default=None)
    ap.add_argument('--transactions', help='Final consolidated transactions Excel; required for five-column transaction tables')
    a = ap.parse_args()
    aliases = json.load(open(a.aliases, encoding='utf-8')) if os.path.exists(a.aliases) else {}
    man = build(a.master, a.outdir, aliases, a.transactions, a.only)
    if not a.no_refine:
        refine_positions(man)
    man = slim(man, a.outdir)
    with open(os.path.join(a.outdir, 'manifest.json'), 'w', encoding='utf-8') as stream:
        json.dump(man, stream, ensure_ascii=False, indent=1)
    for it in man:
        print('%-45s tasks=%d projects=%d transactions=%d updates=%d support=%d challenges=%d slides=%d'
              % (it['sector'], *it['n'], len(it['keep'])))

if __name__ == '__main__':
    main()

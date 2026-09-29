#!/usr/bin/env python3
# /// script
# requires-python = ">=3.12"
# dependencies = [
#   "pandas>=2.2,<3",
#   "openpyxl>=3.1,<4",
#   "xlrd>=2.0.1,<3",
# ]
# ///
"""Consolidate per-sector Arabic transaction exports into one Excel report.

Usage:
    uv run build_report.py --input-dir DIR --output FILE.xlsx [--dedupe]

Reads every .xls/.xlsx in DIR, normalises statuses onto the three ministry
categories, and writes a workbook with a four-line summary on top, a status
breakdown by sector beneath it, and then a table of the delayed transactions,
sorted by sector then planned completion date.

Never run a LibreOffice recalculation on the output: it turns the text
transaction numbers into real numbers (381117010117 -> 381117010117.00).
"""

import argparse
import glob
import os
import re
import sys
import tempfile

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

CTRL = dict.fromkeys(
    map(ord, "\u202a\u202b\u202c\u202d\u202e\u200e\u200f\u061c"), None
)

STATUS_MAP = {
    "مكتمل": "مكتملة", "مكتملة": "مكتملة", "منجز": "مكتملة", "منجزة": "مكتملة",
    "على المخطط": "على المخطط", "علي المخطط": "على المخطط",
    "متأخر": "متأخرة", "متاخر": "متأخرة", "متأخرة": "متأخرة", "المتأخر": "متأخرة",
}
CATEGORIES = ["مكتملة", "على المخطط", "متأخرة"]
SECTOR_HEADERS = ["مكتملة", "على المخطط", "متأخر"]

COL = {"status": 19, "planned": 23, "source": 26, "created": 27,
       "subject": 28, "number": 30}
SECTOR_CELL = (1, 12)

GREEN, GOLD, LIGHT, WHITE = "0F5C3F", "D7A562", "F2F6F4", "FFFFFF"
FONT, FONT_SIZE = "Abar Mid", 10

def clean(value):
    if pd.isna(value):
        return ""
    return re.sub(r"\s+", " ", str(value).translate(CTRL).strip())

def norm_date(value):
    text = clean(value)
    match = re.match(r"^(\d{1,2})[/-](\d{1,2})[/-](\d{4})$", text)
    if match:
        day, month, year = (int(g) for g in match.groups())
        return f"{day:02d}/{month:02d}/{year}", (year, month, day)
    return text, (9999, 99, 99)

def read_files(input_dir):
    paths = sorted(
        p for p in glob.glob(os.path.join(input_dir, "*"))
        if p.lower().endswith((".xls", ".xlsx"))
    )
    if not paths:
        sys.exit(f"No .xls/.xlsx files found in {input_dir}")

    rows, unmapped, sectors = [], set(), []
    for path in paths:
        sheet = pd.read_excel(path, header=None, dtype=object)
        header = None
        for i, row in sheet.iterrows():
            if any("رقم المعاملة" in clean(v) for v in row if pd.notna(v)):
                header = i
                break
        if header is None:
            sys.exit(f"Missing رقم المعاملة header in {os.path.basename(path)}")
        required_columns = max(SECTOR_CELL[1], *COL.values()) + 1
        if sheet.shape[1] < required_columns or sheet.shape[0] <= SECTOR_CELL[0]:
            sys.exit(f"Missing required source cells in {os.path.basename(path)}")

        sector = clean(sheet.iloc[SECTOR_CELL]) or clean(
            os.path.basename(path).rsplit(".", 1)[0].replace("_", " ")
        )
        for _, row in sheet.iloc[header + 1:].iterrows():
            number = clean(row[COL["number"]])
            if number.endswith(".0"):
                number = number[:-2]
            if not number:
                continue
            raw_status = clean(row[COL["status"]])
            status = STATUS_MAP.get(raw_status)
            if status is None:
                unmapped.add(raw_status)
            created, _ = norm_date(row[COL["created"]])
            planned, planned_key = norm_date(row[COL["planned"]])
            rows.append({
                "num": number,
                "subject": clean(row[COL["subject"]]),
                "sector": sector,
                "source": clean(row[COL["source"]]),
                "created": created,
                "planned": planned,
                "plan_key": planned_key,
                "status": status,
                "raw_status": raw_status,
            })
        if sector not in sectors:
            sectors.append(sector)
        print(f"  read {os.path.basename(path)} -> {sector}")
    # Keep the schema when every valid sector export contains zero records.
    columns = ["num", "subject", "sector", "source", "created", "planned",
               "plan_key", "status", "raw_status"]
    return pd.DataFrame(rows, columns=columns), unmapped, len(paths), sorted(sectors)

def style_cell(cell, *, bold=False, color="000000", fill=None,
               align="right", wrap=False, border=None):
    cell.font = Font(name=FONT, size=FONT_SIZE, bold=bold, color=color)
    if fill:
        cell.fill = PatternFill("solid", fgColor=fill)
    cell.alignment = Alignment(
        horizontal=align, vertical="center", wrap_text=wrap,
        indent=1 if align == "right" else 0,
    )
    if border:
        cell.border = border

def section_band(ws, row, text, last_col):
    for col in range(1, last_col + 1):
        style_cell(ws.cell(row=row, column=col), bold=True, color=WHITE,
                   fill=GREEN)
    ws.cell(row=row, column=1, value=text)
    ws.row_dimensions[row].height = 26

def build(df, sectors, output, dedupe, n_files, duplicate_records):
    counts = {c: int((df.status == c).sum()) for c in CATEGORIES}
    total = len(df)
    delayed = df[df.status == "متأخرة"].sort_values(by=["sector", "plan_key"])
    assert sum(counts.values()) == total, "counts do not balance"
    breakdown = {sec: {c: int(((df.sector == sec) & (df.status == c)).sum())
                       for c in CATEGORIES} for sec in sectors}
    assert sum(sum(v.values()) for v in breakdown.values()) == total, \
        "sector breakdown does not add up to the total"
    for c in CATEGORIES:
        assert sum(v[c] for v in breakdown.values()) == counts[c], \
            f"sector breakdown does not match summary for {c}"

    thin = Side(style="thin", color="BFBFBF")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    wb = Workbook()
    ws = wb.active
    ws.title = "المعاملات"
    ws.sheet_view.rightToLeft = True

    section_band(ws, 1, "الملخص العام للمعاملات", 2)
    band = ws["A1"]
    band.comment = Comment(
        f"المصدر: {n_files} ملف اكسل لقطاعات الوزارة.\n"
        + ("تم احتساب كل رقم معاملة مرة واحدة (استبعاد السجلات المكررة).\n"
           if dedupe else
           "الارقام محتسبة من كامل السجلات دون استبعاد المكرر.\n")
        + f"عدد السجلات المكررة لنفس رقم المعاملة: {duplicate_records}.\n"
        "توحيد الحالات: مكتمل = مكتملة، على المخطط، متأخر = متأخرة.",
        "Claude",
    )

    labels = {"مكتملة": "المعاملات المكتملة",
              "على المخطط": "المعاملات على المخطط",
              "متأخرة": "المعاملات المتأخرة"}
    summary = [("اجمالي عدد المعاملات", total)] + [
        (labels[c], counts[c]) for c in CATEGORIES
    ]
    for i, (label, value) in enumerate(summary, start=2):
        style_cell(ws.cell(row=i, column=1, value=label), bold=(i == 2),
                   fill=LIGHT, border=border)
        cell = ws.cell(row=i, column=2, value=value)
        style_cell(cell, bold=(i == 2), color=(GREEN if i == 2 else "000000"),
                   align="center", border=border)
        cell.number_format = "0"
        ws.row_dimensions[i].height = 20

    band_row = len(summary) + 3
    section_band(ws, band_row, "تفصيل حالات المعاملات حسب القطاع", 4)

    head_row = band_row + 1
    for j, title in enumerate(["القطاع"] + SECTOR_HEADERS, start=1):
        style_cell(ws.cell(row=head_row, column=j, value=title), bold=True,
                   color=GREEN, fill=GOLD, align="center", border=border)
    ws.row_dimensions[head_row].height = 24

    row_i = head_row + 1
    for n, sector in enumerate(sectors):
        fill = LIGHT if n % 2 == 1 else None
        style_cell(ws.cell(row=row_i, column=1, value=sector), fill=fill,
                   border=border, wrap=True)
        for j, cat in zip(range(2, 5), CATEGORIES):
            value = breakdown[sector][cat]
            cell = ws.cell(row=row_i, column=j, value=value or None)
            style_cell(cell, align="center", fill=fill, border=border)
            cell.number_format = "0"
        ws.row_dimensions[row_i].height = 22
        row_i += 1

    band_row = row_i + 1
    section_band(ws, band_row, "جدول المعاملات المتأخرة", 6)

    head_row = band_row + 1
    headers = ["رقم المعاملة", "موضوع المعاملة", "القطاع",
               "الجهة الوارد منها المعاملة", "تاريخ إنشاء المعاملة",
               "تاريخ الإنجاز المخطط"]
    for j, title in enumerate(headers, start=1):
        cell = ws.cell(row=head_row, column=j, value=title)
        style_cell(cell, bold=True, color=GREEN, fill=GOLD, align="center",
                   wrap=True, border=border)
    ws.row_dimensions[head_row].height = 32

    row_i = head_row + 1
    for n, (_, item) in enumerate(delayed.iterrows()):
        values = [str(item["num"]), item["subject"], item["sector"],
                  item["source"], item["created"], item["planned"]]
        for j, value in enumerate(values, start=1):
            cell = ws.cell(row=row_i, column=j)
            text = str(value).strip()
            if text:
                cell.value = text
                cell.data_type = "s"
            cell.number_format = "@"
            style_cell(cell, border=border,
                       align="right" if j in (2, 3, 4) else "center",
                       wrap=j in (2, 3, 4),
                       fill=LIGHT if n % 2 == 1 else None)
        ws.row_dimensions[row_i].height = 30
        row_i += 1

    for column, width in {"A": 44, "B": 50, "C": 28, "D": 28,
                          "E": 16, "F": 16}.items():
        ws.column_dimensions[column].width = width
    ws.print_options.horizontalCentered = True
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True

    os.makedirs(os.path.dirname(os.path.abspath(output)), exist_ok=True)
    wb.save(output)
    return total, counts, delayed

def verify(output, expected_delayed, expected_sectors):
    ws = load_workbook(output).active
    last = ws.max_row
    col_a = {ws.cell(row=r, column=1).value: r for r in range(1, last + 1)}
    sector_band = col_a["تفصيل حالات المعاملات حسب القطاع"]
    delayed_band = col_a["جدول المعاملات المتأخرة"]
    sector_rows = range(sector_band + 2, delayed_band - 1)

    summary = [ws.cell(row=r, column=2).value
               for r in range(3, sector_band - 1)]
    col_sums = [sum(ws.cell(row=r, column=c).value or 0 for r in sector_rows)
                for c in range(2, 5)]
    sectors_ok = (len(sector_rows) == expected_sectors
                  and col_sums == summary
                  and sum(col_sums) == ws["B2"].value)

    first_data = delayed_band + 2
    rows = last - first_data + 1
    bad = [r for r in range(first_data, last + 1)
           if not isinstance(ws.cell(row=r, column=1).value, str)]
    wrong_font = [
        (r, c) for r in range(1, last + 1) for c in range(1, 7)
        if ws.cell(row=r, column=c).value is not None
        and (ws.cell(row=r, column=c).font.name != FONT
             or ws.cell(row=r, column=c).font.sz != FONT_SIZE)
    ]
    merged = [str(m) for m in ws.merged_cells.ranges]
    print(f"verify: merged cell ranges = {len(merged)} {merged if merged else ''}")
    print(f"verify: sector rows = {len(sector_rows)} (expected {expected_sectors})"
          f" | sector column sums = {col_sums} vs summary {summary}"
          f" -> {'OK' if sectors_ok else 'MISMATCH'}")
    print(f"verify: delayed rows in sheet = {rows} (expected {expected_delayed})"
          f" | non-text ID cells = {len(bad)}"
          f" | cells not {FONT} {FONT_SIZE}pt = {len(wrong_font)}")
    return (sectors_ok and rows == expected_delayed and not bad
            and not wrong_font and not merged)

def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--input-dir", default="/mnt/user-data/uploads")
    parser.add_argument("--output",
                        default="/mnt/user-data/outputs/تقرير_المعاملات_الموحد.xlsx")
    parser.add_argument("--dedupe", action="store_true",
                        help="count each transaction number once")
    args = parser.parse_args()

    df, unmapped, n_files, sectors = read_files(args.input_dir)
    duplicate_records = len(df) - df.num.nunique()
    print(f"\nraw records: {len(df)} | unique numbers: {df.num.nunique()}"
          f" | duplicate records: {duplicate_records}")
    if unmapped:
        sys.exit(f"Unmapped status values, stop and ask the user: {sorted(unmapped)}")
    if args.dedupe:
        conflicts = df.groupby("num").agg(statuses=("status", "nunique"),
                                           sectors=("sector", "nunique"))
        ambiguous = conflicts[(conflicts.statuses > 1) | (conflicts.sectors > 1)]
        if not ambiguous.empty:
            sys.exit("Conflicting duplicate transaction numbers need a dedupe decision: "
                     + ", ".join(ambiguous.index.astype(str)))
        df = df.drop_duplicates(subset="num", keep="first")

    output_dir = os.path.dirname(os.path.abspath(args.output))
    os.makedirs(output_dir, exist_ok=True)
    fd, candidate = tempfile.mkstemp(prefix=".transactions-unverified-",
                                      suffix=".xlsx", dir=output_dir)
    os.close(fd)
    try:
        total, counts, delayed = build(df, sectors, candidate, args.dedupe,
                                       n_files, duplicate_records)
        print(f"\nاجمالي {total} = "
              + " + ".join(f"{c} {counts[c]}" for c in CATEGORIES)
              + f"  -> {'OK' if sum(counts.values()) == total else 'MISMATCH'}")
        print("sector breakdown table: sorted Arabic alphabetical,", len(sectors), "sectors")
        print("delayed by sector:", delayed.sector.value_counts().to_dict())
        if not verify(candidate, len(delayed), len(sectors)):
            sys.exit("Workbook verification failed; output was not published")
        os.replace(candidate, args.output)
        print("saved: " + args.output)
    finally:
        if os.path.exists(candidate):
            os.remove(candidate)

if __name__ == "__main__":
    main()

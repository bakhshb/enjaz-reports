"""Template-preserving helpers and readers for the existing final Excel schemas."""
import copy, datetime as dt, math, re, statistics
import openpyxl
from lxml import etree
from pptx.oxml.ns import qn
from pptx.oxml.xmlchemy import OxmlElement
TASK_SUMMARY_TITLES=('المهام المكتملة','المهام المتأخرة','المهام المعلقة','طلبات الدعم')
SUHAIL_SUMMARY_TITLES=('أبرز التحديثات','المشاريع المتأخرة','التحديات')
MONTHS=['يناير','فبراير','مارس','أبريل','مايو','يونيو','يوليو','أغسطس','سبتمبر','أكتوبر','نوفمبر','ديسمبر']


def val(v):
    if v is None: return ''
    if isinstance(v, (dt.datetime, dt.date)): return v.strftime('%d-%m-%Y')
    if isinstance(v, float) and v.is_integer(): return str(int(v))
    return str(v)


def set_text(shape, value):
    tf=shape.text_frame; text=val(value)
    lines=text.splitlines() or ['']
    while len(tf.paragraphs)<len(lines):
        src=tf.paragraphs[1]._p if len(tf.paragraphs)>1 else tf.paragraphs[0]._p
        tf._txBody.append(copy.deepcopy(src))
    for i,p in enumerate(tf.paragraphs):
        line=lines[i] if i<len(lines) else ''
        if p.runs:
            p.runs[0].text=line
            for r in p.runs[1:]: r.text=''
        elif line:
            p.add_run().text=line


def clear_cell_bullets(cell):
    for p in cell.text_frame.paragraphs:
        pPr=p._p.get_or_add_pPr()
        for child in list(pPr):
            if child.tag in {qn('a:buChar'),qn('a:buAutoNum'),qn('a:buBlip'),qn('a:buNone')}:
                pPr.remove(child)
        pPr.insert_element_before(OxmlElement('a:buNone'),'a:tabLst','a:defRPr','a:extLst')


def set_cell_text(cell,value):
    set_text(cell,value); clear_cell_bullets(cell)


def shape_by_id(slide, sid):
    def walk(items):
        for sh in items:
            if sh.shape_id==sid: return sh
            if getattr(sh,'shape_type',None)==6:
                x=walk(sh.shapes)
                if x: return x
    x=walk(slide.shapes)
    if not x: raise KeyError(f'shape {sid} not found')
    return x


def recursive_text(shape):
    out=[]
    if getattr(shape,'has_text_frame',False) and shape.text.strip(): out.append(shape.text.strip())
    if getattr(shape,'shape_type',None)==6:
        for child in shape.shapes: out.extend(recursive_text(child))
    return out


def slide_texts(slide):
    out=[]
    for sh in slide.shapes: out.extend(recursive_text(sh))
    return out


def resize_table(table, rows_needed, body_template_index=None):
    tbl=table._tbl
    if body_template_index is None:
        body_template_index=max(0,len(table.rows)-1)
    while len(table.rows)<rows_needed:
        src=table.rows[min(body_template_index,len(table.rows)-1)]._tr
        tbl.append(copy.deepcopy(src))
    while len(table.rows)>rows_needed:
        tbl.remove(tbl.tr_lst[-1])


def _body_base_height(table, fixed_rows):
    hs=[table.rows[i].height for i in range(fixed_rows,len(table.rows)) if table.rows[i].height]
    if hs: return int(statistics.median(hs))
    return 520000


def fill_titled_table(shape, rows, title=None, weights=None):
    t=shape.table
    base=_body_base_height(t,2)
    title_h=t.rows[0].height; header_h=t.rows[1].height
    resize_table(t,2+len(rows),2 if len(t.rows)>2 else len(t.rows)-1)
    if title is not None:
        for j,c in enumerate(t.rows[0].cells): set_cell_text(c,title if j==0 else '')
    for i,row in enumerate(rows,2):
        for j,x in enumerate(row): set_cell_text(t.cell(i,j),x)
    if weights is None: weights=[1]*len(rows)
    for i,w in enumerate(weights,2): t.rows[i].height=max(base,int(base*w))
    if title_h: t.rows[0].height=title_h
    if header_h: t.rows[1].height=header_h
    shape.height=sum((r.height or 0) for r in t.rows)


def estimate_lines(text,width_inches,font_pt):
    text=val(text).strip()
    if not text: return 1
    # Arabic at the approved font sizes is typically around 10-13 visible
    # characters per inch; use a conservative estimate so long records move
    # to a continuation page before they become cramped.
    chars_per_line=max(6,int(width_inches*(12 if font_pt<=11 else 10)))
    return sum(max(1,math.ceil(len(part)/chars_per_line)) for part in text.splitlines() or [''])


def row_units(row, table, font_pt=11):
    widths=[c.width/914400 for c in table.columns]
    lines=max(estimate_lines(x,widths[i] if i<len(widths) else 1.5,font_pt) for i,x in enumerate(row))
    return max(1,math.ceil(lines/3))


def slide_index(prs,target):
    for i,sl in enumerate(prs.slides):
        if sl.part is target.part: return i
    raise ValueError('slide is not in presentation')


def clone_slide(prs,source,after=None,shape_predicate=None):
    dest=prs.slides.add_slide(source.slide_layout)
    # Preserve the source root and shape-tree metadata/namespaces as a unit.
    # Rebuilding only its shapes can discard Office compatibility metadata.
    root=copy.deepcopy(source._element)
    # Office tags and creation IDs describe the original object's identity;
    # sharing them between slides can make PowerPoint reject the copy.
    for element in list(root.iter()):
        if etree.QName(element).localname in {'custDataLst','creationId'}:
            element.getparent().remove(element)
    dest._element=root;dest.part._element=root
    dest.__dict__.pop('shapes',None)
    if shape_predicate is not None:
        keep={sh.shape_id for sh in source.shapes if shape_predicate(sh)}
        for sh in list(dest.shapes):
            if sh.shape_id not in keep:sh.element.getparent().remove(sh.element)
    # Copied shapes may contain image or navigation references. Their rIds are
    # local to the source slide and must be recreated in the destination part.
    relationship_ns='{http://schemas.openxmlformats.org/officeDocument/2006/relationships}'
    remapped={}
    for element in dest._element.iter():
        for key,rid in list(element.attrib.items()):
            if not key.startswith(relationship_ns) or not rid:continue
            if rid not in remapped:
                rel=source.part.rels[rid]
                target=rel.target_ref if rel.is_external else rel.target_part
                remapped[rid]=dest.part.relate_to(target,rel.reltype,is_external=rel.is_external)
            element.set(key,remapped[rid])
    if after is not None:
        lst=prs.slides._sldIdLst; el=lst[-1]; lst.remove(el); lst.insert(slide_index(prs,after)+1,el)
    return dest


def remove_slide(prs,slide):
    idx=slide_index(prs,slide); sldId=prs.slides._sldIdLst[idx]; rid=sldId.rId
    prs.slides._sldIdLst.remove(sldId); prs.part.drop_rel(rid)


def tables(slide,cols=None):
    out=[sh for sh in slide.shapes if getattr(sh,'has_table',False)]
    return [sh for sh in out if cols is None or len(sh.table.columns)==cols]


def table_title(shape):
    return next((c.text.strip() for c in shape.table.rows[0].cells if c.text.strip()),'')


def detail_tables(slide,kind):
    marker='المهمة' if kind=='task' else 'اسم المشروع'
    return [sh for sh in tables(slide,6) if len(sh.table.rows)>1
            and sh.table.cell(1,0).text.strip()=='#'
            and sh.table.cell(1,1).text.strip()==marker]


def detail_pattern_slides(prs,kind):
    return [sl for sl in prs.slides if detail_tables(sl,kind)]


def first_table(slide,cols):
    xs=tables(slide,cols)
    if not xs: raise ValueError(f'{cols}-column table pattern missing')
    return xs[0]


def find_slide_with_text(prs,text,require_table=None):
    for sl in prs.slides:
        if text in slide_texts(sl):
            if require_table is None or bool(tables(sl))==require_table: return sl
    raise ValueError(f'slide pattern missing: {text}')


def find_header(ws,labels,start=1):
    labels=[x.strip() for x in labels]
    for r in range(start,ws.max_row+1):
        got=[val(ws.cell(r,c).value).strip() for c in range(1,len(labels)+1)]
        if got==labels:return r
    raise ValueError(f'header not found: {labels}')


def read_block(ws,header_row,cols,stop_titles=()):
    out=[]
    for r in range(header_row+1,ws.max_row+1):
        first=val(ws.cell(r,1).value).strip()
        if first in stop_titles or (first and all(ws.cell(r,c).value is None for c in range(2,cols+1))):break
        row=[ws.cell(r,c).value for c in range(1,cols+1)]
        if any(x not in (None,'') for x in row):out.append(row)
    return out


def detail_sections(ws,cols):
    out=[]; r=1
    expected_task=['#','المهمة','القطاع','تاريخ الانجاز المخطط','الحالة','ملاحظات','الجهات ذات العلاقة']
    expected_proj=['#','اسم المشروع','تاريخ البداية','تاريخ النهاية','الحالة','ما تم حتى تاريخه']
    expected=expected_task if cols==7 else expected_proj
    while r<=ws.max_row:
        got=[val(ws.cell(r,c).value).strip() for c in range(1,cols+1)]
        if got==expected:
            title=val(ws.cell(r-1,1).value).strip(); title=re.sub(r'^مهام مصدر\s*\((.*)\)$',r'\1',title)
            out.append((title,read_block(ws,r,cols)))
        r+=1
    return out


def summary_records(ws,title,headers):
    """Read one titled Excel section; example/empty markers are not records."""
    for r in range(1,ws.max_row):
        heading=val(ws.cell(r,1).value).strip()
        if heading==title or heading.endswith(': '+title):
            got=[val(ws.cell(r+1,c).value).strip() for c in range(1,len(headers)+1)]
            if got!=headers:continue
            out=[]
            for rr in range(r+2,ws.max_row+1):
                first=val(ws.cell(rr,1).value).strip()
                if not first or first=='لا يوجد':break
                if not re.fullmatch(r'\d+(?:\.0)?',first):break
                out.append([ws.cell(rr,c).value for c in range(1,len(headers)+1)])
            return out
    raise ValueError(f'missing summary section: {title}')


def task_data(path):
    wb=openpyxl.load_workbook(path,data_only=True); s=wb['ملخص المهام']
    k=[s.cell(r,2).value for r in range(2,7)]
    hr=find_header(s,['القطاع','مكتملة','على المخطط','متأخر','معلق']); chart=read_sector_counts(s,hr)
    specs=[('#','المهمة','القطاع','ملاحظات')]*3+[('#','المهمة','القطاع','طلب الدعم')]
    summaries={title:summary_records(s,title,list(headers)) for title,headers in zip(TASK_SUMMARY_TITLES,specs)}
    return k,chart,summaries,detail_sections(wb['تفاصيل المهام'],7)


def suhail_data(path):
    wb=openpyxl.load_workbook(path,data_only=True); s=wb['ملخص مشاريع سهيل']
    k=[s.cell(3,c).value for c in range(1,6)]
    hr=find_header(s,['القطاع','مكتملة','على المخطط','متأخر','لم تبدأ']); chart=read_sector_counts(s,hr)
    specs=[['#','المشروع','القطاع','التحديث'],['#','المشروع','القطاع','ملاحظات'],['#','المشروع','القطاع','التحدي']]
    summaries={title:summary_records(s,title,headers) for title,headers in zip(SUHAIL_SUMMARY_TITLES,specs)}
    return k,chart,summaries,detail_sections(wb['تفاصيل مشاريع سهيل'],6)


def read_sector_counts(ws,header):
    records=[]
    for r in range(header+1,ws.max_row+1):
        sector=ws.cell(r,1).value
        if sector in (None,''):break
        values=[ws.cell(r,c).value for c in range(2,6)]
        if any(v is not None and not isinstance(v,(int,float)) for v in values):
            raise ValueError(f'Invalid sector counts at row {r}')
        records.append([sector,*values])
    return records


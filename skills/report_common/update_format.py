"""Source-preserving Suhail update text, shared by report generators."""
import re
from copy import deepcopy
from pptx.oxml.xmlchemy import OxmlElement
from pptx.oxml.ns import qn

DATE = re.compile(r'^(\s*تاريخ التحديث\s*[:：]?\s*(?:\d{1,4}[/.-]\d{1,2}[/.-]\d{1,4}|\d{1,2}\s+[^\W\d_]+(?:\s+[^\W\d_]+)?\s+\d{4})(?:\s*[هـم])?)(?=\s|$|[،؛:])')
MARKER = re.compile(r'^\s*(?:[•●▪◦*-]|\d+[.)\-]|[٠-٩]+[.)\-])\s+')

SUMMARY_UPDATE_TITLES = {'أبرز التحديثات', 'المشاريع المتأخرة'}

def update_column(table):
    """Only Suhail update columns, never task notes or manual challenges."""
    headers = [c.text.strip() for c in table.rows[1].cells]
    if 'ما تم حتى تاريخه' in headers:
        return headers.index('ما تم حتى تاريخه')
    title = next((c.text.strip() for c in table.rows[0].cells if c.text.strip()), '')
    if title in SUMMARY_UPDATE_TITLES:
        return next((i for i, h in enumerate(headers) if h in {'التحديث', 'ملاحظات'}), None)
    return None

def format_table_updates(table, bullet_style, paragraph_items=False):
    column = update_column(table)
    if column is not None:
        for row in list(table.rows)[2:]:
            format_update(row.cells[column], bullet_style, paragraph_items)

def check_update(cell, source, errors, paragraph_items=False):
    parts = update_parts(source, paragraph_items)
    paragraphs = cell.text_frame.paragraphs
    if len(parts) != len(paragraphs):
        errors.append('Suhail update paragraph count differs'); return
    for p, (value, bold_end, bullet) in zip(paragraphs, parts):
        if p.text != value: errors.append('Suhail update text differs')
        props = p._p.get_or_add_pPr()
        if (props.find(qn('a:buChar')) is not None) != bullet:
            errors.append('Suhail update bullets differ')
        if props.get('rtl') != '1': errors.append('Suhail update paragraph is not RTL')
        offset = 0
        for run in p.runs:
            for char in run.text:
                if bool(run.font.bold) != (offset < bold_end):
                    errors.append('Suhail update date/body bold differs')
                offset += 1

def update_parts(text, paragraph_items=False):
    lines = text.replace('\r\n', '\n').replace('\v', '\n').split('\n')
    # Weekly uses explicit markers; Monday uses source update paragraphs.
    if paragraph_items:
        lines = list(map(str.strip, lines))
        while lines and not lines[0]: lines.pop(0)
        while lines and not lines[-1]: lines.pop()
        if not lines: lines = ['']
    marked = [i for i, line in enumerate(lines) if MARKER.match(line)]
    if paragraph_items and lines and DATE.match(lines[0]):
        marked = [i for i, line in enumerate(lines) if i > 0 and line.strip()]
    bullets = set(marked) if len(marked) > 1 else set()
    result = []
    for i, line in enumerate(lines):
        bullet = i in bullets
        value = MARKER.sub('', line, count=1) if bullet or paragraph_items else line
        date = DATE.match(value) if i == 0 else None
        result.append((value, date.end() if date else 0, bullet))
    return result

def display_text(text, paragraph_items=False):
    return '\n'.join(value for value, _, _ in update_parts(text, paragraph_items))

def format_update(cell, bullet_style=None, paragraph_items=False):
    parts = update_parts(cell.text, paragraph_items)
    frame = cell.text_frame
    base = next((deepcopy(r._r.get_or_add_rPr()) for p in frame.paragraphs for r in p.runs), None)
    paragraph_style = deepcopy(frame.paragraphs[0]._p.get_or_add_pPr())
    frame.clear()
    for i, (value, bold_end, bullet) in enumerate(parts):
        p = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
        old = p._p.find(qn('a:pPr'))
        if old is not None: p._p.remove(old)
        props = deepcopy(paragraph_style)
        for node in list(props):
            if node.tag in {qn('a:buNone'), qn('a:buChar'), qn('a:buAutoNum'), qn('a:buBlip')}:
                props.remove(node)
        if bullet:
            if bullet_style is None:
                raise ValueError('Approved bullet style is required for Suhail list items')
            for attr in ('marL', 'marR', 'indent'):
                if bullet_style.get(attr) is not None: props.set(attr, bullet_style.get(attr))
        marker = OxmlElement('a:buChar' if bullet else 'a:buNone')
        if bullet: marker.set('char', '•')
        # Bullets precede tabLst/defRPr/extLst in DrawingML schema order.
        at = next((n for n, el in enumerate(props) if el.tag in {qn('a:tabLst'), qn('a:defRPr'), qn('a:extLst')}), len(props))
        props.insert(at, marker)
        p._p.insert(0, props)
        for text, bold in ((value[:bold_end], True), (value[bold_end:], False)):
            if not text: continue
            run = p.add_run(); run.text = text
            if base is not None: run._r.insert(0, deepcopy(base))
            run.font.bold = bold

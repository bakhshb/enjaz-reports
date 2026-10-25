"""User-approved PowerPoint table direction, font size and column alignment."""
from pptx.util import Pt
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn

TABLE_PT = 11
AGENDA_HEADERS = ['م','جدول الأعمال','المسؤول','المدة الزمنية (بالدقيقة)']

def reverse_table_columns(table):
    """Mirror the grid and merged groups together, retaining field formatting."""
    grid = table._tbl.tblGrid
    for col in reversed(list(grid)):
        grid.append(col)
    for row in table.rows:
        groups = []
        cells = list(row._tr.tc_lst)
        i = 0
        while i < len(cells):
            span = int(cells[i].get('gridSpan', '1'))
            groups.append(cells[i:i + span])
            i += span
        for group in reversed(groups):
            for cell in group:
                row._tr.insert_element_before(cell, 'a:extLst')
                props = cell.find(qn('a:tcPr'))
                if props is not None:
                    left, right = props.find(qn('a:lnL')), props.find(qn('a:lnR'))
                    for node in (left, right):
                        if node is not None: props.remove(node)
                    if left is not None:
                        left.tag = qn('a:lnR'); props.insert(0,left)
                    if right is not None:
                        right.tag = qn('a:lnL'); props.insert(0,right)

def rtl_table(table):
    """Convert LTR storage to RTL without changing column content or widths."""
    if table._tbl.tblPr.get('rtl') not in ('1', 'true'):
        reverse_table_columns(table)
    table._tbl.tblPr.set('rtl', '1')

def set_cell_size(cell, font_pt=TABLE_PT):
    for node in cell._tc.iter():
        if node.tag in {qn('a:rPr'), qn('a:defRPr'), qn('a:endParaRPr')}:
            node.set('sz', str(font_pt * 100))
    for paragraph in cell.text_frame.paragraphs:
        paragraph.font.size = Pt(font_pt)
        for run in paragraph.runs:
            run.font.size = Pt(font_pt)

def column_alignment(header):
    header = ' '.join(header.replace('أ','ا').replace('إ','ا').split())
    if header in {'القطاع','تاريخ الانجاز المخطط','تاريخ البداية','تاريخ النهاية','الحالة','حالة المشروع'}:
        return PP_ALIGN.CENTER
    if header in {'المهمة','المشروع','اسم المشروع','ملاحظات','الملاحظات','ما تم حتى تاريخه','ما تم حتى آخر تحديث','التحديث','التحديثات','طلب الدعم','طلبات الدعم','التحدي','التحديات'}:
        return PP_ALIGN.RIGHT
    return None

def apply_table(table, font_pt=TABLE_PT):
    headers = [c.text for c in table.rows[1].cells] if len(table.rows)>1 else []
    for i,row in enumerate(table.rows):
        for j,cell in enumerate(row.cells):
            set_cell_size(cell,font_pt)
            alignment = column_alignment(headers[j]) if i>=1 and j<len(headers) else None
            if alignment is not None:
                for paragraph in cell.text_frame.paragraphs:
                    paragraph.alignment = alignment
                    paragraph._p.get_or_add_pPr().set('rtl','1')

def all_tables(prs):
    def walk(shapes):
        for shape in shapes:
            if shape.has_table: yield shape.table
            if hasattr(shape,'shapes'): yield from walk(shape.shapes)
    for slide in prs.slides: yield from walk(slide.shapes)

def table_font_size(table, agenda_pt=TABLE_PT):
    return agenda_pt if [c.text.strip() for c in table.rows[0].cells]==AGENDA_HEADERS else TABLE_PT


def set_table_sizes(prs, agenda_pt=TABLE_PT):
    """Apply the approved size and per-column alignment, preserving merged titles."""
    for table in all_tables(prs): apply_table(table,table_font_size(table,agenda_pt))

def set_table_directions(prs):
    for table in all_tables(prs): rtl_table(table)

def check_table_typography(prs, errors, agenda_pt=TABLE_PT):
    for table in all_tables(prs):
        if table._tbl.tblPr.get('rtl') not in ('1', 'true'):
            errors.append('Table direction must be RTL')
        font_pt=table_font_size(table,agenda_pt)
        headers = [c.text for c in table.rows[1].cells] if len(table.rows)>1 else []
        for i,row in enumerate(table.rows):
            for j,cell in enumerate(row.cells):
                expected = column_alignment(headers[j]) if i>=1 and j<len(headers) else None
                for p in cell.text_frame.paragraphs:
                    if expected is not None and p.alignment != expected:
                        errors.append('Table column alignment differs: '+headers[j])
                    for run in p.runs:
                        if run.text and (run.font.size is None or run.font.size.pt!=font_pt):
                            errors.append(f'Table font size must be {font_pt} pt')

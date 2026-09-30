"""Read-only source, package and native-chart checks used by both report gates."""
import hashlib, io, posixpath, re, zipfile
from pathlib import Path, PurePosixPath
from xml.etree import ElementTree as ET
from openpyxl import load_workbook
from openpyxl.utils.cell import range_boundaries
CHART_NS={'c':'http://schemas.openxmlformats.org/drawingml/2006/chart'}
RID='{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id'


def norm(v):
    if v is None:return ""
    if isinstance(v,float) and v.is_integer():v=int(v)
    return re.sub(r"\s+"," ",str(v).replace("\u00a0"," ")).strip()


def sha(path):
    h=hashlib.sha256()
    with open(path,"rb") as f:
        for b in iter(lambda:f.read(1<<20),b""):h.update(b)
    return h.hexdigest()


def header_row(ws,required,max_scan=30):
    for r in range(1,min(ws.max_row,max_scan)+1):
        m={norm(ws.cell(r,c).value):c for c in range(1,ws.max_column+1)}
        if required.issubset(m):return r,m
    return None,None


def read_tasks(path):
    wb=load_workbook(io.BytesIO(Path(path).read_bytes()),read_only=True,data_only=False); ws=wb["تفاصيل المهام"]
    out=[]; r=1
    while r<=ws.max_row:
        m=re.fullmatch(r"مهام مصدر \((.+)\)",norm(ws.cell(r,1).value))
        if not m:r+=1;continue
        sec=norm(m.group(1)); r+=2
        while r<=ws.max_row:
            a=norm(ws.cell(r,1).value)
            if re.fullmatch(r"مهام مصدر \((.+)\)",a):break
            if re.fullmatch(r"\d+(?:\.0)?",a):
                vals=[norm(ws.cell(r,c).value) for c in range(1,7)]
                out.append((sec,*vals))
            r+=1
    return out


def read_suhail(path):
    wb=load_workbook(io.BytesIO(Path(path).read_bytes()),read_only=True,data_only=False); ws=wb["تفاصيل مشاريع سهيل"]
    out=[]; r=1
    while r<=ws.max_row:
        vals=[norm(ws.cell(r,c).value) for c in range(1,7)]
        if vals[0] and not any(vals[1:]) and r<ws.max_row:
            nxt=[norm(ws.cell(r+1,c).value) for c in range(1,7)]
            if nxt[:2]==["#","اسم المشروع"]:
                sec=vals[0]; r+=2
                while r<=ws.max_row:
                    row=[norm(ws.cell(r,c).value) for c in range(1,7)]
                    if not re.fullmatch(r"\d+(?:\.0)?",row[0]):break
                    out.append((sec,*row)); r+=1
                continue
        r+=1
    return out


def package(path,errors):
    try:
        with zipfile.ZipFile(path) as z:
            bad=z.testzip(); names=set(z.namelist())
            if bad:errors.append(f"ZIP CRC failure: {bad}")
            if len(names)!=len(z.namelist()):errors.append('Duplicate package members')
            relns="{http://schemas.openxmlformats.org/package/2006/relationships}"
            for n in names:
                if n.endswith((".xml",".rels")):
                    try:ET.fromstring(z.read(n))
                    except Exception as e:errors.append(f"malformed XML {n}: {e}")
                if n.endswith('.xml'):
                    relpath=posixpath.join(posixpath.dirname(n),'_rels',posixpath.basename(n)+'.rels')
                    ids={r.get('Id') for r in ET.fromstring(z.read(relpath))} if relpath in names else set()
                    for el in ET.fromstring(z.read(n)).iter():
                        if el.tag=='{http://schemas.openxmlformats.org/drawingml/2006/main}pPr':
                            tags=[child.tag.rsplit('}',1)[-1] for child in el]
                            if 'buNone' in tags and any(tag in tags and tags.index(tag)>tags.index('buNone') for tag in ('lnSpc','spcBef','spcAft','buFont','buFontTx')):
                                errors.append(f'invalid bullet property order: {n}')
                        for key,value in el.attrib.items():
                            if key.startswith('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}') and value and value not in ids:
                                errors.append(f'missing relationship ID: {n} -> {value}')
                if not n.endswith(".rels"):continue
                root=ET.fromstring(z.read(n)); p=PurePosixPath(n)
                source=PurePosixPath("") if p.name==".rels" and str(p.parent)=="_rels" else p.parent.parent/p.name[:-5]
                for rel in root.findall(f"{relns}Relationship"):
                    if rel.attrib.get("TargetMode")=="External":errors.append(f"external relationship: {n} -> {rel.attrib.get('Target')}")
                    else:
                        target=rel.attrib.get('Target','').split('#',1)[0]
                        resolved=posixpath.normpath(posixpath.join(str(source.parent),target)).lstrip('/')
                        if resolved not in names:errors.append(f'missing relationship target: {n} -> {target}')
    except Exception as e:errors.append(f"invalid PPTX package: {e}")


def source_metrics(path, sheet, statuses, kpi_cells):
    """Read the final Excel summary independently of the report builder."""
    with open(path,'rb') as stream:
        wb=load_workbook(stream,data_only=True)
        ws=wb[sheet]
        header,columns=header_row(ws,{'القطاع',*statuses},ws.max_row)
        if header is None:raise ValueError(f'{sheet}: sector status header missing')
        records=[]
        for row in range(header+1,ws.max_row+1):
            sector=norm(ws.cell(row,columns['القطاع']).value)
            if not sector:break
            vals=[ws.cell(row,columns[s]).value for s in statuses]
            if any(v is not None and (not isinstance(v,(int,float)) or v<0 or int(v)!=v) for v in vals):
                raise ValueError(f'{sheet}: invalid sector counts at row {row}')
            records.append([sector,*[v if v else None for v in vals]])
        kpis=[ws[cell].value for cell in kpi_cells]
        totals=[sum(r[i+1] or 0 for r in records) for i in range(len(statuses))]
        if kpis!=[sum(totals),*totals]:raise ValueError(f'{sheet}: KPI and sector counts disagree')
        wb.close()
        return records,kpis


def formula_values(wb,formula):
    sheet,area=formula.rsplit('!',1)
    sheet=sheet.strip("'").replace("''", "'")
    c1,r1,c2,r2=range_boundaries(area.replace('$',''))
    if c1!=c2 and r1!=r2:raise ValueError(f'non-linear chart range: {formula}')
    return [wb[sheet].cell(r,c).value for r in range(r1,r2+1) for c in range(c1,c2+1)]


def cached_values(ref):
    cache=next((x for x in ref if x.tag.endswith('Cache')),None)
    if cache is None:raise ValueError('chart cache missing')
    count=cache.find('c:ptCount',CHART_NS)
    if count is None:raise ValueError('chart cache count missing')
    values=[None]*int(count.get('val'))
    seen=set()
    for point in cache.findall('c:pt',CHART_NS):
        idx=int(point.get('idx'))
        if idx in seen or not 0<=idx<len(values):raise ValueError('invalid chart cache index')
        seen.add(idx); values[idx]=point.findtext('c:v',namespaces=CHART_NS)
    return values


def chart_metrics(path, chart_name, records, statuses, errors):
    """Require source Excel, chart caches and referenced embedded cells to agree."""
    try:
        records=records or [[None]*(len(statuses)+1)]
        with zipfile.ZipFile(path) as z:
            root=ET.fromstring(z.read(chart_name))
            rel_path=posixpath.join(posixpath.dirname(chart_name),'_rels',posixpath.basename(chart_name)+'.rels')
            rels=ET.fromstring(z.read(rel_path))
            external=root.find('c:externalData',CHART_NS)
            if external is None:raise ValueError('embedded chart workbook reference missing')
            rel=next((x for x in rels if x.get('Id')==external.get(RID)),None)
            if rel is None or rel.get('TargetMode')=='External':raise ValueError('chart workbook is not embedded')
            book=posixpath.normpath(posixpath.join(posixpath.dirname(chart_name),rel.get('Target'))).lstrip('/')
            wb=load_workbook(io.BytesIO(z.read(book)),data_only=False)
            series=root.findall('.//c:ser',CHART_NS)
            if len(series)!=len(statuses):raise ValueError('wrong number of chart series')
            for i,ser in enumerate(series):
                for tag,expected in [('tx',[statuses[i]]),('cat',[r[0] for r in records]),('val',[r[i+1] for r in records])]:
                    holder=ser.find('c:'+tag,CHART_NS)
                    if holder is None:raise ValueError(f'series {i}: missing {tag}')
                    ref=next((x for x in holder if x.tag.endswith('Ref')),None)
                    if ref is None:raise ValueError(f'series {i}: missing {tag} reference')
                    formula=ref.findtext('c:f',namespaces=CHART_NS)
                    actual=cached_values(ref)
                    cells=formula_values(wb,formula)
                    normalize=lambda seq:[norm(v) if v is not None else '' for v in seq]
                    if normalize(actual)!=normalize(expected):errors.append(f'{chart_name} series {i} {tag}: cache differs from source')
                    if normalize(cells)!=normalize(expected):errors.append(f'{chart_name} series {i} {tag}: embedded workbook differs from source')
            gap=root.find('.//c:dispBlanksAs',CHART_NS)
            if gap is None or gap.get('val')!='gap':errors.append(f'{chart_name}: blanks must display as gaps')
            wb.close()
    except Exception as exc:errors.append(f'{chart_name}: {exc}')


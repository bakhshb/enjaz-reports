# /// script
# requires-python = ">=3.11"
# dependencies = ["python-pptx>=1.0.2", "openpyxl>=3.1.5"]
# ///
"""Mandatory post-build delivery gate for monday-meeting-report.

Run after build_report.py and before delivery. This script normalizes approved
table typography and blocks delivery when the deck is structurally invalid,
uses stale/current-input-mismatched data, loses table titles/headers, or contains
external PowerPoint relationships.
"""
from __future__ import annotations
import argparse, collections, hashlib, io, os, posixpath, re, shutil, tempfile, zipfile
from pathlib import Path, PurePosixPath
from xml.etree import ElementTree as ET
from openpyxl import load_workbook
from openpyxl.utils.cell import range_boundaries
from pptx import Presentation
import sys
sys.path.insert(0,str(Path(__file__).resolve().parent))
import monday_status_styles as status_styles
from pptx.util import Pt
from pptx.oxml.ns import qn
sys.path.insert(0,str(Path(__file__).resolve().parents[2]))
from report_common.validation import (norm, sha, header_row, read_tasks, read_suhail, package, source_metrics, formula_values, cached_values, chart_metrics)

from report_common.update_format import update_parts, display_text
from report_common.pptx_helpers import suhail_data, val

AGENDA_HEADERS=["م","جدول الأعمال","المسؤول","المدة الزمنية (بالدقيقة)"]
TASK_HEADERS=["#","المهمة","القطاع","تاريخ الإنجاز المخطط","الحالة","ملاحظات"]
SUHAIL_HEADERS=["#","اسم المشروع","تاريخ البداية","تاريخ النهاية","الحالة","ما تم حتى تاريخه"]
SUMMARY_HEADERS={
    "المهام المكتملة":["ملاحظات","القطاع","المهمة","#"],
    "المهام المتأخرة":["ملاحظات","القطاع","المهمة","#"],
    "المهام المعلقة":["ملاحظات","القطاع","المهمة","#"],
    "طلبات الدعم":["طلب الدعم","القطاع","المهمة","#"],
    "أبرز التحديثات":["التحديث","القطاع","المشروع","#"],
    "المشاريع المتأخرة":["ملاحظات","القطاع","المشروع","#"],
    "التحديات":["التحدي","القطاع","المشروع","#"],
}
FONT="Abar Mid"; ALLOWED={"Abar Mid","Abar Mid SemiBold"}
AGENDA_PT=14.0; TABLE_PT=11.0




def read_agenda(path):
    wb=load_workbook(io.BytesIO(Path(path).read_bytes()),read_only=True,data_only=False); ws=wb["Sheet1"]
    req={"الترتيب","الموضوع","المتحدث","الوقت المطلوب من المالك"}; hr,c=header_row(ws,req)
    if not hr:raise RuntimeError("topics workbook headers not found")
    out=[]; stop=False
    for r in range(hr+1,ws.max_row+1):
        order=ws.cell(r,c["الترتيب"]).value; topic=norm(ws.cell(r,c["الموضوع"]).value)
        if order in (None,"") or not topic:continue
        try:o=f"{int(float(order)):02d}"
        except:o=norm(order)
        dur=ws.cell(r,c["الوقت المطلوب من المالك"]).value
        if isinstance(dur,float) and dur.is_integer():dur=int(dur)
        out.append((o,topic,norm(ws.cell(r,c["المتحدث"]).value),norm(dur)))
        if topic=="ملخص الاجتماع":stop=True;break
    if not stop:raise RuntimeError("topics workbook missing final item ملخص الاجتماع")
    return out



def normalize_typography(prs):
    for sl in prs.slides:
        for sh in sl.shapes:
            if getattr(sh,"has_table",False):
                tt=sh.table; rr=rows(tt); is_ag=bool(rr and rr[0]==AGENDA_HEADERS); pt=AGENDA_PT if is_ag else TABLE_PT
                detail=len(tt.columns)==6 and len(rr)>1 and rr[1][4]=='الحالة'
                for row_index,rw in enumerate(tt.rows):
                    for column_index,cell in enumerate(rw.cells):
                        # Approved status examples must be checked as generated;
                        # normalization must not hide an incorrect status style.
                        if detail and row_index>=2 and column_index==4:continue
                        for p in cell.text_frame.paragraphs:
                            for run in p.runs:
                                if run.text:
                                    if run.font.name not in ALLOWED:run.font.name=FONT
                                    run.font.size=Pt(pt)
            elif getattr(sh,"has_text_frame",False):
                for p in sh.text_frame.paragraphs:
                    for run in p.runs:
                        if run.text and run.font.name in {"Arial","Calibri","+mn-lt","+mj-lt"}:run.font.name=FONT


def canon(v):
    s=norm(v)
    s=re.sub(r"تاريخ التحديث\s*[:：]\s*","تاريخ التحديث ",s)
    return norm(s)

def canon_record(rec):
    return tuple(canon(x) for x in rec)

def rows(t):return [[norm(c.text) for c in r.cells] for r in t.rows]
def isnum(s):return bool(re.fullmatch(r"\d+(?:\.0)?",norm(s)))


def read_summary(path,sheet,titles):
    ws=load_workbook(io.BytesIO(Path(path).read_bytes()),read_only=True,data_only=False)[sheet]
    result={title:[] for title in titles}
    found=set()
    for r in range(1,ws.max_row):
        heading=norm(ws.cell(r,1).value)
        title=next((t for t in titles if heading==t or heading.endswith(': '+t)),None)
        if title is None:continue
        if ws.cell(r,2).value is not None:continue  # KPI label, not a section band.
        expected=list(reversed(SUMMARY_HEADERS[title]))
        got=[norm(ws.cell(r+1,c).value) for c in range(1,5)]
        if got!=expected:raise ValueError(f'{sheet}: invalid headers for {title}')
        if title in found:raise ValueError(f'{sheet}: duplicate summary section {title}')
        found.add(title)
        for rr in range(r+2,ws.max_row+1):
            first=norm(ws.cell(rr,1).value)
            if not isnum(first):break
            result[title].append(tuple(norm(ws.cell(rr,c).value) for c in range(4,0,-1)))
    if set(titles)!=found:raise ValueError(f'{sheet}: missing summary sections {sorted(set(titles)-found)}')
    return result


def collect(prs,errors,template=status_styles.MASTER):
    ag=[]; task=[]; suh=[]; suh_seq=[]; summaries={title:[] for title in SUMMARY_HEADERS}
    approved=status_styles.examples(template)
    for slide in prs.slides:
        for shape in slide.shapes:
            if not shape.has_table or len(shape.table.columns)!=6:continue
            if len(shape.table.rows)<2 or shape.table.cell(1,4).text.strip()!='الحالة':continue
            for row in list(shape.table.rows)[2:]:
                cell=row.cells[4];status=status_styles.canonical(cell.text)
                allowed=status_styles.TASK_STATUSES if shape.table.cell(1,1).text.strip()=='المهمة' else status_styles.PROJECT_STATUSES
                if status not in allowed:errors.append('Unsupported detail status: '+status)
                if status not in approved:
                    errors.append('Missing approved template status style: '+status)
                elif status_styles.signature(cell)!=status_styles.signature(approved[status]):
                    errors.append('detail status fill or typography differs from approved template: '+status)
    for si,sl in enumerate(prs.slides,1):
        for sh in sl.shapes:
            if not getattr(sh,"has_table",False):continue
            t=sh.table; rr=rows(t); cols=len(t.columns)
            for row in t.rows:
                seen_extension=False
                for child in row._tr:
                    if child.tag==qn("a:extLst"):seen_extension=True
                    elif child.tag==qn("a:tc") and seen_extension:errors.append("Invalid table row XML order")
            if t._tbl.tblPr.get("rtl") not in ("1","true"):errors.append(f"slide {si} table is not RTL")
            if sh.top+sum(r.height for r in t.rows)>prs.slide_height-int(50*12700):errors.append(f"slide {si} table crosses footer clearance")
            if cols==4 and rr and rr[0]==AGENDA_HEADERS:
                ag.extend(tuple(x) for x in rr[1:] if any(x))
            elif cols==6 and len(rr)>=2 and rr[1][:2]==["#","المهمة"]:
                if not any(x.startswith("تفاصيل المهام (") for x in rr[0]):errors.append(f"slide {si} task table missing title row")
                if rr[1]!=TASK_HEADERS:errors.append(f"slide {si} task headers missing/corrupt")
                m=re.fullmatch(r"تفاصيل المهام \((.+)\)",rr[0][0] if rr else ""); sec=norm(m.group(1)) if m else ""
                for x in rr[2:]:
                    if x and isnum(x[0]):task.append((sec,*x[:6]))
            elif cols==6 and len(rr)>=2 and rr[1][:2]==["#","اسم المشروع"]:
                if not rr[0][0]:errors.append(f"slide {si} Suhail table missing sector/title row")
                if rr[1]!=SUHAIL_HEADERS:errors.append(f"slide {si} Suhail headers missing/corrupt")
                sec=re.sub(r"\s*\(تابع\)\s*$","",rr[0][0]).strip()
                if not suh_seq or suh_seq[-1]!=sec:suh_seq.append(sec)
                for x in rr[2:]:
                    if x and isnum(x[0]):suh.append((sec,*x[:6]))
            elif cols==4 and rr:
                title=next((x for x in rr[0] if x in SUMMARY_HEADERS),None)
                if title:
                    if t._tbl.tblPr.get("rtl") in ("1","true"):
                        rr=[rr[0],*[list(reversed(row)) for row in rr[1:]]]
                    if len(rr)<2 or rr[1]!=SUMMARY_HEADERS[title]:errors.append(f"slide {si} {title} headers missing/corrupt")
                    for x in rr[2:]:
                        if x and isnum(x[-1]):summaries[title].append(tuple(x))
                    if not any(x and isnum(x[-1]) for x in rr[2:]):
                        errors.append(f'slide {si} retains empty summary table {title}')
            is_ag=cols==4 and rr and rr[0]==AGENDA_HEADERS; want=AGENDA_PT if is_ag else TABLE_PT
            for rw in t.rows:
                for cell in rw.cells:
                    for p in cell.text_frame.paragraphs:
                        for run in p.runs:
                            if not run.text.strip():continue
                            if run.font.name not in ALLOWED:errors.append(f"slide {si} unapproved font {run.font.name!r}: {run.text[:30]!r}")
                            sz=run.font.size.pt if run.font.size else None
                            if sz is None or abs(sz-want)>.01:errors.append(f"slide {si} wrong table font size {sz!r}; expected {want:g}")
    if ag and ag[-1][1]!="ملخص الاجتماع":errors.append("agenda does not end at ملخص الاجتماع")
    if len(ag)!=len(set(ag)):errors.append("agenda contains duplicate rows")
    if len(suh_seq)!=len(set(suh_seq)):
        errors.append(f"Suhail sector continuation is non-contiguous: {suh_seq!r}")
    return ag,task,suh,summaries

def verify_update_format(prs, expected, errors):
    actual=[]
    for slide in prs.slides:
        shapes=[sh for sh in slide.shapes if sh.has_table]
        for i,a in enumerate(shapes):
            for b in shapes[i+1:]:
                if min(a.left+a.width,b.left+b.width)>max(a.left,b.left) and min(a.top+a.height,b.top+b.height)>max(a.top,b.top):
                    errors.append('Dynamic tables overlap')
        for sh in shapes:
            if len(sh.table.columns)==6 and sh.table.cell(1,1).text.strip()=='اسم المشروع':
                actual.extend(row.cells[5] for row in list(sh.table.rows)[2:])
    if len(actual)!=len(expected):return
    for cell,record in zip(actual,expected):
        parts=update_parts(record[-1],paragraph_items=True)
        paragraphs=cell.text_frame.paragraphs
        if len(parts)!=len(paragraphs):
            errors.append('Suhail update paragraph count differs');continue
        for p,(value,bold_end,bullet) in zip(paragraphs,parts):
            if p.text!=value:errors.append('Suhail update text differs')
            props=p._p.get_or_add_pPr()
            if (props.find(qn('a:buChar')) is not None)!=bullet:errors.append('Suhail update bullets differ')
            if props.get('rtl')!='1':errors.append('Suhail update paragraph is not RTL')
            offset=0
            for run in p.runs:
                for char in run.text:
                    if bool(run.font.bold)!=(offset<bold_end):errors.append('Suhail update date/body bold differs');break
                    offset+=1


def cmp(label,exp,act,errors):
    ce=collections.Counter(canon_record(x) for x in exp); ca=collections.Counter(canon_record(x) for x in act)
    if ce==ca:
        if [canon_record(x) for x in exp]!=[canon_record(x) for x in act]:
            errors.append(f'{label}: source record order changed')
        return
    miss=list((ce-ca).elements()); extra=list((ca-ce).elements())
    if miss:errors.append(f"{label}: missing {len(miss)} current-input record(s); first={miss[0]!r}")
    if extra:errors.append(f"{label}: stale/extra {len(extra)} record(s); first={extra[0]!r}")


CHART_NS={'c':'http://schemas.openxmlformats.org/drawingml/2006/chart'}
REL_NS='{http://schemas.openxmlformats.org/package/2006/relationships}'
RID='{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id'










def verify_metrics(path, prs, tasks_path, suhail_path, errors):
    specs=[('ملخص المهام',tasks_path,['مكتملة','على المخطط','متأخر','معلق'],
            ['B2','B3','B4','B5','B6'],[21,23,25,19,18],'ppt/charts/chart1.xml'),
           ('ملخص مشاريع سهيل',suhail_path,['مكتملة','على المخطط','متأخر','لم تبدأ'],
            ['A3','B3','C3','D3','E3'],[5,8,11,17,13],'ppt/charts/chart2.xml')]
    for title,source,statuses,cells,ids,chart in specs:
        try:
            records,kpis=source_metrics(source,title,statuses,cells)
            details=read_tasks(source) if title=='ملخص المهام' else read_suhail(source)
            source_counts=collections.Counter()
            for row in details:
                sector=row[3] if title=='ملخص المهام' else row[0]
                status=row[5]
                status={'مكتمل':'مكتملة','متأخرة':'متأخر','معلقة':'معلق'}.get(status,status)
                if status not in statuses:raise ValueError(f'unknown detail status {status}')
                source_counts[(sector,status)]+=1
            chart_counts={(norm(row[0]),status):int(row[i+1] or 0) for row in records for i,status in enumerate(statuses)}
            if any(source_counts[key]!=chart_counts.get(key,0) for key in set(source_counts)|set(chart_counts)):
                errors.append(f'{title}: sector summary differs from detail records')
            chart_metrics(path,chart,[row for row in records if any(float(value or 0)>0 for value in row[1:])],statuses,errors)
            candidates=[s for s in prs.slides if any(sh.has_chart for sh in s.shapes)
                        and any(sh.has_text_frame and title in sh.text for sh in s.shapes)]
            if len(candidates)!=1:raise ValueError('expected exactly one summary chart slide')
            def walk(shapes):
                for shape in shapes:
                    yield shape
                    if hasattr(shape,'shapes'):yield from walk(shape.shapes)
            shapes={sh.shape_id:sh for sh in walk(candidates[0].shapes)}
            for sid,value in zip(ids,kpis):
                if sid not in shapes or norm(shapes[sid].text)!=norm(value):
                    errors.append(f'{title}: KPI shape {sid} differs from source')
        except Exception as exc:errors.append(f'{title}: {exc}')

def main(argv=None):
    ap=argparse.ArgumentParser(); ap.add_argument("--report",required=True); ap.add_argument("--topics",required=True); ap.add_argument("--tasks",required=True); ap.add_argument("--suhail",required=True); ap.add_argument("--template",default=str(status_styles.MASTER)); a=ap.parse_args(argv)
    for p in [a.report,a.topics,a.tasks,a.suhail]:
        if not os.path.isfile(p):raise SystemExit(f"missing input: {p}")
    exp_ag=read_agenda(a.topics); exp_task=read_tasks(a.tasks); exp_suh=read_suhail(a.suhail)
    task_titles=list(SUMMARY_HEADERS)[:4];suhail_titles=list(SUMMARY_HEADERS)[4:]
    exp_summaries={**read_summary(a.tasks,"ملخص المهام",task_titles),**read_summary(a.suhail,"ملخص مشاريع سهيل",suhail_titles)}
    report=Path(a.report)
    with tempfile.TemporaryDirectory(prefix=".mm-gate-",dir=report.resolve().parent) as td:
        tmp=Path(td)/report.name; shutil.copy2(report,tmp)
        prs=Presentation(str(tmp)); normalize_typography(prs); prs.save(str(tmp))
        errors=[]; package(tmp,errors); prs2=None
        try:prs2=Presentation(str(tmp)); ag,task,suh,summaries=collect(prs2,errors,a.template)
        except Exception as e:errors.append(f"PowerPoint reopen failed: {e}");ag=task=suh=[];summaries={}
        if ag!=exp_ag:errors.append(f"agenda mismatch: expected={exp_ag!r}; actual={ag!r}")
        cmp("task details",exp_task,task,errors)
        raw_suh=[(section,*row) for section,records in suhail_data(a.suhail)[3] for row in records]
        rendered_suh=[(*row[:-1],display_text(val(raw[-1]),paragraph_items=True)) for row,raw in zip(exp_suh,raw_suh)]
        cmp("Suhail details",rendered_suh,suh,errors)
        if prs2 is not None:verify_update_format(prs2,[(val(r[-1]),) for r in raw_suh],errors)
        for title,expected in exp_summaries.items():cmp(title,expected,summaries.get(title,[]),errors)
        if prs2 is not None:verify_metrics(tmp,prs2,a.tasks,a.suhail,errors)
        print("REPORT_SHA256",sha(tmp)); print("TOPICS_SHA256",sha(a.topics)); print("TASKS_SHA256",sha(a.tasks)); print("SUHAIL_SHA256",sha(a.suhail))
        if errors:
            print(f"QA FAILED: {len(errors)} blocking issue(s)")
            for e in errors:print("ERROR:",e)
            return 1
        status_styles.publish(tmp,report)
    print("PROGRAMMATIC QA PASSED - render and inspect these final bytes before delivery")
    return 0
if __name__=="__main__":raise SystemExit(main())

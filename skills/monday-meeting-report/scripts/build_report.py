#!/usr/bin/env python3
# /// script
# requires-python = ">=3.12"
# dependencies = [
#   "lxml>=5.3",
#   "openpyxl>=3.1",
#   "python-pptx>=1.0",
# ]
# ///
import argparse, copy, datetime as dt, hashlib, math, re, shutil, statistics, tempfile, zipfile, os, subprocess, sys
from pathlib import Path
from lxml import etree
import openpyxl
from pptx import Presentation
sys.path.insert(0,str(Path(__file__).resolve().parent))
import monday_status_styles as status_styles
from pptx.oxml.xmlchemy import OxmlElement
from pptx.oxml.ns import qn
sys.path.insert(0,str(Path(__file__).resolve().parents[2]))
from report_common.pptx_helpers import (val, set_text, clear_cell_bullets, set_cell_text, shape_by_id, recursive_text, slide_texts, resize_table, _body_base_height, fill_titled_table, estimate_lines, row_units, slide_index, clone_slide, remove_slide, tables, table_title, detail_tables, detail_pattern_slides, first_table, find_slide_with_text, find_header, read_block, detail_sections, summary_records, task_data, suhail_data, read_sector_counts)

import monday_layout as layout
from report_common.table_typography import set_table_sizes


NS_C = {'c':'http://schemas.openxmlformats.org/drawingml/2006/chart'}
MONTHS = ['يناير','فبراير','مارس','أبريل','مايو','يونيو','يوليو','أغسطس','سبتمبر','أكتوبر','نوفمبر','ديسمبر']
AGENDA_HEADERS=['م','جدول الأعمال','المسؤول','المدة الزمنية (بالدقيقة)']
MASTER_SHA256='1b7c72de3a5034ee0352a8b6a67f3f519a2e4bad5d95ad93424acdb5353ef255'
TASK_SUMMARY_TITLES=('المهام المكتملة','المهام المتأخرة','المهام المعلقة','طلبات الدعم')
SUHAIL_SUMMARY_TITLES=('أبرز التحديثات','المشاريع المتأخرة','التحديات')




















def fill_agenda_table(shape, rows, weights=None):
    t=shape.table
    base=_body_base_height(t,1); header_h=t.rows[0].height
    resize_table(t,1+len(rows),1 if len(t.rows)>1 else 0)
    for j,x in enumerate(AGENDA_HEADERS): set_cell_text(t.cell(0,j),x)
    for i,row in enumerate(rows,1):
        for j,x in enumerate(row): set_cell_text(t.cell(i,j),x)
    if weights is None: weights=[1]*len(rows)
    for i,w in enumerate(weights,1): t.rows[i].height=max(base,int(base*w))
    if header_h: t.rows[0].height=header_h
    shape.height=sum((r.height or 0) for r in t.rows)






def paginate_rows(rows, shape, fixed_rows=2, font_pt=11, capacity_units=None):
    t=shape.table
    if capacity_units is None: capacity_units=max(1,len(t.rows)-fixed_rows)
    pages=[]; cur=[]; weights=[]; used=0
    for row in rows:
        w=row_units(row,t,font_pt)
        if w>capacity_units:
            raise ValueError('A single row needs %d units but the template page holds %d; '
                             'shorten the source text or use a taller approved template row.'
                             % (w,capacity_units))
        if cur and used+w>capacity_units:
            pages.append((cur,weights)); cur=[]; weights=[]; used=0
        cur.append(row); weights.append(w); used+=w
    if cur or not pages: pages.append((cur,weights))
    return pages




















def fill_summary_tables(prs,source_slide,sections,titles,section_title,measurements,bullets=None):
    patterns={table_title(sh):sh for sh in tables(source_slide) if table_title(sh) in titles}
    start=min(sh.top for sh in patterns.values())
    plans=[]
    for title in titles:
        if sections[title]:
            header,heights=measurements['summary:'+title]
            plans.append((title,sections[title],header,heights))
    pages=layout.pack(plans,start,prs.slide_height-int(55*12700))
    def keep(sh):
        if getattr(sh,'has_table',False):return table_title(sh) in titles
        if getattr(sh,'has_text_frame',False) and sh.text.strip()==section_title:return True
        return getattr(sh,'shape_type',None)==9 and sh.top<int(.65*914400)
    targets=[source_slide];last=source_slide
    for _ in pages[1:]:
        last=clone_slide(prs,source_slide,after=last,shape_predicate=keep);targets.append(last)
    for page_idx,(sl,items) in enumerate(zip(targets,pages)):
        by_title={table_title(sh):sh for sh in tables(sl) if table_title(sh) in titles}
        used_titles=set()
        for title,chunk,header,heights,top in items:
            sh=by_title[title]
            layout.populate(sh,chunk,bullets=bullets)
            layout.size_shape(sh,header,heights,top)
            used_titles.add(title)
        for title,sh in by_title.items():
            if title not in used_titles:sh.element.getparent().remove(sh.element)
        if page_idx:
            for sh in sl.shapes:
                if getattr(sh,'has_text_frame',False) and sh.text.strip()==section_title:
                    set_text(sh,section_title+' (تابع)')
    return targets


def validate_master_patterns(prs):
    agenda=[sh for sl in prs.slides for sh in tables(sl,4) if len(sh.table.rows)>=1]
    task_detail=detail_pattern_slides(prs,'task'); suhail_detail=detail_pattern_slides(prs,'suhail')
    if not agenda: raise ValueError('agenda table pattern missing from master')
    if not task_detail: raise ValueError('task detail table pattern missing from master')
    if not suhail_detail: raise ValueError('Suhail detail table pattern missing from master')
    if not any(getattr(sh,'has_chart',False) for sl in prs.slides for sh in sl.shapes):
        raise ValueError('summary chart patterns missing from master')
    for titles in (TASK_SUMMARY_TITLES,SUHAIL_SUMMARY_TITLES):
        missing=[title for title in titles if not any(table_title(sh)==title for sl in prs.slides for sh in tables(sl))]
        if missing:raise ValueError(f'summary table patterns missing: {missing}')


def header_map(ws,row):
    return {str(c.value).strip():c.column for c in ws[row] if c.value is not None}


def agenda_rows(path):
    ws=openpyxl.load_workbook(path,data_only=True)['Sheet1']; hm=header_map(ws,2); out=[]; found=False
    for r in range(3,ws.max_row+1):
        topic=val(ws.cell(r,hm['الموضوع']).value).strip(); order=ws.cell(r,hm['الترتيب']).value
        if order not in (None,''):
            out.append((f'{int(order):02d}',topic,val(ws.cell(r,hm['المتحدث']).value),val(ws.cell(r,hm['الوقت المطلوب من المالك']).value)))
        if re.sub(r'\s+',' ',topic)=='ملخص الاجتماع': found=True; break
    if not found: raise ValueError('ملخص الاجتماع is missing from topics workbook')
    return out
















def save_chart_payload(rows,path):
    wb=openpyxl.load_workbook(path); ws=wb.active; maxr=max(ws.max_row,len(rows)+1)
    for r in range(2,maxr+1):
        for c in range(1,6):ws.cell(r,c).value=None
    for r,row in enumerate(rows,2):
        for c,x in enumerate(row,1):ws.cell(r,c).value=None if c>1 and x==0 else x
    wb.save(path)


def update_chart_xml(xml_bytes,rows):
    root=etree.fromstring(xml_bytes); root_range=len(rows)+1
    for ser_i,ser in enumerate(root.xpath('.//c:ser',namespaces=NS_C)):
        if ser_i>=4:continue
        # Series-name formulas refer to the header row and must not grow with data.
        for f in ser.xpath('./c:cat//c:f | ./c:val//c:f',namespaces=NS_C):
            f.text=re.sub(r'\$\d+$',f'${root_range}',f.text)
        cats=[val(r[0]) for r in rows]; nums=[r[ser_i+1] if r[ser_i+1] else None for r in rows]
        for cache,values in [(ser.find('.//c:cat//c:strCache',NS_C),cats),(ser.find('.//c:val//c:numCache',NS_C),nums)]:
            if cache is None:continue
            for p in list(cache.findall('c:pt',NS_C)):cache.remove(p)
            pc=cache.find('c:ptCount',NS_C)
            if pc is not None:pc.set('val',str(len(values)))
            for idx,x in enumerate(values):
                if x in (None,''):continue
                pt=etree.SubElement(cache,'{%s}pt'%NS_C['c'],idx=str(idx)); v=etree.SubElement(pt,'{%s}v'%NS_C['c']); v.text=val(x)
    disp=root.find('.//c:dispBlanksAs',NS_C)
    if disp is None:disp=etree.SubElement(root,'{%s}dispBlanksAs'%NS_C['c'])
    disp.set('val','gap')
    return etree.tostring(root,xml_declaration=True,encoding='UTF-8',standalone=True)


def patch_charts(pptx,task_rows,suhail_rows):
    with tempfile.TemporaryDirectory(prefix='.monday-charts-') as temporary:
        tmp=Path(temporary); src=tmp/'src'; src.mkdir()
        with zipfile.ZipFile(pptx) as z:z.extractall(src)
        for chart_name,book_name,rows in [('chart1.xml','Microsoft_Excel_Worksheet.xlsx',task_rows),('chart2.xml','Microsoft_Excel_Worksheet1.xlsx',suhail_rows)]:
            rows=[row for row in rows if any(float(value or 0)>0 for value in row[1:])]
            rows=rows or [[None]*5]  # Valid blank chart range when the source has no sectors.
            save_chart_payload(rows,src/'ppt'/'embeddings'/book_name)
            cp=src/'ppt'/'charts'/chart_name; cp.write_bytes(update_chart_xml(cp.read_bytes(),rows))
        out=tmp/'out.pptx'
        with zipfile.ZipFile(out,'w',zipfile.ZIP_DEFLATED) as z:
            for p in src.rglob('*'):
                if p.is_file():z.write(p,p.relative_to(src))
        shutil.copy2(out,pptx)


def build_dynamic_detail_section(prs,pattern_slides,sections,kind,title_prefix=None,approved_styles=None,measurements=None,bullets=None):
    pattern=max(pattern_slides,key=lambda sl:max(len(sh.table.rows) for sh in detail_tables(sl,kind)))
    pattern_shape=max(detail_tables(pattern,kind),key=lambda sh:len(sh.table.rows))
    plans=[];labels={}
    for i,(section,rows) in enumerate(sections):
        if not rows:continue
        key=f'{kind}:{i}'
        labels[key]=f'{title_prefix} ({section})' if title_prefix else section
        header,heights=measurements[key]
        plans.append((key,[row[:6] for row in rows],header,heights))
    if not plans:return [],list(pattern_slides)
    pages=layout.pack(plans,pattern_shape.top,prs.slide_height-int(55*12700))
    targets=[pattern];last=pattern
    seed=copy.deepcopy(pattern_shape.element)
    for _ in pages[1:]:
        last=clone_slide(prs,pattern,after=last);targets.append(last)
    for sl,items in zip(targets,pages):
        for sh in detail_tables(sl,kind):sh.element.getparent().remove(sh.element)
        for key,chunk,header,heights,top in items:
            element=copy.deepcopy(seed)
            element.nvGraphicFramePr.cNvPr.set('id',str(sl.shapes._next_shape_id))
            sl.shapes._spTree.insert_element_before(element,'p:extLst')
            sh=sl.shapes[-1]
            layout.populate(sh,chunk,labels[key],kind,approved_styles,bullets)
            layout.size_shape(sh,header,heights,top)
    return targets,[sl for sl in pattern_slides if sl.part is not pattern.part]


def parse_args(argv=None):
    ap=argparse.ArgumentParser(); ap.add_argument('--topics',required=True); ap.add_argument('--tasks',required=True); ap.add_argument('--suhail',required=True); ap.add_argument('--output',required=True); ap.add_argument('--template'); ap.add_argument('--layout-engine',choices=['auto','native','estimate'],default=os.environ.get('MONDAY_LAYOUT_ENGINE','auto'))
    return ap.parse_args(argv)


def build_candidate(a):
    template=Path(a.template) if a.template else Path(__file__).resolve().parents[1]/'assets'/'monday-meeting-master.pptx'
    if not template.exists():raise FileNotFoundError(template)
    if not a.template and hashlib.sha256(template.read_bytes()).hexdigest()!=MASTER_SHA256:
        raise ValueError('Bundled Monday meeting template checksum mismatch')

    agenda=agenda_rows(a.topics); tk,tchart,task_summaries,tdetails=task_data(a.tasks); sk,schart,suhail_summaries,sdetails=suhail_data(a.suhail)
    out=Path(a.output); prs=Presentation(template); set_table_sizes(prs); validate_master_patterns(prs)
    approved=status_styles.examples(template)

    # Capture all master patterns before any slide is cloned or removed.
    agenda_slide=next(sl for sl in prs.slides if any(len(sh.table.columns)==4 and sh.table.cell(0,0).text.strip()=='م' for sh in tables(sl)))
    agenda_shape=next(sh for sh in tables(agenda_slide,4) if sh.table.cell(0,0).text.strip()=='م')
    task_summary=find_slide_with_text(prs,'ملخص المهام')
    suhail_summary=find_slide_with_text(prs,'ملخص مشاريع سهيل')
    task_patterns=detail_pattern_slides(prs,'task'); suhail_patterns=detail_pattern_slides(prs,'suhail')

    bullets=layout.bullet_example(prs)
    specs=[]
    for slide,sections,titles in [(task_summary,task_summaries,TASK_SUMMARY_TITLES),(suhail_summary,suhail_summaries,SUHAIL_SUMMARY_TITLES)]:
        patterns={table_title(sh):sh for sh in tables(slide)}
        specs.extend(('summary:'+title,patterns[title],sections[title],None,None) for title in titles)
    for kind,patterns,sections,prefix in [('task',task_patterns,tdetails,'تفاصيل المهام'),('suhail',suhail_patterns,sdetails,None)]:
        shape=max((sh for sl in patterns for sh in detail_tables(sl,kind)),key=lambda sh:len(sh.table.rows))
        for i,(section,rows) in enumerate(sections):
            title=f'{prefix} ({section})' if prefix else section
            specs.append((f'{kind}:{i}',shape,[row[:6] for row in rows],title,kind))
    measurements=layout.measure(specs,template,approved,bullets,getattr(a,'layout_engine','auto'))

    # Cover + agenda. Agenda automatically adds continuation pages when needed.
    set_text(shape_by_id(prs.slides[0],4),f'{dt.datetime.now().astimezone().date().day} {MONTHS[dt.datetime.now().astimezone().date().month-1]} {dt.datetime.now().astimezone().date().year}')
    agenda_pages=paginate_rows(agenda,agenda_shape,1,11)
    agenda_targets=[agenda_slide]; last=agenda_slide
    for _ in range(1,len(agenda_pages)):
        last=clone_slide(prs,agenda_slide,after=last); agenda_targets.append(last)
    for sl,(chunk,weights) in zip(agenda_targets,agenda_pages):fill_agenda_table(first_table(sl,4),chunk,weights)

    # Task summary can expand down the left side; when that area fills, create a
    # clean continuation page containing only the approved heading/table pattern.
    for sid,x in zip([21,23,25,19,18],tk):set_text(shape_by_id(task_summary,sid),x)
    fill_summary_tables(prs,task_summary,task_summaries,TASK_SUMMARY_TITLES,'ملخص المهام',measurements)

    # Task details are data-driven. The master contributes one approved detail
    # pattern; every source and overflow page is generated from it.
    task_generated,task_unused=build_dynamic_detail_section(prs,task_patterns,tdetails,'task','تفاصيل المهام',approved_styles=approved,measurements=measurements,bullets=bullets)

    # Suhail summary and updates.
    for sid,x in zip([5,8,11,17,13],sk):set_text(shape_by_id(suhail_summary,sid),x)
    set_text(shape_by_id(suhail_summary,14),'لم تبدأ')
    fill_summary_tables(prs,suhail_summary,suhail_summaries,SUHAIL_SUMMARY_TITLES,'ملخص مشاريع سهيل',measurements,bullets)

    # Suhail details are also data-driven: new sectors and any amount of projects
    # simply create more pages from the approved six-column detail pattern.
    suhail_generated,suhail_unused=build_dynamic_detail_section(prs,suhail_patterns,sdetails,'suhail',None,approved_styles=approved,measurements=measurements,bullets=bullets)

    # Only after every dynamic page has been allocated do we remove the unused
    # seed-pattern slides from the master.
    for sl in task_unused+suhail_unused:
        remove_slide(prs,sl)

    set_table_sizes(prs)
    prs.save(out); patch_charts(out,tchart,schart)


def validate_candidate(a):
    command=[sys.executable,str(Path(__file__).with_name('report_gate_v2.py')),
             '--report',str(a.output),'--topics',str(a.topics),'--tasks',str(a.tasks),
             '--suhail',str(a.suhail)]
    if a.template:command.extend(['--template',str(a.template)])
    result=subprocess.run(command,capture_output=True,text=True,encoding='utf-8',
                          env={**os.environ,'PYTHONIOENCODING':'utf-8'})
    if result.returncode:
        raise ValueError('Candidate validation failed:\n'+result.stdout+result.stderr)


def main(argv=None):
    args=parse_args(argv)
    output=Path(args.output).resolve()
    template=Path(args.template) if args.template else Path(__file__).resolve().parents[1]/'assets/monday-meeting-master.pptx'
    protected=[template,Path(args.topics),Path(args.tasks),Path(args.suhail)]
    if any(output==p.resolve() or (output.exists() and p.exists() and output.samefile(p)) for p in protected):
        raise ValueError('Output must not overwrite the template or an input workbook')
    output.parent.mkdir(parents=True,exist_ok=True)
    with tempfile.TemporaryDirectory(prefix='.monday-build-',dir=output.parent) as temporary:
        candidate=copy.copy(args)
        candidate.output=str(Path(temporary)/'report.pptx')
        build_candidate(candidate)
        validate_candidate(candidate)
        status_styles.publish(candidate.output,output)
    print(output)
    return 0


if __name__=='__main__':raise SystemExit(main())

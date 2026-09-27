#!/usr/bin/env python3
# /// script
# requires-python = ">=3.12"
# dependencies = [
#   "lxml>=5.3",
#   "python-pptx>=1.0",
# ]
# ///
import argparse, os, re, zipfile
from lxml import etree
from pptx import Presentation

def main():
    ap=argparse.ArgumentParser(); ap.add_argument('--report',required=True); a=ap.parse_args(); errors=[]
    try: prs=Presentation(a.report)
    except Exception as e: raise SystemExit(f'PPTX parse failed: {e}')
    with zipfile.ZipFile(a.report) as z:
        names=set(z.namelist())
        for n in ['ppt/charts/chart1.xml','ppt/charts/chart2.xml','ppt/embeddings/Microsoft_Excel_Worksheet.xlsx','ppt/embeddings/Microsoft_Excel_Worksheet1.xlsx']:
            if n not in names: errors.append(f'missing {n}')
        for n in names:
            if n.endswith('.rels'):
                root=etree.fromstring(z.read(n))
                for rel in root:
                    if rel.get('TargetMode')=='External': errors.append(f'external relationship in {n}: {rel.get("Target")}')
        for n in ['ppt/charts/chart1.xml','ppt/charts/chart2.xml']:
            if n in names:
                root=etree.fromstring(z.read(n)); ns={'c':'http://schemas.openxmlformats.org/drawingml/2006/chart'}
                d=root.find('.//c:dispBlanksAs',ns)
                if d is None or d.get('val')!='gap': errors.append(f'{n} does not use gap for blanks')
    agenda_rows=[]
    headers=['م','جدول الأعمال','المسؤول','المدة الزمنية (بالدقيقة)']
    for sl in prs.slides:
        for sh in sl.shapes:
            if not sh.has_table or len(sh.table.columns)!=4: continue
            t=sh.table
            if [t.cell(0,c).text.strip() for c in range(4)]!=headers: continue
            agenda_rows.extend([t.cell(r,c).text.strip() for c in range(4)] for r in range(1,len(t.rows)))
    if not agenda_rows or agenda_rows[-1][1]!='ملخص الاجتماع': errors.append('agenda does not end with ملخص الاجتماع')
    if errors: raise SystemExit('\n'.join(errors))
    print(f'OK: {os.path.abspath(a.report)}')
if __name__=='__main__': main()
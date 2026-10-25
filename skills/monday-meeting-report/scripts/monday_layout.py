"""Content-sized rows and continuous table flow for the Monday master."""
import copy
import json
import math
import os
from pathlib import Path
import subprocess
import tempfile
import warnings

from pptx import Presentation
from pptx.oxml.ns import qn
from pptx.util import Pt
from report_common.pptx_helpers import fill_titled_table, remove_slide, resize_table, set_cell_text
from report_common.update_format import format_table_updates
from report_common.table_typography import apply_table, rtl_table

AGENDA_PT=14


def bullet_example(prs):
    for slide in prs.slides:
        for shape in slide.shapes:
            if not shape.has_table or len(shape.table.columns) != 6:
                continue
            if shape.table.cell(1,1).text.strip() != 'اسم المشروع':
                continue
            for row in list(shape.table.rows)[2:]:
                for p in row.cells[5].text_frame.paragraphs:
                    props = p._p.get_or_add_pPr()
                    if props.find(qn('a:buChar')) is not None:
                        return copy.deepcopy(props)
    raise ValueError('Suhail master has no approved bullet paragraph')


def populate(shape, rows, title=None, kind=None, styles=None, bullets=None):
    rtl_table(shape.table)
    fixed=1 if kind=='agenda' else 2
    font_pt=AGENDA_PT if kind=='agenda' else 11
    if kind=='agenda':
        resize_table(shape.table,1+len(rows),1)
        for i,values in enumerate(rows,1):
            for j,value in enumerate(values):set_cell_text(shape.table.cell(i,j),value)
        for row in shape.table.rows:
            for cell in row.cells:
                cell.margin_top=cell.margin_bottom=Pt(2)
    else:
        fill_titled_table(shape, rows, title=title)
    for i, values in enumerate(rows, fixed):
        for j, value in enumerate(values):
            cell = shape.table.cell(i, j)
            # Remove surplus template paragraphs, which otherwise create empty lines.
            wanted = max(1, len(str(value if value is not None else '').splitlines()))
            while len(cell.text_frame.paragraphs) > wanted:
                p = cell.text_frame.paragraphs[-1]._p
                p.getparent().remove(p)
        if kind in ('task','suhail'):
            import monday_status_styles as status_styles
            cell = shape.table.cell(i, 4)
            status = status_styles.canonical(cell.text)
            allowed = status_styles.TASK_STATUSES if kind == 'task' else status_styles.PROJECT_STATUSES
            if status not in allowed or status not in styles:
                raise ValueError('Missing or unsupported approved status: ' + status)
            text = cell.text
            status_styles.apply(cell, styles[status])
            set_cell_text(cell, text)
    format_table_updates(shape.table, bullets, paragraph_items=True)
    for row in shape.table.rows:
        for cell in row.cells:
            for p in cell.text_frame.paragraphs:
                p._p.get_or_add_pPr().set('rtl','1')
                for run in p.runs:
                    if run.font.name not in ('Abar Mid', 'Abar Mid SemiBold'):
                        run.font.name = 'Abar Mid'
                    run.font.size = Pt(font_pt)

    apply_table(shape.table,font_pt)


def measure(specs, template, styles, bullets, engine='auto'):
    """Measure fresh disposable tables; no reusable/stale measurement cache."""
    native = engine == 'native' or (engine == 'auto' and os.name == 'nt')
    probe = Presentation(template)
    for slide in list(probe.slides):
        remove_slide(probe, slide)
    estimates = {}
    for key, shape, rows, title, kind in specs:
        if not rows:
            continue
        slide = probe.slides.add_slide(probe.slide_layouts[0])
        for sh in list(slide.shapes):
            sh.element.getparent().remove(sh.element)
        element=copy.deepcopy(shape.element)
        # Match clone_slide: source-local Office tags/IDs cannot follow a table
        # into a fresh slide, where the same rId can mean a slide layout.
        for node in list(element.iter()):
            if node.tag.rsplit('}',1)[-1] in {'custDataLst','creationId'}:
                node.getparent().remove(node)
        slide.shapes._spTree.insert_element_before(element, 'p:extLst')
        sh = slide.shapes[-1]
        sh.name = key
        populate(sh, rows, title, kind, styles, bullets)
        fixed=1 if kind=='agenda' else 2
        scale=(AGENDA_PT if kind=='agenda' else 11)/11
        heights = []
        for row in list(sh.table.rows)[fixed:]:
            height = 20*scale
            for j, cell in enumerate(row.cells):
                width = max(20, sh.table.columns[j].width / 12700 - (cell.margin_left + cell.margin_right) / 12700)
                lines = sum(max(1, math.ceil(len(line) * 4.8 * scale / width)) for line in cell.text.splitlines())
                height = max(height, lines * 14 * scale + (cell.margin_top + cell.margin_bottom) / 12700 + 2)
            heights.append(int(Pt(height)))
            row.height = Pt(1) if native else Pt(height)
        sh.height = sum(r.height for r in sh.table.rows)
        estimates[key] = ([r.height for r in list(sh.table.rows)[:fixed]], heights)
    if not native or not estimates:
        return estimates
    try:
        with tempfile.TemporaryDirectory(prefix='monday-layout-') as temporary:
            path = Path(temporary) / 'probe.pptx'
            probe.save(path)
            result = subprocess.run(['powershell.exe','-NoProfile','-NonInteractive',
                                     '-ExecutionPolicy','Bypass','-File',
                                     str(Path(__file__).with_name('measure_rows.ps1')),str(path)],
                                    capture_output=True, text=True, encoding='utf-8', errors='replace', timeout=60)
            if result.returncode:
                raise ValueError((result.stdout + result.stderr).strip() or f'PowerShell exit {result.returncode}')
            actual = json.loads(path.with_suffix('.json').read_text(encoding='utf-8-sig'))
        if not isinstance(actual,dict) or set(actual)!=set(estimates):
            raise ValueError('PowerPoint measurements do not cover every current table')
        for key,values in actual.items():
            if (not isinstance(values,list) or len(values)!=sum(map(len,estimates[key]))
                    or any(type(v) not in (int,float) or not math.isfinite(v) or v<=0 for v in values)):
                raise ValueError('Invalid PowerPoint row measurements: '+key)
        measured={}
        for key,values in actual.items():
            fixed=len(estimates[key][0])
            measured[key]=([max(int(Pt(v)), estimates[key][0][i]) for i,v in enumerate(values[:fixed])],
                           [max(int(Pt(v + 2)), int(Pt(20))) for v in values[fixed:]])
        return measured
    except (OSError, subprocess.SubprocessError, ValueError, TypeError) as exc:
        message='Native PowerPoint row measurement failed: '+str(exc)
        if engine=='native':
            raise RuntimeError(message) from exc
        warnings.warn(message+'; using conservative estimates. Render and inspect before delivery.', RuntimeWarning)
        return estimates


def pack(sections, start, bottom, gap=int(Pt(9)), continuation=None):
    """Pack whole records, counting a repeated title/header only once per table."""
    continuation = start if continuation is None else continuation
    pages = [[]]
    used = start
    for key, rows, header, heights in sections:
        if len(rows) != len(heights):
            raise ValueError('Row measurements do not match current records')
        chunk, sizes = [], []
        fixed = sum(header)
        for row, height in zip(rows, heights):
            if fixed + height > bottom - continuation:
                raise ValueError('A single record is too tall for an approved page: ' + key)
            needed = height if chunk else fixed + height + (gap if pages[-1] else 0)
            if used + needed > bottom:
                if chunk:
                    pages[-1].append((key, chunk, header, sizes, top))
                    chunk, sizes = [], []
                pages.append([])
                used = continuation
            if not chunk:
                top = used + (gap if pages[-1] else 0)
                used = top + fixed
            chunk.append(row)
            sizes.append(height)
            used += height
        if chunk:
            pages[-1].append((key, chunk, header, sizes, top))
    return pages


def size_shape(shape, header, heights, top):
    shape.top = top
    for row, height in zip(shape.table.rows, [*header, *heights]):
        row.height = height
    shape.height = sum(r.height for r in shape.table.rows)

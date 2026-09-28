#!/usr/bin/env python3
# /// script
# requires-python = ">=3.12"
# dependencies = [
#   "openpyxl>=3.1,<4",
# ]
# ///
"""Build the two-sheet «ملخص مشاريع سهيل» workbook from the raw Suhail export.

Usage:
    uv run build_report.py <source.xlsx> [output.xlsx]

Auto-detects the highest-numbered "Wxx" week for both the status column
(«الحالة الفعلية Wxx») and the progress column («ما تم حتى تاريخه Wxx»).
"""
import re
import sys
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

SRC = sys.argv[1]
OUT = sys.argv[2] if len(sys.argv) > 2 else '/mnt/user-data/outputs/ملخص_مشاريع_سهيل.xlsx'
BLANK_STATUS_DEFAULT = 'على المخطط'
STATUSES = ['مكتملة', 'على المخطط', 'متأخر', 'لم تبدأ']

SUMMARY_SECTORS = [
    'مركز امتثال أعمال',
    'وكالة التوعية وتنمية المهارات',
    'مركز منصة نسك الرقمية',
    'وكالة التعاون الدولي والشراكات',
    'وكالة تطوير الأعمال وخدمة المستفيدين',
    'وكالة التحول الرقمي وتقنية المعلومات',
    'وكالة شؤون الحج',
    'وكالة شؤون العمرة والزيارة',
    'مركز التخطيط الاستراتيجي',
]

DETAIL_SECTORS = [
    'وكالة شؤون الحج',
    'وكالة تطوير الأعمال وخدمة المستفيدين',
    'وكالة التعاون الدولي والشراكات',
    'وكالة التوعية وتنمية المهارات',
    'وكالة التحول الرقمي وتقنية المعلومات',
    'مركز منصة نسك الرقمية',
    'مركز امتثال أعمال',
]

SECT = {
    'مركز امتثال الأعمال': 'مركز امتثال أعمال',
    'وكالة تطوير الاعمال وخدمة المستفيدين': 'وكالة تطوير الأعمال وخدمة المستفيدين',
    'وكالة التوعية و تنمية المهارات': 'وكالة التوعية وتنمية المهارات',
    'وكالة شؤون العمرة': 'وكالة شؤون العمرة والزيارة',
    'وكالة شئون العمرة والزيارة': 'وكالة شؤون العمرة والزيارة',
    'وكالة شئون الحج': 'وكالة شؤون الحج',
}

STAT = {
    'مكتمل': 'مكتملة', 'مكتملة': 'مكتملة', 'منجز': 'مكتملة',
    'على المخطط': 'على المخطط', 'حسب المخطط': 'على المخطط',
    'متأخر': 'متأخر', 'متأخرة': 'متأخر',
    'لم يبدأ': 'لم تبدأ', 'لم تبدأ': 'لم تبدأ', 'لم يتم البدء': 'لم تبدأ',
}

GREEN, GOLD, DARK = '038060', 'D7A562', '164767'
STATUS_FILL = {'على المخطط': GREEN, 'مكتملة': DARK, 'متأخر': 'CFA76C', 'لم تبدأ': '8C8C8C'}
F = 'Abar Mid'
SZ = 10

def norm(v):
    return re.sub(r'\s+', ' ', str(v)).strip() if v is not None else ''

src = openpyxl.load_workbook(SRC, data_only=True)
proj_sheet = next((s for s in src.sheetnames if 'سهيل' in s and 'ملخص' not in s), src.sheetnames[0])
ws = src[proj_sheet]
hdr = [norm(c.value) for c in ws[1]]

def latest(match):
    best, bi = -1, None
    for i, h in enumerate(hdr):
        m = re.search(r'W\s*(\d+)', h)
        if m and match(h) and int(m.group(1)) > best:
            best, bi = int(m.group(1)), i
    return best, bi

w_stat, i_stat = latest(lambda h: 'الحالة الفعلية' in h)
w_done, i_done = latest(lambda h: re.search(r'ما\s*تم\s*حتى', h))
if i_stat is None or i_done is None:
    sys.exit('Could not find «الحالة الفعلية Wxx» and/or «ما تم حتى تاريخه Wxx» columns.')

def col(parts):
    for i, h in enumerate(hdr):
        if all(p in h for p in parts):
            return i + 1
    return None

c_sector = col(['القطاع'])
c_name = col(['اسم المشروع'])
c_start = col(['تاريخ البداية'])
c_end = col(['تاريخ النهاية'])
c_notes = col(['ملاحظات'])
c_challenge = col(['التحدي']) or col(['التحديات'])
if c_notes is None:
    print('warning: no project notes column found; delayed-project notes will be blank', file=sys.stderr)

rows = []
for r in range(2, ws.max_row + 1):
    name = norm(ws.cell(r, c_name).value)
    if not name:
        continue
    sec = norm(ws.cell(r, c_sector).value)
    rows.append({
        'sector': SECT.get(sec, sec),
        'name': name,
        'start': norm(ws.cell(r, c_start).value),
        'end': norm(ws.cell(r, c_end).value),
        'status': STAT.get(norm(ws.cell(r, i_stat + 1).value), BLANK_STATUS_DEFAULT),
        'done': ws.cell(r, i_done + 1).value or '',
        'notes': ws.cell(r, c_notes).value if c_notes else None,
        'challenge': ws.cell(r, c_challenge).value if c_challenge else None,
    })

for x in rows:
    if x['sector'] not in SUMMARY_SECTORS:
        SUMMARY_SECTORS.append(x['sector'])
    if x['sector'] not in DETAIL_SECTORS:
        DETAIL_SECTORS.append(x['sector'])

thin = Side(style='thin', color='D9D9D9')
BORD = Border(left=thin, right=thin, top=thin, bottom=thin)
CEN = Alignment(horizontal='center', vertical='center', wrap_text=True, readingOrder=2)
RIG = Alignment(horizontal='right', vertical='center', wrap_text=True, readingOrder=2)

out = openpyxl.Workbook()

def head(sheet, row, labels):
    for j, lab in enumerate(labels, start=1):
        c = sheet.cell(row, j, lab)
        c.font = Font(name=F, bold=True, size=SZ, color='FFFFFF')
        c.fill = PatternFill('solid', fgColor=GREEN)
        c.alignment = CEN
        c.border = BORD
    sheet.row_dimensions[row].height = 26

def title(sheet, row, text, span):
    sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=span)
    c = sheet.cell(row, 1, text)
    c.font = Font(name=F, bold=True, size=SZ, color=DARK)
    c.alignment = Alignment(horizontal='right', vertical='center', readingOrder=2)
    sheet.row_dimensions[row].height = 24

s1 = out.active
s1.title = 'ملخص مشاريع سهيل'
s1.sheet_view.rightToLeft = True
s1.sheet_view.showGridLines = False

r = 1
title(s1, r, 'أولا: الملخص العام لمشاريع سهيل', 5)
r += 1
kpis = [('إجمالي مشاريع سهيل', len(rows))] + [
    (lab, sum(1 for x in rows if x['status'] == st))
    for lab, st in [('المشاريع المكتملة', 'مكتملة'),
                    ('المشاريع على المخطط', 'على المخطط'),
                    ('المشاريع المتأخرة', 'متأخر'),
                    ('المشاريع لم تبدأ', 'لم تبدأ')]]
for j, (lab, val) in enumerate(kpis, start=1):
    c = s1.cell(r, j, lab)
    c.font = Font(name=F, bold=True, size=SZ, color='FFFFFF')
    c.fill = PatternFill('solid', fgColor=GREEN)
    c.alignment = CEN
    c.border = BORD
    v = s1.cell(r + 1, j, val)
    v.font = Font(name=F, bold=True, size=SZ, color=DARK)
    v.alignment = CEN
    v.border = BORD
s1.row_dimensions[r].height = 30
s1.row_dimensions[r + 1].height = 28
r += 3

title(s1, r, 'ثانيا: تفصيل حالات المشاريع حسب القطاع', 5)
r += 1
head(s1, r, ['القطاع'] + STATUSES)
r += 1
for sec in SUMMARY_SECTORS:
    c = s1.cell(r, 1, sec)
    c.font = Font(name=F, size=SZ)
    c.alignment = RIG
    c.border = BORD
    for j, st in enumerate(STATUSES, start=2):
        n = sum(1 for x in rows if x['sector'] == sec and x['status'] == st)
        cc = s1.cell(r, j, n if n else None)
        cc.font = Font(name=F, size=SZ, bold=True)
        cc.alignment = CEN
        cc.border = BORD
    s1.row_dimensions[r].height = 22
    r += 1
r += 1

def simple_table(r, heading, status, include_notes=False):
    width = 4 if include_notes else 3
    title(s1, r, heading, width)
    r += 1
    head(s1, r, ['#', 'المشروع', 'القطاع'] + (['ملاحظات'] if include_notes else []))
    r += 1
    sel = [x for x in rows if x['status'] == status]
    if not sel:
        s1.merge_cells(start_row=r, start_column=1, end_row=r, end_column=width)
        c = s1.cell(r, 1, 'لا يوجد')
        c.font = Font(name=F, size=SZ, italic=True)
        c.alignment = CEN
        c.border = BORD
        return r + 2
    for i, x in enumerate(sel, start=1):
        s1.cell(r, 1, i).alignment = CEN
        s1.cell(r, 2, x['name']).alignment = RIG
        s1.cell(r, 3, x['sector']).alignment = RIG
        if include_notes:
            s1.cell(r, 4, x['notes'] if x['notes'] is not None else '').alignment = RIG
        for j in range(1, width + 1):
            s1.cell(r, j).font = Font(name=F, size=SZ)
            s1.cell(r, j).border = BORD
        s1.row_dimensions[r].height = 70 if include_notes else 24
        r += 1
    return r + 1

r = simple_table(r, 'ثالثا: المشاريع المكتملة', 'مكتملة')
r = simple_table(r, 'رابعا: المشاريع المتأخرة', 'متأخر', include_notes=True)

title(s1, r, 'خامسا: أبرز التحديثات', 4)
r += 1
head(s1, r, ['#', 'المشروع', 'القطاع', 'التحديث'])
r += 1
upd_name = next((s for s in src.sheetnames if 'التحديثات' in s), None)
i = 1
if upd_name:
    for row in src[upd_name].iter_rows(min_row=2, values_only=True):
        if not row or not row[1]:
            continue
        sec = norm(row[2])
        s1.cell(r, 1, i).alignment = CEN
        s1.cell(r, 2, norm(row[1])).alignment = RIG
        s1.cell(r, 3, SECT.get(sec, sec)).alignment = RIG
        s1.cell(r, 4, str(row[3]).strip() if row[3] else '').alignment = RIG
        for j in range(1, 5):
            s1.cell(r, j).font = Font(name=F, size=SZ)
            s1.cell(r, j).border = BORD
        s1.row_dimensions[r].height = 70
        i += 1
        r += 1

r += 1
title(s1, r, 'سادسا: التحديات', 4)
r += 1
head(s1, r, ['#', 'المشروع', 'القطاع', 'التحدي'])
r += 1
challenge_sheet = next((s for s in src.sheetnames if 'تحدي' in s and s != proj_sheet), None)
challenges = []
if challenge_sheet:
    challenge_ws = src[challenge_sheet]
    challenge_headers = [norm(c.value) for c in challenge_ws[1]]
    def challenge_col(names):
        return next((i for i, h in enumerate(challenge_headers) if any(name in h for name in names)), None)
    i_ch_name = challenge_col(['اسم المشروع', 'المشروع'])
    i_ch_sector = challenge_col(['القطاع'])
    i_ch_text = challenge_col(['التحدي', 'التحديات'])
    if None in (i_ch_name, i_ch_sector, i_ch_text):
        sys.exit('Challenges sheet needs project, sector, and challenge columns.')
    for values in challenge_ws.iter_rows(min_row=2, values_only=True):
        if not values or not values[i_ch_text]:
            continue
        sector = norm(values[i_ch_sector])
        challenges.append((norm(values[i_ch_name]), SECT.get(sector, sector), values[i_ch_text]))
elif c_challenge:
    challenges = [(x['name'], x['sector'], x['challenge']) for x in rows if x['challenge']]
else:
    print('warning: no challenges sheet or challenge column found; challenges table is empty', file=sys.stderr)

if challenges:
    for i, (name, sector, challenge) in enumerate(challenges, start=1):
        for j, value in enumerate((i, name, sector, challenge), start=1):
            cell = s1.cell(r, j, value)
            cell.font = Font(name=F, size=SZ)
            cell.alignment = CEN if j == 1 else RIG
            cell.border = BORD
        s1.row_dimensions[r].height = 70
        r += 1
else:
    s1.merge_cells(start_row=r, start_column=1, end_row=r, end_column=4)
    cell = s1.cell(r, 1, 'لا يوجد')
    cell.font = Font(name=F, size=SZ, italic=True)
    cell.alignment = CEN
    cell.border = BORD

for c_, w in zip('ABCDE', [34, 34, 34, 60, 16]):
    s1.column_dimensions[c_].width = w

s2 = out.create_sheet('تفاصيل مشاريع سهيل')
s2.sheet_view.rightToLeft = True
s2.sheet_view.showGridLines = False
ORDER = {'متأخر': 0, 'على المخطط': 1, 'لم تبدأ': 2, 'مكتملة': 3}

r = 1
for sec in DETAIL_SECTORS:
    sel = [x for x in rows if x['sector'] == sec]
    if not sel:
        continue
    sel = [x for _, x in sorted(enumerate(sel),
                                key=lambda t: (ORDER.get(t[1]['status'], 4), t[0]))]
    title(s2, r, sec, 6)
    r += 1
    head(s2, r, ['#', 'اسم المشروع', 'تاريخ البداية', 'تاريخ النهاية', 'الحالة', 'ما تم حتى تاريخه'])
    r += 1
    for i, x in enumerate(sel, start=1):
        s2.cell(r, 1, i).alignment = CEN
        s2.cell(r, 2, x['name']).alignment = RIG
        s2.cell(r, 3, x['start']).alignment = CEN
        s2.cell(r, 4, x['end']).alignment = CEN
        s2.cell(r, 5, x['status']).alignment = CEN
        s2.cell(r, 6, str(x['done']).strip()).alignment = RIG
        for j in range(1, 7):
            s2.cell(r, j).font = Font(name=F, size=SZ)
            s2.cell(r, j).border = BORD
        s2.cell(r, 5).fill = PatternFill('solid', fgColor=STATUS_FILL[x['status']])
        s2.cell(r, 5).font = Font(name=F, size=SZ, bold=True, color='FFFFFF')
        s2.row_dimensions[r].height = 78
        r += 1
    r += 1

for c_, w in zip('ABCDEF', [5, 40, 15, 15, 15, 75]):
    s2.column_dimensions[c_].width = w

out.save(OUT)

counts = {st: sum(1 for x in rows if x['status'] == st) for st in STATUSES}
print(f'week used: W{w_stat} (status) / W{w_done} (progress)')
print(f'projects: {len(rows)}  {counts}')
print(f'saved: {OUT}')

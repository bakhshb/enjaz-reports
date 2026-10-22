"""Preserve and validate inherited weekly detail artwork and navigation."""
from collections import Counter
from copy import deepcopy
from hashlib import sha256
from lxml import etree
from pptx.oxml.ns import qn
from pptx.oxml.xmlchemy import OxmlElement
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from report_common import pptx_helpers as helpers

NAVIGATION = {'ملخص المهام': 'إجمالي المهام',
              'ملخص مشاريع سهيل': 'إجمالي مشاريع سهيل',
              'المعاملات': 'إجمالي العاملات', 'التفاصيل': None}
DIGITAL_NAME = 'التجربة الرقمية'

def label(text):
    return ' '.join(text.split())

def detail_kind(slide):
    return next((kind for kind in ('task', 'project') if helpers.detail_tables(slide, kind)), None)

def all_shapes(shapes):
    for shape in shapes:
        yield shape
        if hasattr(shape, 'shapes'):
            yield from all_shapes(shape.shapes)

def next_shape_id(slide):
    # Include group descendants, not just the top-level group IDs.
    return max((int(node.get('id')) for node in slide._element.iter(qn('p:cNvPr'))), default=0) + 1

def targets(prs):
    result = {}
    for name, marker in NAVIGATION.items():
        matches = [sl for sl in prs.slides if (marker in helpers.slide_texts(sl) if marker else
                   {'تفاصيل المهام', 'تفاصيل مشاريع سهيل'}.issubset(helpers.slide_texts(sl)))]
        if len(matches) != 1:
            raise ValueError('Cannot identify navigation destination: ' + name)
        result[name] = matches[0].part
    digital = [sl for sl in prs.slides if sum(sh.has_chart for sh in sl.shapes) == 5]
    if digital:
        if len(digital) != 1:
            raise ValueError('Cannot identify digital-maturity navigation destination')
        result[DIGITAL_NAME] = digital[0].part
    return result

def shape_collections(shapes):
    yield shapes
    for shape in shapes:
        if hasattr(shape, 'shapes'):
            yield from shape_collections(shape.shapes)

def buttons(layout, names=None):
    """Return the label and its matching background, without the full-page fill."""
    for name in (names or NAVIGATION):
        matches = [(sh, siblings) for siblings in shape_collections(layout.shapes)
                   for sh in siblings if getattr(sh, 'has_text_frame', False) and label(sh.text) == name]
        if len(matches) != 1:
            raise ValueError('Missing/duplicate layout navigation label: ' + name)
        text, siblings = matches[0]
        backgrounds = [sh for sh in siblings if sh != text and
                       abs(sh.left-text.left) < 12700 and abs(sh.width-text.width) < 12700
                       and sh.top <= text.top and sh.top+sh.height >= text.top+text.height]
        if len(backgrounds) != 1:
            raise ValueError('Missing/duplicate navigation background: ' + name)
        yield name, (text, backgrounds[0])

def prepare_navigation(prs):
    destinations = targets(prs)
    layouts = {sl.slide_layout.part: sl.slide_layout for sl in prs.slides
               if detail_kind(sl) or (DIGITAL_NAME in destinations and sl.part == destinations[DIGITAL_NAME])}
    for layout in layouts.values():
        for name, shapes in buttons(layout, destinations):
            rid = layout.part.relate_to(destinations[name], RT.SLIDE)
            for shape in shapes:
                props = shape.element.find('.//' + qn('p:cNvPr'))
                for old in list(props.findall(qn('a:hlinkClick'))):
                    props.remove(old)
                link = OxmlElement('a:hlinkClick')
                link.set(qn('r:id'), rid)
                link.set('action', 'ppaction://hlinksldjump')
                props.insert(0, link)

def layer_signature(layer):
    """Compare artwork, geometry and image bytes; links are validated separately."""
    root = deepcopy(layer._element)
    for node in list(root.iter()):
        if etree.QName(node).localname in {'hlinkClick', 'hlinkMouseOver', 'creationId'}:
            node.getparent().remove(node)
    for node in root.iter():
        for key, rid in list(node.attrib.items()):
            if key.startswith('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}'):
                rel = layer.part.rels[rid]
                value = rel.target_ref if rel.is_external else (
                    'slideLayout' if rel.reltype == RT.SLIDE_LAYOUT else sha256(rel.target_part.blob).hexdigest())
                node.set(key, value)
    root.attrib.pop('showMasterSp', None)
    return etree.tostring(root, method='c14n')

def check_detail_artwork(prs, master, errors):
    patterns = {kind: helpers.detail_pattern_slides(master, kind)[0] for kind in ('task', 'project')}
    destinations = targets(prs)
    if DIGITAL_NAME in destinations:
        digital = next(sl for sl in prs.slides if sl.part == destinations[DIGITAL_NAME])
        try:
            for name, shapes in buttons(digital.slide_layout, destinations):
                for shape in shapes:
                    links = list(shape.element.iter(qn('a:hlinkClick')))
                    if len(links) != 1 or digital.slide_layout.part.rels[links[0].get(qn('r:id'))].target_part != destinations[name]:
                        errors.append('Digital slide navigation differs: ' + name)
        except (KeyError, ValueError) as exc:
            errors.append('Digital slide navigation is invalid: ' + str(exc))
    checked = set()
    for index, slide in enumerate(prs.slides, 1):
        kind = detail_kind(slide)
        if kind is None:
            continue
        prefix = f'Slide {index}: '
        if slide._element.get('showMasterSp', '1') not in ('1', 'true'):
            errors.append(prefix + 'inherited layout artwork is hidden')
        ids = [sh.shape_id for sh in all_shapes(slide.shapes)]
        if any(count > 1 for count in Counter(ids).values()):
            errors.append(prefix + 'duplicate shape IDs, including group descendants')
        layout = slide.slide_layout
        expected = patterns[kind].slide_layout
        if (layout.part, kind) in checked:
            continue
        checked.add((layout.part, kind))
        if layout._element.get('showMasterSp', '1') not in ('1', 'true'):
            errors.append(prefix + 'master artwork is hidden by layout')
        try:
            for actual, reference in ((layout, expected), (layout.slide_master, expected.slide_master)):
                if layer_signature(actual) != layer_signature(reference):
                    errors.append(prefix + 'layout/master artwork differs from approved template')
            for name, shapes in buttons(layout, destinations):
                for shape in shapes:
                    links = list(shape.element.iter(qn('a:hlinkClick')))
                    if len(links) != 1:
                        errors.append(prefix + 'navigation link missing/duplicated: ' + name)
                        continue
                    link = links[0]
                    rel = layout.part.rels[link.get(qn('r:id'))]
                    if rel.is_external or rel.target_part != destinations[name] or link.get('action') != 'ppaction://hlinksldjump':
                        errors.append(prefix + 'navigation destination differs: ' + name)
        except (KeyError, ValueError) as exc:
            errors.append(prefix + 'invalid inherited artwork/navigation: ' + str(exc))

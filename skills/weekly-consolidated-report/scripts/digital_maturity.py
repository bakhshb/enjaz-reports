# /// script
# requires-python = ">=3.12"
# dependencies = ["pdfplumber>=0.11", "python-pptx>=1.0.2", "openpyxl>=3.1", "lxml>=5.3"]
# ///
"""Read the digital-maturity PDF summary and fill its approved weekly slide."""
import argparse
import hashlib
import json
import re
import sys
from collections import defaultdict
from copy import deepcopy
from pathlib import Path

import pdfplumber
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Pt

sys.path.insert(0, str(Path(__file__).resolve().parents[2]))
from report_common import pptx_helpers as helpers
from weekly_data import content_bottom, set_cell_exact


MASTER = Path(__file__).resolve().parents[1] / 'assets/weekly-report-digital-master.pptx'
MASTER_HASH = '4212057961c371b3a63b8a5676ebe42582ab4d5db98fff983168f0fe610a689c'
CALCULATION_NOTE = ('آلية الاحتساب: 10% للجاهزية وتشمل المتطلبات والاجتماعات والتنسيق، '
                    'و5% لتحديد التواريخ المستهدفة، والنسبة المتبقية وفق الإنجاز الفعلي لمتطلبات الدليل.')
CALCULATION_NOTE_ID = 7
AXES = ('experience', 'satisfaction', 'complaints', 'tools')
LABELS = ('تجربة المستخدم', 'رضا المستفيد', 'التعامل مع الشكاوى', 'التقنيات والأدوات')
DETAIL_CARD_IDS = dict(zip(AXES, (6, 8, 10, 12)))
DETAIL_CARD_BOXES = {
    'experience': (405, 50, 535, 115),
    'satisfaction': (272, 50, 405, 115),
    'complaints': (140, 50, 272, 115),
    'tools': (5, 50, 140, 115),
}
COLORS = {
    'experience': (0.898, 0.737, 0.553),
    'satisfaction': (0.220, 0.435, 0.404),
    'complaints': (0.706, 0.482, 0.373),
    'tools': (0.557, 0.729, 0.859),
}
GRAY = (0.898, 0.890, 0.878)
CARDS = {
    'experience': (27, 31, 45, 47),
    'satisfaction': (5, 9, 21, 23),
    'complaints': (73, 77, 89, 91),
    'tools': (51, 55, 67, 69),
}


def _rtl(value):
    """This PDF stores Arabic glyphs in reverse order within each visual line."""
    return ' '.join(line[::-1].strip() for line in (value or '').splitlines()).strip()


def _percent(value):
    match = re.fullmatch(r'\s*(?:%\s*([\d.,]+)|([\d.,]+)\s*%)\s*', value or '')
    if not match:
        raise ValueError('Cannot read a digital-maturity percentage: ' + repr(value))
    number = float((match.group(1) or match.group(2)).replace(',', '.'))
    if not 0 <= number <= 100:
        raise ValueError('Digital-maturity percentage outside 0–100')
    return int(number) if number.is_integer() else number


def _chart_number(value):
    return int(value) if float(value).is_integer() else value


def _one_percent(page, box):
    text = page.crop(box).extract_text() or ''
    matches = re.findall(r'%\s*[\d.,]+|[\d.,]+\s*%', text)
    if len(matches) != 1:
        raise ValueError('Expected one percentage in a digital-maturity card')
    return _percent(matches[0])


def _axis_from_fill(color):
    if not isinstance(color, (tuple, list)) or len(color) != 3:
        return None
    for axis, expected in COLORS.items():
        if max(abs(float(a) - b) for a, b in zip(color, expected)) < .025:
            return axis
    if max(abs(float(a) - b) for a, b in zip(color, GRAY)) < .025:
        return 'note'
    return None


def extract_pdf(path):
    """Extract summary and six detail drafts; every Arabic field needs visual review."""
    path = Path(path)
    with pdfplumber.open(path) as pdf:
        if not pdf.pages:
            raise ValueError('Digital-maturity PDF has no pages')
        page = pdf.pages[0]
        if not 530 <= page.width <= 550 or not 950 <= page.height <= 975:
            raise ValueError('Digital-maturity PDF page-one layout differs from the approved source')
        metrics = {
            'satisfaction': _one_percent(page, (25, 75, 270, 270)),
            'experience': _one_percent(page, (270, 75, 525, 270)),
            'tools': _one_percent(page, (25, 280, 270, 470)),
            'complaints': _one_percent(page, (270, 280, 525, 470)),
        }
        bars = defaultdict(list)
        for rect in page.rects:
            if 515 <= rect['top'] <= 640 and rect['width'] > 20 and 10 <= rect['height'] <= 20:
                axis = _axis_from_fill(rect['non_stroking_color'])
                if axis:
                    bars[round(rect['top'])].append((axis, rect))
        if not 1 <= len(bars) <= 5:
            raise ValueError('Digital-maturity sector bars are missing or exceed template capacity')
        sectors = []
        for top, segments in sorted(bars.items()):
            rect = segments[0][1]
            name = _rtl(page.crop((372, rect['top']-2, page.width, rect['bottom']+2)).extract_text())
            if not name:
                raise ValueError('Digital-maturity sector label is missing')
            values = {axis: None for axis in AXES}
            note = None
            for axis, segment in segments:
                label = page.crop((segment['x0'], segment['top'], segment['x1'], segment['bottom'])).extract_text()
                if axis == 'note':
                    note = _rtl(label)
                else:
                    if values[axis] is not None:
                        raise ValueError('Duplicate digital-maturity axis in one sector')
                    values[axis] = _percent(label)
            if note and any(value is not None for value in values.values()):
                raise ValueError('Digital-maturity sector has both a note and percentages')
            sectors.append({'name': name, 'metrics': values, 'note': note})
        tables = [table for table in page.extract_tables()
                  if len(table) >= 2 and len(table[0]) == 3 and _rtl(table[0][0]) == 'التحدي']
        if len(tables) != 1:
            raise ValueError('Cannot identify the digital-maturity challenges table')
        challenges = []
        for index, row in enumerate(tables[0][1:], 1):
            if len(row) != 3 or (row[2] or '').strip() != str(index):
                raise ValueError('Digital-maturity challenge numbering is inconsistent')
            perspective, text = _rtl(row[1]), _rtl(row[0])
            if not perspective or not text:
                raise ValueError('Digital-maturity challenge is incomplete')
            challenges.append({'perspective': perspective, 'text': text})
        if not 1 <= len(challenges) <= 4:
            raise ValueError('Digital-maturity challenges exceed the approved slide capacity')
        top_text = page.crop((0, 0, page.width, 75)).extract_text() or ''
        date = re.search(r'14\d\d/\d{1,2}/\d{1,2}', top_text)
        if not date:
            raise ValueError('Digital-maturity PDF update date is missing')
        if len(pdf.pages) != 7:
            raise ValueError('The approved digital detail master requires exactly seven PDF pages')
        detail_pages = []
        for number, detail in enumerate(pdf.pages[1:], 2):
            if not 530 <= detail.width <= 550 or not 950 <= detail.height <= 975:
                raise ValueError(f'Digital-maturity PDF page {number} has an unapproved layout')
            found = detail.find_tables()
            if len(found) != 1:
                raise ValueError(f'Expected one detail table on PDF page {number}')
            raw_table = found[0].extract()
            if len(raw_table[0]) not in (4, 5) or _rtl(raw_table[0][0]) != 'أساس التقييم':
                raise ValueError(f'Unexpected detail headers on PDF page {number}')
            title_lines = (detail.crop((270, 0, 540, 50)).extract_text() or '').splitlines()
            if len(title_lines) < 2:
                raise ValueError(f'Digital-maturity detail title is missing on PDF page {number}')
            draft_rows = [[_rtl(value) for value in row[:4]] for row in raw_table[1:]]
            draft_rows = [row for row in draft_rows if any(row)]
            if not draft_rows:
                raise ValueError(f'Digital-maturity detail table is empty on PDF page {number}')
            cards = {}
            for axis, box in DETAIL_CARD_BOXES.items():
                lines = (detail.crop(box).extract_text() or '').splitlines()
                if len(lines) < 2:
                    raise ValueError(f'Digital-maturity detail card is missing on PDF page {number}: {axis}')
                raw_value = lines[1].strip()
                cards[axis] = f'{_percent(raw_value):g}%' if '%' in raw_value else _rtl(raw_value)
            below = detail.crop((0, found[0].bbox[3], detail.width, 900)).extract_text() or ''
            draft_challenges = [_rtl(line.replace('•', '').strip()) for line in below.splitlines() if '•' in line]
            detail_pages.append({'page': number, 'title': _rtl(title_lines[1]),
                                 'cards': cards, 'rows': draft_rows,
                                 'challenges': draft_challenges})
    return {'source_pdf_sha256': hashlib.sha256(path.read_bytes()).hexdigest(),
            'source_date_hijri': date.group(), 'metrics': metrics,
            'sectors': sectors, 'challenges': challenges,
            'detail_pages': detail_pages, 'reviewed': False}


def load_reviewed(path, pdf):
    data = json.loads(Path(path).read_text(encoding='utf-8-sig'))
    if data.get('source_pdf_sha256') != hashlib.sha256(Path(pdf).read_bytes()).hexdigest():
        raise ValueError('Reviewed digital-maturity data does not match the supplied PDF')
    if data.get('reviewed') is not True:
        raise ValueError('Review the extracted digital-maturity data against PDF page one before building')
    if set(data.get('metrics', {})) != set(AXES):
        raise ValueError('Four digital-maturity metrics are required')
    for value in data['metrics'].values():
        if not isinstance(value, (int, float)) or isinstance(value, bool) or not 0 <= value <= 100:
            raise ValueError('Invalid digital-maturity metric')
    sectors = data.get('sectors', [])
    if not 1 <= len(sectors) <= 5 or len({s['name'] for s in sectors}) != len(sectors):
        raise ValueError('Digital-maturity sectors must be unique and fit the approved chart')
    for sector in sectors:
        if not sector.get('name') or set(sector.get('metrics', {})) != set(AXES):
            raise ValueError('Incomplete digital-maturity sector')
        numbers = list(sector['metrics'].values())
        if any(v is not None and (not isinstance(v, (int, float)) or isinstance(v, bool) or not 0 <= v <= 100) for v in numbers):
            raise ValueError('Invalid digital-maturity sector percentage')
        if sector.get('note') and any(v is not None for v in numbers):
            raise ValueError('A digital-maturity sector note cannot have percentages')
    if not 1 <= len(data.get('challenges', [])) <= 4:
        raise ValueError('Digital-maturity challenges exceed the approved table')
    if any(not c.get('perspective') or not c.get('text') for c in data['challenges']):
        raise ValueError('Incomplete digital-maturity challenge')
    pages = data.get('detail_pages', [])
    if len(pages) != 6 or [page.get('page') for page in pages] != list(range(2, 8)):
        raise ValueError('Six reviewed digital-maturity detail pages are required in PDF order')
    for page in pages:
        if not page.get('title') or set(page.get('cards', {})) != set(AXES):
            raise ValueError('Digital-maturity detail title or cards are incomplete')
        for value in page['cards'].values():
            if not isinstance(value, str) or not value.strip():
                raise ValueError('Digital-maturity detail card is empty')
            if '%' in value:
                _percent(value)
            elif value not in ('معلق', 'لا ينطبق'):
                raise ValueError('Unapproved digital-maturity detail status: ' + value)
        rows = page.get('rows', [])
        if not rows or any(len(row) != 4 or not isinstance(row[0], str) for row in rows):
            raise ValueError('Digital-maturity detail rows are incomplete')
        for row in rows:
            if not (all(isinstance(cell, str) and cell.strip() for cell in row) or
                    (row[0].strip() and not any(row[1:]))):
                raise ValueError('Digital-maturity detail row must be complete or a section heading')
        if not isinstance(page.get('challenges'), list) or any(not isinstance(c, str) or not c.strip() for c in page['challenges']):
            raise ValueError('Digital-maturity detail challenges are incomplete')
    return data


def _chart(shape, categories, series, percent_labels=False):
    chart_data = CategoryChartData()
    chart_data.categories = categories
    for name, values in series:
        chart_data.add_series(name, values)
    shape.chart.replace_data(chart_data)
    if percent_labels:
        shape.chart.plots[0].data_labels.number_format = '0%'
        for chart_series in shape.chart.series:
            chart_series.data_labels.number_format = '0%'
    blank = shape.chart._chartSpace.find('.//{http://schemas.openxmlformats.org/drawingml/2006/chart}dispBlanksAs')
    if blank is not None:
        blank.set('val', 'gap')


def digital_slide(prs):
    matches = [sl for sl in prs.slides if sum(sh.has_chart for sh in sl.shapes) == 5
               and any(sh.has_table and len(sh.table.columns) == 3 for sh in sl.shapes)]
    if len(matches) != 1:
        raise ValueError('Expected one approved digital-maturity slide')
    return matches[0]


def calculation_note(slide):
    note = helpers.shape_by_id(slide, CALCULATION_NOTE_ID)
    if ' '.join(note.text.split()) != ' '.join(CALCULATION_NOTE.split()):
        raise ValueError('Approved digital-maturity calculation note is missing or changed')
    return note


def detail_slides(prs):
    """Find the six approved detail pages without relying on slide numbers."""
    matches = []
    for slide in prs.slides:
        if any(sh.has_table and len(sh.table.columns) == 4 and len(sh.table.rows) >= 2
               and sh.table.cell(1, 0).text == 'أساس التقييم' for sh in slide.shapes):
            matches.append(slide)
    if len(matches) != 6:
        raise ValueError('Expected six approved digital-maturity detail slides')
    return matches


def _detail_table(slide):
    return next(sh.table for sh in slide.shapes if sh.has_table and len(sh.table.columns) == 4
                and len(sh.table.rows) >= 2 and sh.table.cell(1, 0).text == 'أساس التقييم')


def _challenge_table(slide):
    return next((sh.table for sh in slide.shapes if sh.has_table and len(sh.table.columns) == 2
                 and sh.table.cell(0, 0).text == 'التحديات'), None)


def _protect_codes(value):
    """Keep Latin task IDs readable inside the master's RTL table cells."""
    return re.sub(r'\b[A-Z]{1,5}-\d{2,3}\b', lambda match: '\u202d' + match.group() + '\u202c', value)


def _plain(value):
    return value.replace('\u202d', '').replace('\u202c', '')


def _set_challenge_text(cell, value):
    set_cell_exact(cell, value)
    parts = re.split(r'(\b[A-Z]{1,5}-\d{2,3}\b)', value)
    if len(parts) == 1:
        return
    paragraph = cell.text_frame.paragraphs[0]
    first = paragraph.runs[0]
    template = deepcopy(first._r.get_or_add_rPr())
    first.text = parts[0]
    for index, part in enumerate(parts[1:], 1):
        run = paragraph.add_run()
        run.text = part
        props = deepcopy(template)
        if index % 2:
            props.set('rtl', '0')
            props.set('lang', 'en-US')
        run._r.insert(0, props)


def populate_details(prs, data):
    for slide, detail in zip(detail_slides(prs), data['detail_pages']):
        for axis, shape_id in DETAIL_CARD_IDS.items():
            shape = helpers.shape_by_id(slide, shape_id)
            paragraphs = shape.text_frame.paragraphs
            if len(paragraphs) != 2 or len(paragraphs[1].runs) != 1:
                raise ValueError('Approved digital-maturity card formatting changed')
            paragraphs[1].runs[0].text = detail['cards'][axis]
        table = _detail_table(slide)
        if len(detail['rows']) > len(table.rows)-2:
            raise ValueError(f'Digital-maturity PDF page {detail["page"]} exceeds its detail table capacity')
        set_cell_exact(table.cell(0, 0), detail['title'])
        for index, values in enumerate(detail['rows'], 2):
            for cell, value in zip(table.rows[index].cells, values):
                set_cell_exact(cell, _protect_codes(value))
            if re.fullmatch(r'[A-Z]{1,5}-\d{2,3}', values[3]):
                code_cell = table.cell(index, 3)
                code_cell.text_frame.paragraphs[0]._p.get_or_add_pPr().set('rtl', '0')
        for row in list(table.rows)[2+len(detail['rows']):]:
            row._tr.getparent().remove(row._tr)
        challenges = _challenge_table(slide)
        if challenges is None:
            if detail['challenges']:
                raise ValueError(f'Digital-maturity PDF page {detail["page"]} has challenges without an approved destination')
            continue
        if not 1 <= len(detail['challenges']) <= len(challenges.rows)-2:
            raise ValueError(f'Digital-maturity PDF page {detail["page"]} exceeds its challenges capacity')
        for index, challenge in enumerate(detail['challenges'], 2):
            set_cell_exact(challenges.cell(index, 0), str(index-1))
            _set_challenge_text(challenges.cell(index, 1), challenge)
        for row in list(challenges.rows)[2+len(detail['challenges']):]:
            row._tr.getparent().remove(row._tr)
        # Keep the approved 11 pt type while clearing the inherited footer.
        challenge_shape = next(sh for sh in slide.shapes if sh.has_table and sh.table._tbl is challenges._tbl)
        excess = challenge_shape.top + challenge_shape.height - content_bottom(prs, slide) + 30000
        if excess > 0:
            body = list(challenges.rows)[2:]
            reduction = (excess + len(body) - 1) // len(body)
            if any(row.height - reduction < 350000 for row in body):
                raise ValueError(f'Digital-maturity PDF page {detail["page"]} exceeds its footer clearance')
            for row in body:
                row.height -= reduction
            challenge_shape.height = sum(row.height for row in challenges.rows)


def populate(prs, data):
    slide = digital_slide(prs)
    calculation_note(slide)
    for axis, (chart_id, label_id, track_id, fill_id) in CARDS.items():
        value = data['metrics'][axis]
        _chart(helpers.shape_by_id(slide, chart_id), ['المنجز', 'المتبقي'],
               [('نسبة الإنجاز', [_chart_number(value), _chart_number(100-value)])])
        label_shape = helpers.shape_by_id(slide, label_id)
        helpers.set_text(label_shape, f'{value:g}%')
        label_shape.text_frame.paragraphs[0]._p.get_or_add_pPr().set('rtl', '0')
        for run in label_shape.text_frame.paragraphs[0].runs:
            run._r.get_or_add_rPr().set('lang', 'en-US')
        track, fill = (helpers.shape_by_id(slide, sid) for sid in (track_id, fill_id))
        fill.left = track.left
        fill.width = round(track.width * value / 100)
    sector_chart = helpers.shape_by_id(slide, 93)
    _chart(sector_chart, [s['name'] for s in data['sectors']],
           [(label, [_chart_number(s['metrics'][axis]/100) if s['metrics'][axis] is not None else None
                     for s in data['sectors']])
            for axis, label in zip(AXES, LABELS)], percent_labels=True)
    if any(s.get('note') for s in data['sectors']):
        row_spacing = 1524000 / len(data['sectors'])
        for index, sector in enumerate(data['sectors']):
            if not sector.get('note'):
                continue
            left = sector_chart.left
            top = sector_chart.top + round(155000 + (5-len(data['sectors']))*70000 + index * row_spacing)
            box = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, 4900000, 155000)
            box.name = f'Digital sector note {index+1}'
            box.fill.solid()
            box.fill.fore_color.rgb = RGBColor(229, 227, 224)
            box.line.fill.background()
            box.text = sector['note']
            for paragraph in box.text_frame.paragraphs:
                paragraph.alignment = PP_ALIGN.CENTER
                paragraph.font.size = Pt(11)
                paragraph.font.name = 'Abar Mid'
                paragraph.font.color.rgb = RGBColor(55, 55, 55)
    table = next(sh.table for sh in slide.shapes if sh.has_table and len(sh.table.columns) == 3)
    for index, challenge in enumerate(data['challenges'], 2):
        for cell, value in zip(table.rows[index].cells,
                               (str(index-1), challenge['perspective'], challenge['text'])):
            set_cell_exact(cell, value)
    for row in list(table.rows)[2+len(data['challenges']):]:
        row._tr.getparent().remove(row._tr)
    populate_details(prs, data)
    return slide


def check(prs, report_path, data, errors):
    from report_common import validation
    slide = digital_slide(prs)
    try:
        calculation_note(slide)
    except ValueError as exc:
        errors.append(str(exc))
    for axis, (chart_id, label_id, track_id, fill_id) in CARDS.items():
        value = data['metrics'][axis]
        if helpers.shape_by_id(slide, label_id).text != f'{value:g}%':
            errors.append('Digital-maturity card label differs: ' + axis)
        track, fill = (helpers.shape_by_id(slide, sid) for sid in (track_id, fill_id))
        if abs(fill.width - round(track.width*value/100)) > 2 or fill.left != track.left:
            errors.append('Digital-maturity progress bar differs: ' + axis)
        chart = helpers.shape_by_id(slide, chart_id)
        validation.chart_metrics(report_path, str(chart.chart.part.partname).lstrip('/'),
                                 [['المنجز', _chart_number(value)], ['المتبقي', _chart_number(100-value)]],
                                 ['نسبة الإنجاز'], errors)
    sector_chart = helpers.shape_by_id(slide, 93)
    if any(series.data_labels.number_format != '0%' for series in sector_chart.chart.series):
        errors.append('Digital-maturity sector labels must display percentages')
    rows = [[s['name'], *[_chart_number(s['metrics'][axis]/100) if s['metrics'][axis] is not None else None
                         for axis in AXES]] for s in data['sectors']]
    validation.chart_metrics(report_path, str(sector_chart.chart.part.partname).lstrip('/'), rows, LABELS, errors)
    table = next(sh.table for sh in slide.shapes if sh.has_table and len(sh.table.columns) == 3)
    actual = [[c.text for c in row.cells] for row in list(table.rows)[2:]]
    expected = [[str(i), c['perspective'], c['text']] for i, c in enumerate(data['challenges'], 1)]
    if actual != expected:
        errors.append('Digital-maturity challenges differ from reviewed PDF data')
    notes = {s.name: s.text for s in slide.shapes if s.name.startswith('Digital sector note ')}
    expected_notes = {f'Digital sector note {i}': s['note'] for i, s in enumerate(data['sectors'], 1) if s.get('note')}
    if notes != expected_notes:
        errors.append('Digital-maturity sector notes differ from reviewed PDF data')
    for slide, detail in zip(detail_slides(prs), data['detail_pages']):
        for axis, shape_id in DETAIL_CARD_IDS.items():
            actual = helpers.shape_by_id(slide, shape_id).text_frame.paragraphs[1].text
            if actual != detail['cards'][axis]:
                errors.append(f'Digital-maturity detail card differs on PDF page {detail["page"]}: {axis}')
        table = _detail_table(slide)
        if table.cell(0, 0).text != detail['title']:
            errors.append(f'Digital-maturity detail title differs on PDF page {detail["page"]}')
        actual_rows = [[_plain(cell.text) for cell in row.cells] for row in list(table.rows)[2:]]
        if actual_rows != detail['rows']:
            errors.append(f'Digital-maturity detail rows differ on PDF page {detail["page"]}')
        challenges = _challenge_table(slide)
        actual_challenges = [] if challenges is None else [_plain(row.cells[1].text) for row in list(challenges.rows)[2:]]
        if actual_challenges != detail['challenges']:
            errors.append(f'Digital-maturity detail challenges differ on PDF page {detail["page"]}')


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--pdf', required=True)
    parser.add_argument('--output', required=True)
    args = parser.parse_args(argv)
    output = Path(args.output)
    if output.resolve() == Path(args.pdf).resolve():
        raise ValueError('Extraction output must not overwrite the PDF')
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(json.dumps(extract_pdf(args.pdf), ensure_ascii=False, indent=2), encoding='utf-8')
    print(output)


if __name__ == '__main__':
    main()

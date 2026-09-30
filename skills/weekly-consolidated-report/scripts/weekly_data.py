"""Weekly-only template, transaction schema and approved status mapping."""
from collections import Counter
from copy import deepcopy
from pathlib import Path
from openpyxl import load_workbook
from pptx.table import _Cell
from pptx.dml.color import RGBColor
from pptx import Presentation
from report_common import pptx_helpers as helpers
from report_common import status_styles, validation

MASTER=Path(__file__).resolve().parents[1]/'assets/weekly-report-master.pptx'
MASTER_HASH='3e9b018c223c76cbb0e435cd543732801ec8685fc4ec827d4a7f99b29e0423e5'
TRANS_FIELDS=('رقم المعاملة','موضوع المعاملة','الجهة الوارد منها المعاملة','تاريخ إنشاء المعاملة','تاريخ الإنجاز المخطط')
# User confirmed on 2026-09-29 that Monday status colors apply to weekly too.
# Weekly typography, margins, alignment and borders still come from weekly cells.
STATUS_FILLS={'مكتملة':'DCE6F2','على المخطط':'EBF1DE','متأخر':'F2DCDB','معلق':'F2F2F2','لم تبدأ':'F2F2F2'}

def field_key(name):
    return validation.norm(name).replace('إنشاء','انشاء').replace('الانجاز','الإنجاز').replace('المخطط له','المخطط')

def content_bottom(prs,slide):
    # Locate actual footer artwork in this template's inherited layout.
    # Final rendered inspection remains authoritative for visual clearance.
    layers=(slide,slide.slide_layout,slide.slide_layout.slide_master)
    footer=[shape.top for layer in layers for shape in layer.shapes
            if shape.top>prs.slide_height*.85 and shape.height<prs.slide_height*.1]
    if not footer:raise ValueError('Cannot identify approved footer artwork')
    return min(footer)-90000

def approved_styles(template=MASTER,kind='task'):
    result={}
    presentation=template if hasattr(template,'slides') else Presentation(template)
    for slide in helpers.detail_pattern_slides(presentation,kind):
        for shape in helpers.detail_tables(slide,kind):
            for row in list(shape.table.rows)[2:]:
                cell=row.cells[4];status=status_styles.canonical(cell.text)
                if not status:continue
                if status not in STATUS_FILLS:raise ValueError('Unknown weekly master status: '+status)
                if status in result and status_styles.signature(cell)!=status_styles.signature(result[status]):
                    raise ValueError('Conflicting weekly status examples: '+status)
                result.setdefault(status,cell)
    if 'مكتملة' not in result:raise ValueError('Weekly master missing base status-cell example')
    for status,color in STATUS_FILLS.items():
        if status in result:
            fill=result[status].fill
            if fill.type is None or str(fill.fore_color.rgb)!=color:
                raise ValueError('Weekly status example conflicts with approved mapping: '+status)
        else:
            cell=_Cell(deepcopy(result['مكتملة']._tc),result['مكتملة']._parent)
            cell.fill.solid();cell.fill.fore_color.rgb=RGBColor.from_string(color)
            helpers.set_cell_text(cell,status)
            result[status]=cell
    return result

def summary_columns(title,headers,is_task):
    last='طلب الدعم' if title=='طلبات الدعم' else 'التحديث' if title=='أبرز التحديثات' else 'التحدي' if title=='التحديات' else 'ملاحظات'
    source=['#','المهمة' if is_task else 'المشروع','القطاع',last]
    # The approved suspended-task pattern labels its notes destination التحديث.
    # It still receives the independent suspended-task notes, never project updates.
    keys=[last if title=='المهام المعلقة' and h=='التحديث' else h for h in headers]
    if sorted(keys)!=sorted(source):raise ValueError('Unexpected summary headers: '+title)
    return [source.index(h) for h in keys]

def set_cell_exact(cell,value):
    text=helpers.val(value).replace('\r\n','\n')
    helpers.set_cell_text(cell,text)
    count=len(text.split('\n'))
    while len(cell.text_frame.paragraphs)>count:
        paragraph=cell.text_frame.paragraphs[-1]._p
        paragraph.getparent().remove(paragraph)
    while len(cell.text_frame.paragraphs)<count:cell.text_frame.add_paragraph()

def transaction_data(path):
    chart,kpis=validation.source_metrics(path,'المعاملات',['مكتملة','على المخطط','متأخر'],['B2','B3','B4','B5'])
    wb=load_workbook(path,data_only=True)
    try:
        ws=wb['المعاملات']
        expected=['رقم المعاملة','موضوع المعاملة','القطاع','الجهة الوارد منها المعاملة','تاريخ إنشاء المعاملة','تاريخ الإنجاز المخطط']
        header=None
        for r in range(1,ws.max_row+1):
            got={field_key(ws.cell(r,c).value):c for c in range(1,ws.max_column+1) if ws.cell(r,c).value is not None}
            if all(field_key(s) in got for s in expected):header=r;columns=got;break
        if header is None:raise ValueError('Transactions detail header missing')
        records=[]
        for r in range(header+1,ws.max_row+1):
            values=[ws.cell(r,columns[field_key(s)]).value for s in expected]
            if values[0] in (None,'','لا يوجد'):continue
            records.append([helpers.val(v) for v in values])
        if len(records)!=kpis[3]:raise ValueError('Delayed transaction total differs from detail records')
        actual=Counter(row[2] for row in records)
        counts={str(row[0]):int(row[3] or 0) for row in chart}
        if any(actual[s]!=counts.get(s,0) for s in actual.keys()|counts.keys()):
            raise ValueError('Delayed transactions differ from sector counts')
        return kpis,chart,records
    finally:wb.close()

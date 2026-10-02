# /// script
# requires-python = ">=3.12"
# dependencies = ["openpyxl>=3.1", "python-pptx>=1.0.2", "lxml>=5.3"]
# ///
"""Generate a validated weekly candidate; rendered layout review remains required."""
import argparse, copy, datetime as dt, hashlib, json, math, sys, tempfile, zipfile
from copy import deepcopy
from pathlib import Path
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.oxml.ns import qn
sys.path.insert(0,str(Path(__file__).resolve().parents[2]))
sys.path.insert(0,str(Path(__file__).resolve().parent))
from report_common import pptx_helpers as builder
from report_common import status_styles
from report_common.package_tools import reachable_parts
from weekly_data import transaction_data, field_key, MASTER, MASTER_HASH, approved_styles, set_cell_exact, summary_columns, content_bottom

from update_format import format_update

def remove(shape):shape._element.getparent().remove(shape._element)

def flow(prs,seed,sections,summary=False,measurements=None):
    """Initial placement only; final visual acceptance must confirm pagination."""
    patterns={builder.table_title(sh):sh for sh in builder.tables(seed)}
    top=min(sh.top for sh in patterns.values())
    bottom=content_bottom(prs,seed)
    # Continuations preserve layout/navigation, but not the source KPI/chart.
    def keep(sh):
        outside_summary = sh.top + sh.height <= 700000 or sh.top >= bottom
        return (not sh.has_table and not sh.has_chart
                and not getattr(sh,'has_text_frame',False)
                and (not summary or outside_summary))
    current=seed;y=top;targets=[seed]
    for sh in list(builder.tables(seed)):remove(sh)
    for title,records,pattern in sections:
        if not records:continue
        fixed=sum(r.height for r in list(pattern.table.rows)[:2])
        base=builder._body_base_height(pattern.table,2)
        if measurements is not None and title not in measurements:raise ValueError('Missing layout measurements: '+title)
        measured=(measurements or {}).get(title)
        if measured and len(measured['body'])!=len(records):raise ValueError('Measured record count differs: '+title)
        if measured and (len(measured['header'])!=2 or any(not isinstance(h,(int,float)) or not math.isfinite(h) or h<=0 for h in measured['header']+measured['body'])):
            raise ValueError('Invalid native row heights: '+title)
        fixed_heights=[math.ceil(h*12700) for h in measured['header']] if measured else None
        if fixed_heights:fixed=sum(fixed_heights)
        chunk=[];weights=[];heights=[]
        def place():
            nonlocal y,chunk,weights,heights
            element=deepcopy(pattern._element)
            identity=element.find('.//'+qn('p:cNvPr'))
            identity.set('id',str(max((s.shape_id for s in current.shapes),default=0)+1))
            for node in list(element.iter()):
                if node.tag.rsplit('}',1)[-1] in ('creationId','custDataLst'):node.getparent().remove(node)
            current.shapes._spTree.insert_element_before(element,'p:extLst')
            sh=list(current.shapes)[-1];sh.top=y
            builder.fill_titled_table(sh,chunk,title=title,weights=weights)
            for row,values in zip(list(sh.table.rows)[2:],chunk):
                for cell,value in zip(row.cells,values):set_cell_exact(cell,value)
            if fixed_heights:
                for row,height in zip(sh.table.rows,fixed_heights+heights):row.height=height
                sh.height=sum(row.height for row in sh.table.rows)
            y+=sh.height+90000;chunk=[];weights=[];heights=[]
        for index,row in enumerate(records):
            weight=builder.row_units(row,pattern.table,11)
            height=math.ceil(measured['body'][index]*12700) if measured else base*weight
            if measured and fixed+height>bottom-(900000 if summary else top):raise ValueError('Rendered record exceeds page capacity with the approved style: '+title)
            used=fixed+sum(heights)+height
            if y+used>bottom:
                if chunk:place()
                current=builder.clone_slide(prs,seed,after=current,shape_predicate=keep)
                targets.append(current);y=900000 if summary else top
            chunk.append(row);weights.append(weight);heights.append(height)
        if chunk:place()
    return targets

def update_chart(shape,records,statuses):
    data=CategoryChartData()
    # Retain the template category order for matching sectors, then append new ones.
    order=[str(cat.label) for cat in shape.chart.plots[0].categories]
    rank={name:i for i,name in enumerate(order)}
    records=sorted([row for row in records if any(float(value or 0)>0 for value in row[1:])],key=lambda row:rank.get(str(row[0]),len(rank)))
    rows=records or [[None]*(len(statuses)+1)]
    data.categories=[row[0] or '' for row in rows]
    for c,status in enumerate(statuses,1):data.add_series(status,[row[c] or None for row in rows])
    shape.chart.replace_data(data)
    chart=shape.chart._chartSpace
    blank=chart.find('.//'+qn('c:dispBlanksAs'))
    if blank is not None:blank.set('val','gap')
    return records

def build_candidate(args):
    template=Path(args.template) if args.template else MASTER
    if not args.template and hashlib.sha256(template.read_bytes()).hexdigest()!=MASTER_HASH:
        raise ValueError('Bundled weekly template checksum mismatch')
    tasks=Path(args.tasks);projects=Path(args.suhail);transactions=Path(args.transactions)
    measurements=None
    if args.measurements:
        measured=json.loads(Path(args.measurements).read_text(encoding='utf-8-sig'))
        for key,path in [('tasks',tasks),('suhail',projects),('transactions',transactions),('template',template)]:
            if hashlib.sha256(path.read_bytes()).hexdigest().lower()!=measured['hashes'][key].lower():raise ValueError('Layout measurements are stale: '+key)
        measurements=measured['tables']
    tk,tc,ts,td=builder.task_data(tasks);pk,pc,ps,pd=builder.suhail_data(projects);xk,xc,xd=transaction_data(transactions)
    prs=Presentation(template);slides=list(prs.slides)
    # Discover roles first; IDs are checked within the identified summary role.
    summary=[builder.find_slide_with_text(prs,label) for label in ('إجمالي المهام','إجمالي مشاريع سهيل','إجمالي العاملات')]
    task_patterns=builder.detail_pattern_slides(prs,'task')
    project_patterns=builder.detail_pattern_slides(prs,'project')
    if not task_patterns or not project_patterns:raise ValueError('Missing weekly detail patterns')
    styles={kind:approved_styles(prs,kind) for kind in ('task','project')}
    bullet_style=next((deepcopy(p._p.get_or_add_pPr()) for sl in project_patterns
                       for sh in builder.detail_tables(sl,'project') for row in sh.table.rows
                       for cell in row.cells for p in cell.text_frame.paragraphs
                       if p._p.find('./'+qn('a:pPr')+'/'+qn('a:buChar')) is not None),None)

    date=dt.datetime.now().astimezone().date()
    builder.set_text(builder.shape_by_id(slides[0],19),f'{date.day} {builder.MONTHS[date.month-1]} {date.year}')
    chart_specs=[]
    for sl,kpis,ids,records,statuses,chart_name in [
        (summary[0],tk,[11,14,16,7,6],tc,['مكتملة','على المخطط','متأخر','معلق'],'chart1.xml'),
        (summary[1],pk,[10,12,14,8,7],pc,['مكتملة','على المخطط','متأخر','لم تبدأ'],'chart2.xml'),
        (summary[2],xk,[24,29,31,26],xc,['مكتملة','على المخطط','متأخر'],'chart3.xml')]:
        for sid,value in zip(ids,kpis):builder.set_text(builder.shape_by_id(sl,sid),value)
        ordered=update_chart(next(sh for sh in sl.shapes if sh.has_chart),records,statuses)
        chart_specs.append((chart_name,ordered,statuses))
    for sl,sections in [(summary[0],ts),(summary[1],ps)]:
        # Map source header order to the approved template header order.
        seq=[]
        for title,records in sections.items():
            pattern=next(sh for sh in builder.tables(sl) if builder.table_title(sh)==title)
            headers=[c.text.strip() for c in pattern.table.rows[1].cells]
            columns=summary_columns(title,headers,sl==summary[0])
            seq.append((title,[[row[i] for i in columns] for row in records],pattern))
        flow(prs,sl,seq,True,measurements)
    tframe=next(sh for sh in builder.tables(summary[2]))
    headings=[field_key(c.text) for c in tframe.table.rows[1].cells]
    xheaders=[field_key(s) for s in ['رقم المعاملة','موضوع المعاملة','القطاع','الجهة الوارد منها المعاملة','تاريخ إنشاء المعاملة','تاريخ الإنجاز المخطط']]
    flow(prs,summary[2],[('المعاملات المتأخرة',[[row[xheaders.index(h)] for h in headings] for row in xd],tframe)],True,measurements)
    generated={};unused=[]
    for kind,patterns,sections in [('task',task_patterns,td),('project',project_patterns,pd)]:
        seed=patterns[0];pattern=builder.detail_tables(seed,kind)[0]
        # Remove any additional sample tables before filling a single flowing stream.
        seq=[((f'تفاصيل المهام ({title})' if kind=='task' else title),[row[:6] for row in rows],pattern) for title,rows in sections]
        generated[kind]=flow(prs,seed,seq,measurements=measurements)
        if not any(rows for _,rows in sections):generated[kind]=[]
        for sl in generated[kind]:
            for sh in builder.tables(sl,6):
                for row in list(sh.table.rows)[2:]:
                    cell=row.cells[4];status=status_styles.canonical(cell.text)
                    allowed=status_styles.TASK_STATUSES if kind=='task' else status_styles.PROJECT_STATUSES
                    if status not in allowed:raise ValueError('Unsupported '+kind+' status: '+status)
                    if status not in styles[kind]:raise ValueError('Missing approved weekly status style: '+status)
                    status_styles.apply(cell,styles[kind][status])
                    set_cell_exact(cell,status)
                    if kind=='project':
                        format_update(row.cells[5],bullet_style)
        unused.extend(patterns[1:])
        if not generated[kind]:unused.append(seed)
    for slide in unused:builder.remove_slide(prs,slide)
    # Repoint all internal slide links from removed examples to their surviving section.
    retained={sl.part for sl in prs.slides}
    for part in list(prs.part.package.iter_parts()):
        for rel in list(part.rels.values()):
            if rel.is_external or not rel.reltype.endswith('/slide') or rel.target_part in retained:continue
            fallback=generated['task'] if rel.target_part in {slide.part for slide in task_patterns} else generated['project']
            if fallback:rel._target=fallback[0].part
            else:
                for el in list(part._element.iter()) if hasattr(part,'_element') else []:
                    if el.get(qn('r:id'))==rel.rId and el.tag in (qn('a:hlinkClick'),qn('a:hlinkMouseOver')):el.getparent().remove(el)
                part.drop_rel(rel.rId)
    for i,sl in enumerate(prs.slides,1):
        for sh in sl.shapes:
            if sh._element.find('.//'+qn('a:fld')) is not None and getattr(sh,'has_text_frame',False):
                for t in sh._element.findall('.//'+qn('a:t')):t.text=str(i)
    deck=Path(args.output);prs.save(deck)
    with tempfile.TemporaryDirectory() as td:
        with zipfile.ZipFile(deck) as z:z.extractall(td)
        included=reachable_parts(td)
        with zipfile.ZipFile(deck,'w',zipfile.ZIP_DEFLATED) as z:
            for name in sorted(included):z.write(Path(td)/name,name)
    return deck

def parse_args(argv=None):
    parser=argparse.ArgumentParser(description=__doc__)
    for name in ('tasks','suhail','transactions','output'):parser.add_argument('--'+name,required=True)
    parser.add_argument('--template',help='Explicitly approved alternative template')
    parser.add_argument('--measurements',help='Native row heights from a rendered candidate with these exact inputs')
    return parser.parse_args(argv)

def main(argv=None):
    args=parse_args(argv)
    output=Path(args.output).resolve()
    protected=[Path(args.template) if args.template else MASTER,Path(args.tasks),Path(args.suhail),Path(args.transactions)]
    if args.measurements:protected.append(Path(args.measurements))
    if any(output==p.resolve() or (output.exists() and p.exists() and output.samefile(p)) for p in protected):
        raise ValueError('Output must not overwrite the template or an input workbook')
    from validate_report import validate
    output.parent.mkdir(parents=True,exist_ok=True)
    with tempfile.TemporaryDirectory(prefix='.weekly-build-',dir=output.parent) as temporary:
        candidate=copy.copy(args);candidate.output=str(Path(temporary)/'report.pptx')
        build_candidate(candidate)
        validate(candidate.output,args.tasks,args.suhail,args.transactions,args.template or MASTER)
        status_styles.publish(candidate.output,output)
    print(str(output)+': programmatic checks passed; render and visual review required')
    return 0

if __name__=='__main__':raise SystemExit(main())

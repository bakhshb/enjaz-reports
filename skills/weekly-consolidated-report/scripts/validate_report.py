# /// script
# requires-python = ">=3.12"
# dependencies = ["openpyxl>=3.1", "python-pptx>=1.0.2", "lxml>=5.3"]
# ///
"""Read-only weekly source reconciliation; visual acceptance is a separate step."""
import argparse, collections, re, sys
from pathlib import Path
from pptx import Presentation
from pptx.oxml.ns import qn
sys.path.insert(0,str(Path(__file__).resolve().parents[2]))
sys.path.insert(0,str(Path(__file__).resolve().parent))
from report_common import pptx_helpers as helpers, validation, status_styles
from weekly_data import MASTER, approved_styles, transaction_data, field_key, TRANS_FIELDS, summary_columns, content_bottom

from update_format import update_parts, display_text

SPECS=[('إجمالي المهام','ملخص المهام',['مكتملة','على المخطط','متأخر','معلق'],['B2','B3','B4','B5','B6'],[11,14,16,7,6]),
       ('إجمالي مشاريع سهيل','ملخص مشاريع سهيل',['مكتملة','على المخطط','متأخر','لم تبدأ'],['A3','B3','C3','D3','E3'],[10,12,14,8,7]),
       ('إجمالي العاملات','المعاملات',['مكتملة','على المخطط','متأخر'],['B2','B3','B4','B5'],[24,29,31,26])]

def exact(value):
    return helpers.val(value).replace('\r\n','\n').replace('\v','\n')

def detail_row(row):
    result=[exact(v) for v in row]
    result[4]=status_styles.canonical(result[4])
    return result

def compare(label,expected,actual,errors):
    if [[exact(v) for v in row] for row in expected]!=[[exact(v) for v in row] for row in actual]:
        errors.append(label+': records, text, grouping or order differ from source')

def font_signature(run):
    cs=run._r.find('./'+qn('a:rPr')+'/'+qn('a:cs'))
    return (run.font.name,run.font.size.pt if run.font.size else None,cs.get('typeface') if cs is not None else run.font.name)

def check_table_style(table,pattern,errors):
    if [c.width for c in table.columns]!=[c.width for c in pattern.columns]:errors.append('Table column widths changed')
    for ri,row in enumerate(table.rows):
        reference=pattern.rows[ri if ri<len(pattern.rows) else 2]
        for ci,cell in enumerate(row.cells):
            base=reference.cells[ci]
            is_status=len(table.columns)==6 and ri>=2 and ci==4
            if not is_status and (cell.margin_left,cell.margin_right,cell.margin_top,cell.margin_bottom)!=(base.margin_left,base.margin_right,base.margin_top,base.margin_bottom):errors.append('Table cell margins changed')
            fonts={font_signature(run) for p in base.text_frame.paragraphs for run in p.runs if run.text.strip()}
            # Empty reference cells retain formatting in their first run.
            if not fonts:fonts={font_signature(run) for p in base.text_frame.paragraphs for run in p.runs}
            for p in cell.text_frame.paragraphs:
                for run in p.runs:
                    if run.text.strip() and font_signature(run) not in fonts:errors.append('Table font differs from weekly template: '+run.text[:25])

def check_navigation(prs,errors):
    slides=list(prs.slides)
    positions=[]
    for label,*_ in SPECS:
        positions.append(next((i for i,sl in enumerate(slides) if label in helpers.slide_texts(sl)),-1))
    dividers=[(i,sl) for i,sl in enumerate(slides) if {'تفاصيل المهام','تفاصيل مشاريع سهيل'}.issubset(helpers.slide_texts(sl))]
    if len(dividers)!=1:errors.append('Expected one detail navigation divider');return
    divider_index,divider=dividers[0]
    if positions!=sorted(positions) or positions[0]<=0 or divider_index<=positions[-1]:errors.append('Summary/divider order changed')
    groups={kind:[(i,sl) for i,sl in enumerate(slides) if helpers.detail_tables(sl,kind)] for kind in ('task','project')}
    for kind,label in [('task','تفاصيل المهام'),('project','تفاصيل مشاريع سهيل')]:
        shape=next(sh for sh in divider.shapes if getattr(sh,'has_text_frame',False) and sh.text==label)
        links=list(shape._element.iter(qn('a:hlinkClick')))
        targets=[divider.part.rels[link.get(qn('r:id'))].target_part for link in links if link.get(qn('r:id'))]
        expected=[groups[kind][0][1].part] if groups[kind] else []
        if targets!=expected:errors.append('Navigation target differs: '+kind)
        if any(i<=divider_index or i>=len(slides)-1 for i,_ in groups[kind]):errors.append('Detail section order changed')
    if groups['task'] and groups['project'] and groups['task'][-1][0]>=groups['project'][0][0]:errors.append('Task/project order changed')
    for i,sl in enumerate(slides):
        for sh in helpers.tables(sl,4):
            title=helpers.table_title(sh)
            start,end=(positions[0],positions[1]) if title in helpers.TASK_SUMMARY_TITLES else (positions[1],positions[2])
            if not start<=i<end:errors.append('Summary table outside its section: '+title)
        for sh in helpers.tables(sl,5):
            if not positions[2]<=i<divider_index:errors.append('Transactions outside their section')

def validate(report,tasks,suhail,transactions,template=MASTER):
    errors=[];validation.package(report,errors)
    final=Presentation(report);master=Presentation(template)
    if (final.slide_width,final.slide_height)!=(master.slide_width,master.slide_height):errors.append('Slide dimensions changed')
    tk,tc,ts,td=helpers.task_data(tasks);pk,pc,ps,pd=helpers.suhail_data(suhail);xk,xc,xd=transaction_data(transactions)
    styles={kind:approved_styles(master,kind) for kind in ('task','project')}
    expected_sections={'task':td,'project':pd}
    sources=[tasks,suhail,transactions]
    for i,(label,sheet,statuses,cells,ids) in enumerate(SPECS):
        rows,kpis=validation.source_metrics(sources[i],sheet,statuses,cells)
        if i<2:
            counts=collections.Counter()
            for section,records in expected_sections['task' if i==0 else 'project']:
                for row in records:
                    status=status_styles.canonical(exact(row[4]))
                    if status not in statuses:errors.append('Unsupported detail status: '+status)
                    counts[(exact(row[2]) if i==0 else section,status)]+=1
            expected={(exact(r[0]),s):int(r[j+1] or 0) for r in rows for j,s in enumerate(statuses)}
            if any(counts[k]!=expected.get(k,0) for k in counts.keys()|expected.keys()):errors.append(sheet+': detail and sector counts disagree')
        matches=[sl for sl in final.slides if label in helpers.slide_texts(sl)]
        if len(matches)!=1:errors.append(label+': expected one summary slide');continue
        slide=matches[0]
        for sid,value in zip(ids,kpis):
            if helpers.shape_by_id(slide,sid).text!=str(value):errors.append(label+': KPI differs')
        charts=[sh for sh in slide.shapes if sh.has_chart]
        if len(charts)!=1:errors.append(label+': expected one native chart');continue
        original=helpers.find_slide_with_text(master,label)
        order={str(c.label):j for j,c in enumerate(next(sh for sh in original.shapes if sh.has_chart).chart.plots[0].categories)}
        rows=sorted([r for r in rows if any(v for v in r[1:])],key=lambda r:order.get(str(r[0]),len(order)))
        validation.chart_metrics(report,str(charts[0].chart.part.partname).lstrip('/'),rows,statuses,errors)
    if sum(sh.has_chart for sl in final.slides for sh in sl.shapes)!=3:errors.append('Expected exactly three native charts')
    patterns={helpers.table_title(sh):sh.table for sl in master.slides for sh in helpers.tables(sl,4)}
    task_pattern=helpers.detail_tables(helpers.detail_pattern_slides(master,'task')[0],'task')[0].table
    project_pattern=helpers.detail_tables(helpers.detail_pattern_slides(master,'project')[0],'project')[0].table
    transaction_pattern=next(sh.table for sl in master.slides for sh in helpers.tables(sl,5))
    actual={'task':[],'project':[]};summaries=collections.defaultdict(list);trans=[]
    for sl in final.slides:
        for sh in helpers.tables(sl):
            table=sh.table
            if sh.top+sh.height>content_bottom(final,sl):errors.append('Table frame reaches footer artwork; rendered correction required')
            if len(table.rows)<3:errors.append('Empty table retained');continue
            rows=[[c.text for c in row.cells] for row in table.rows]
            title=helpers.table_title(sh);headers=rows[1]
            if len(headers)==6:
                kind='task' if 'المهمة' in headers else 'project'
                pattern=task_pattern if kind=='task' else project_pattern
                if headers!=[c.text for c in pattern.rows[1].cells]:errors.append('Detail headers changed')
                section=title
                if kind=='task':
                    match=re.fullmatch(r'تفاصيل المهام \((.*)\)',title)
                    if not match:errors.append('Task section title missing')
                    section=match.group(1) if match else title
                for row in rows[2:]:actual[kind].append([section,*detail_row(row)])
                for row in list(table.rows)[2:]:
                    cell=row.cells[4];status=status_styles.canonical(cell.text)
                    allowed=status_styles.TASK_STATUSES if kind=='task' else status_styles.PROJECT_STATUSES
                    if status not in allowed:errors.append('Unsupported '+kind+' status: '+status)
                    if status not in styles[kind] or status_styles.signature(cell)!=status_styles.signature(styles[kind][status]):errors.append('Status style differs: '+status)
                    if kind=='project':
                        # Expected emphasis is limited to the leading label/date span.
                        for pi,p in enumerate(row.cells[5].text_frame.paragraphs):
                            offset=0
                            date_end=update_parts(p.text)[0][1] if pi==0 else 0
                            for run in p.runs:
                                if run.text and (bool(run.font.bold)!=(offset<date_end) or offset<date_end<offset+len(run.text)):errors.append('Suhail update-date emphasis differs')
                                offset+=len(run.text)
            elif 'رقم المعاملة' in headers:
                pattern=transaction_pattern
                keys=[field_key(h) for h in headers]
                if keys!=[field_key(c.text) for c in pattern.rows[1].cells]:errors.append('Transaction column order changed')
                trans.extend([[row[keys.index(field_key(h))] for h in TRANS_FIELDS] for row in rows[2:]])
            elif title in patterns:
                pattern=patterns[title]
                if headers!=[c.text for c in pattern.rows[1].cells]:errors.append('Summary column order changed')
                columns=summary_columns(title,headers,title in ts)
                summaries[title].extend([[row[columns.index(i)] for i in range(4)] for row in rows[2:]])
            else:errors.append('Unexpected table: '+title);continue
            check_table_style(table,pattern,errors)
    for kind,sections in expected_sections.items():
        expected=[]
        for section,records in sections:
            for record in records:
                values=detail_row(record[:6])
                if kind=='project':values[5]=display_text(values[5])
                expected.append([section,*values])
        compare(kind,expected,actual[kind],errors)
        if kind=='project':
            source_updates=[record[5] for _,records in sections for record in records]
            output_cells=[row.cells[5] for sl in final.slides for sh in helpers.detail_tables(sl,'project') for row in list(sh.table.rows)[2:]]
            for source,cell in zip(source_updates,output_cells):
                parts=update_parts(exact(source))
                paragraphs=cell.text_frame.paragraphs
                if len(parts)!=len(paragraphs):errors.append('Suhail update paragraph count differs');continue
                for (_,_,bullet),p in zip(parts,paragraphs):
                    actual_bullet=p._p.find('./'+qn('a:pPr')+'/'+qn('a:buChar'))
                    if (actual_bullet is not None)!=bullet:errors.append('Suhail update list formatting differs')
    for title,records in {**ts,**ps}.items():compare(title,records,summaries[title],errors)
    compare('transactions',[[row[i] for i in (0,1,3,4,5)] for row in xd],trans,errors)
    # Summary records remain independently sourced; match identities without inventing notes.
    for title,status in [('المهام المكتملة','مكتملة'),('المهام المتأخرة','متأخر'),('المهام المعلقة','معلق'),('المشاريع المتأخرة','متأخر')]:
        project=title=='المشاريع المتأخرة';sections=pd if project else td
        identities=collections.Counter((exact(row[1]),section if project else exact(row[2])) for section,rows in sections for row in rows if status_styles.canonical(exact(row[4]))==status)
        summary=ps[title] if project else ts[title]
        if identities!=collections.Counter((exact(row[1]),exact(row[2])) for row in summary):errors.append(title+': summary/detail identities disagree')
    if helpers.slide_texts(final.slides[-1])!=helpers.slide_texts(master.slides[-1]):errors.append('Closing slide changed or not last')
    check_navigation(final,errors)
    if errors:raise ValueError('\n'.join(dict.fromkeys(errors)))
    return {'programmatic':'passed','visual':'pending','slides':len(final.slides),'report_sha256':validation.sha(report)}

def main(argv=None):
    parser=argparse.ArgumentParser(description=__doc__)
    for name in ('report','tasks','suhail','transactions'):parser.add_argument('--'+name,required=True)
    parser.add_argument('--template',default=str(MASTER));args=parser.parse_args(argv)
    print(validate(args.report,args.tasks,args.suhail,args.transactions,args.template))
    return 0

if __name__=='__main__':raise SystemExit(main())

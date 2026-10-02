"""Rich text for weekly Suhail details; explicit list markers identify lists."""
import re
from copy import deepcopy
from pptx.oxml.xmlchemy import OxmlElement
from pptx.oxml.ns import qn

DATE = re.compile(r'^(\s*تاريخ التحديث\s*[:：]?\s*(?:\d{1,4}[/.-]\d{1,2}[/.-]\d{1,4}|\d{1,2}\s+[^\W\d_]+(?:\s+[^\W\d_]+)?\s+\d{4})(?:\s*[هـم])?)(?=\s|$|[،؛:])')
MARKER = re.compile(r'^\s*(?:[•●▪◦*-]|\d+[.)\-]|[٠-٩]+[.)\-])\s+')

def update_parts(text):
    lines = text.replace('\r\n', '\n').replace('\v', '\n').split('\n')
    # Wrapped paragraphs are not lists. Require at least two explicit items.
    marked = [i for i, line in enumerate(lines) if MARKER.match(line)]
    bullets = set(marked) if len(marked) > 1 else set()
    result = []
    for i, line in enumerate(lines):
        bullet = i in bullets
        value = MARKER.sub('', line, count=1) if bullet else line
        date = DATE.match(value) if i == 0 else None
        result.append((value, date.end() if date else 0, bullet))
    return result

def display_text(text):
    return '\n'.join(value for value, _, _ in update_parts(text))

def format_update(cell, bullet_style=None):
    parts = update_parts(cell.text)
    frame = cell.text_frame
    base = next((deepcopy(r._r.get_or_add_rPr()) for p in frame.paragraphs for r in p.runs), None)
    paragraph_style = deepcopy(frame.paragraphs[0]._p.get_or_add_pPr())
    frame.clear()
    for i, (value, bold_end, bullet) in enumerate(parts):
        p = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
        old = p._p.find(qn('a:pPr'))
        if old is not None: p._p.remove(old)
        props = deepcopy(paragraph_style)
        for node in list(props):
            if node.tag in {qn('a:buNone'), qn('a:buChar'), qn('a:buAutoNum'), qn('a:buBlip')}:
                props.remove(node)
        if bullet:
            if bullet_style is None:
                raise ValueError('Approved weekly bullet style is required for Suhail list items')
            for attr in ('marL', 'marR', 'indent'):
                if bullet_style.get(attr) is not None: props.set(attr, bullet_style.get(attr))
        marker = OxmlElement('a:buChar' if bullet else 'a:buNone')
        if bullet: marker.set('char', '•')
        # Bullets precede tabLst/defRPr/extLst in DrawingML schema order.
        at = next((n for n, el in enumerate(props) if el.tag in {qn('a:tabLst'), qn('a:defRPr'), qn('a:extLst')}), len(props))
        props.insert(at, marker)
        p._p.insert(0, props)
        for text, bold in ((value[:bold_end], True), (value[bold_end:], False)):
            if not text: continue
            run = p.add_run(); run.text = text
            if base is not None: run._r.insert(0, deepcopy(base))
            run.font.bold = bold

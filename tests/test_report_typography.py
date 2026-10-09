"""11 pt and per-column alignment for both report templates."""
import sys, unittest
from pathlib import Path
from pptx import Presentation
from pptx.enum.text import PP_ALIGN
from pptx.util import Pt
ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT/'skills'))
from report_common.table_typography import set_table_sizes, all_tables, check_table_typography

class TableTypographyTests(unittest.TestCase):
    def templates(self):
        for skill,name in [('monday-meeting-report','monday-meeting-master.pptx'),('weekly-consolidated-report','weekly-report-master.pptx')]:
            yield Presentation(ROOT/'skills'/skill/'assets'/name)

    def test_size_alignment_and_other_formatting_preserved(self):
        expected={'القطاع':PP_ALIGN.CENTER,'تاريخ الإنجاز المخطط':PP_ALIGN.CENTER,'تاريخ الانجاز المخطط':PP_ALIGN.CENTER,'تاريخ البداية':PP_ALIGN.CENTER,'تاريخ النهاية':PP_ALIGN.CENTER,'الحالة':PP_ALIGN.CENTER,'المهمة':PP_ALIGN.RIGHT,'اسم المشروع':PP_ALIGN.RIGHT,'المشروع':PP_ALIGN.RIGHT,'ملاحظات':PP_ALIGN.RIGHT,'التحديث':PP_ALIGN.RIGHT,'طلب الدعم':PP_ALIGN.RIGHT,'التحدي':PP_ALIGN.RIGHT,'ما تم حتى تاريخه':PP_ALIGN.RIGHT}
        for prs in self.templates():
            before=[(c.text,c._tc.get_or_add_tcPr().xml,[(r.text,r.font.name,r.font.bold) for p in c.text_frame.paragraphs for r in p.runs]) for t in all_tables(prs) for row in t.rows for c in row.cells]
            set_table_sizes(prs)
            after=[(c.text,c._tc.get_or_add_tcPr().xml,[(r.text,r.font.name,r.font.bold) for p in c.text_frame.paragraphs for r in p.runs]) for t in all_tables(prs) for row in t.rows for c in row.cells]
            self.assertEqual(before,after)
            for t in all_tables(prs):
                for i,row in enumerate(t.rows):
                    for j,c in enumerate(row.cells):
                        for p in c.text_frame.paragraphs:
                            self.assertEqual(p.font.size.pt,11)
                            for r in p.runs:self.assertEqual(r.font.size.pt,11)
                            if i and t.cell(1,j).text in expected:self.assertEqual(p.alignment,expected[t.cell(1,j).text])
            errors=[];check_table_typography(prs,errors);self.assertEqual(errors,[])

    def test_gate_rejects_size_and_alignment_corruption(self):
        for prs in self.templates():
            set_table_sizes(prs)
            table=next(t for t in all_tables(prs) if len(t.columns)==6)
            table.cell(2,2).text_frame.paragraphs[0].alignment=PP_ALIGN.RIGHT
            run=next(r for p in table.cell(2,1).text_frame.paragraphs for r in p.runs if r.text)
            run.font.size=Pt(10)
            errors=[];check_table_typography(prs,errors)
            self.assertIn('Table font size must be 11 pt',errors)
            self.assertTrue(any('alignment' in e for e in errors))

if __name__=='__main__':unittest.main()

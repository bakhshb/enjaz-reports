# /// script
# requires-python = ">=3.12"
# dependencies = ["pandas>=2.2,<3", "openpyxl>=3.1,<4", "python-pptx>=1.0.2", "lxml>=5.3", "pillow"]
# ///
"""Weekly-source handoff and independent sector identity checks."""
import contextlib
from copy import deepcopy
import io
import json
from pathlib import Path
import shutil
import tempfile
import unittest
import zipfile
from pptx import Presentation
from openpyxl import load_workbook

import run_weekly_workflow as workflow

qa=workflow.fixtures.sector_qa
run=workflow.run
SECTOR=workflow.fixtures.SECTOR


class WeeklyChainTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.temp=tempfile.TemporaryDirectory();cls.root=Path(cls.temp.name)
        workflow.scenario(cls.root/'baseline','normal')

    @classmethod
    def tearDownClass(cls):cls.temp.cleanup()

    def setUp(self):
        self.case=tempfile.TemporaryDirectory();self.path=Path(self.case.name)
        self.master=self.root/'baseline/weekly.pptx';self.transactions=self.root/'baseline/transactions.xlsx'
        self.out=self.path/'out';shutil.copytree(self.root/'baseline/sectors',self.out)
        self.deck=next(self.out.glob('*.pptx'))
        manifest=json.loads((self.out/'manifest.json').read_text(encoding='utf-8'))
        for entry in manifest:entry['out']=str(self.deck)
        (self.out/'manifest.json').write_text(json.dumps(manifest,ensure_ascii=False),encoding='utf-8')

    def tearDown(self):self.case.cleanup()

    def check(self):
        with contextlib.redirect_stdout(io.StringIO()):
            return qa.validate(self.out) and qa.check(self.out,self.master,self.transactions)

    def test_normal_chain_and_transaction_identifier(self):
        self.assertTrue(self.check())
        data=qa.actual_records(Presentation(self.deck))
        self.assertEqual(data['المعاملات المتأخرة'][0][0],'000123')

    def test_transaction_table_starts_at_same_height_as_other_sections(self):
        tops={}
        for slide in Presentation(self.deck).slides:
            tables=[sh for sh in slide.shapes if sh.has_table]
            if tables:
                titles={workflow.splitter.table_title(workflow.splitter.table_rows(sh)) for sh in tables}
                section='transactions' if 'المعاملات المتأخرة' in titles else 'tasks' if titles & set(workflow.splitter.TASK_TITLE.values()) else 'projects'
                tops.setdefault(section,[]).append(min(sh.top for sh in tables))
        self.assertEqual(set(tops),{'transactions','tasks','projects'})
        self.assertTrue(all(top==tops['tasks'][0] for values in tops.values() for top in values),tops)

    def test_transaction_position_corruption_rejected(self):
        prs=Presentation(self.deck)
        table=next(sh for sl in prs.slides for sh in sl.shapes if sh.has_table
                   and workflow.splitter.table_title(workflow.splitter.table_rows(sh))=='المعاملات المتأخرة')
        table.top+=100000;prs.save(self.deck)
        self.assertFalse(self.check())

    def test_transaction_title_is_bold_and_nonbold_is_rejected(self):
        prs=Presentation(self.deck)
        table=next(sh.table for sl in prs.slides for sh in sl.shapes if sh.has_table
                   and workflow.splitter.table_title(workflow.splitter.table_rows(sh))=='المعاملات المتأخرة')
        runs=[run for cell in table.rows[0].cells for p in cell.text_frame.paragraphs
              for run in p.runs if run.text.strip()]
        self.assertTrue(runs)
        self.assertTrue(all(run.font.bold for run in runs))
        for run in runs:run.font.bold=False
        prs.save(self.deck);self.assertFalse(self.check())

    def test_transaction_continuations_start_at_section_top(self):
        prs=Presentation(self.master)
        source=prs.slides[workflow.splitter.locate(prs)['trans']]
        table=workflow.splitter.find_tables(source,['رقم المعاملة'])[0][0].table
        row=list(table.rows)[-1]._tr
        for _ in range(60):table._tbl.append(deepcopy(row))
        master=self.path/'large-source.pptx';prs.save(master);out=self.path/'large'
        wb=load_workbook(self.transactions);ws=wb['المعاملات'];rows=list(ws.values)
        header=next(i for i,r in enumerate(rows) if 'رقم المعاملة' in r)
        record=next(r for r in rows[header+1:] if r[rows[header].index('رقم المعاملة')])
        for _ in range(60):ws.append(record)
        transactions=self.path/'large-transactions.xlsx';wb.save(transactions)
        workflow.require(run(SECTOR/'split_report.py',master,'--transactions',transactions,'--outdir',out,'--no-refine'))
        with contextlib.redirect_stdout(io.StringIO()):
            self.assertTrue(qa.check(out,master,transactions))
        tables=[sh for sl in Presentation(next(out.glob('*.pptx'))).slides for sh in sl.shapes
                if sh.has_table and workflow.splitter.table_title(workflow.splitter.table_rows(sh))=='المعاملات المتأخرة']
        self.assertGreater(len(tables),1)
        self.assertTrue(all(sh.top==workflow.splitter.TOP_Y for sh in tables))
        for shape in tables:
            runs=[run for cell in shape.table.rows[0].cells for p in cell.text_frame.paragraphs
                  for run in p.runs if run.text.strip()]
            self.assertTrue(runs)
            self.assertTrue(all(run.font.bold for run in runs))

    def test_cover_corruption_rejected(self):
        prs=Presentation(self.deck)
        group=prs.slides[0].shapes[0]
        next(sh for sh in group.shapes if sh.name=='SectorName').text='قطاع خاطئ'
        prs.save(self.deck);self.assertFalse(self.check())

    def test_task_kpi_corruption_rejected(self):
        prs=Presentation(self.deck)
        next(sh for sh in prs.slides[1].shapes if sh.shape_id==11).text='999'
        prs.save(self.deck);self.assertFalse(self.check())

    def test_transaction_kpi_corruption_rejected(self):
        prs=Presentation(self.deck)
        slide=next(sl for sl in prs.slides if any(getattr(sh,'has_text_frame',False) and sh.text=='إجمالي المعاملات المتأخرة' for sh in sl.shapes))
        next(sh for sh in slide.shapes if sh.shape_id==24).text='999'
        prs.save(self.deck);self.assertFalse(self.check())

    def test_missing_transaction_mapping_rejected(self):
        result=run(SECTOR/'split_report.py',self.master,'--outdir',self.path/'missing','--no-refine')
        self.assertNotEqual(result.returncode,0);self.assertIn('--transactions',result.stderr)

    def test_alias_and_only_selection(self):
        aliases=self.path/'aliases.json';aliases.write_text(json.dumps({'وكالة شؤون الحج':'قطاع موحد'},ensure_ascii=False),encoding='utf-8')
        out=self.path/'alias'
        args=['--transactions',self.transactions,'--aliases',aliases,'--only','قطاع موحد']
        workflow.require(run(SECTOR/'split_report.py',self.master,'--outdir',out,'--no-refine',*args))
        workflow.require(run(SECTOR/'qa_report.py',out,'--master',self.master,*args))
        self.assertEqual(len(list(out.glob('*.pptx'))),1)

    def test_selection_without_records_is_valid(self):
        out=self.path/'empty';args=['--transactions',self.transactions,'--only','لا يطابق أي قطاع']
        workflow.require(run(SECTOR/'split_report.py',self.master,'--outdir',out,'--no-refine',*args))
        workflow.require(run(SECTOR/'qa_report.py',out,'--master',self.master,*args))
        self.assertEqual(list(out.glob('*.pptx')),[])

    def test_unlisted_slide_rejected(self):
        with zipfile.ZipFile(self.deck,'a') as package:
            package.writestr('ppt/slides/slide999.xml',package.read('ppt/slides/slide1.xml'))
        self.assertFalse(self.check())

    def test_extra_deck_rejected(self):
        shutil.copy2(self.deck,self.out/'stale.pptx');self.assertFalse(self.check())


if __name__=='__main__':unittest.main()

"""Observable acceptance of manual tables, planned dates and Suhail rich text."""
import contextlib, io, shutil, tempfile, unittest
from pathlib import Path
from openpyxl import load_workbook
from pptx import Presentation
from pptx.oxml.ns import qn
import run_weekly_workflow as workflow

class RequestedReportChangesTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.temp=tempfile.TemporaryDirectory();cls.root=Path(cls.temp.name)
        with contextlib.redirect_stdout(io.StringIO()):workflow.scenario(cls.root,'normal')

    @classmethod
    def tearDownClass(cls):cls.temp.cleanup()

    def test_sector_delayed_and_on_plan_dates_without_status_in_detail_order(self):
        master=Presentation(self.root/'weekly.pptx')
        for deck in (self.root/'sectors').glob('*.pptx'):
            prs=Presentation(deck)
            seen={title:[] for title in ('المهام على المخطط','المهام المتأخرة')}
            for sl in prs.slides:
                for sh in workflow.weekly.builder.tables(sl):
                    title=workflow.weekly.builder.table_title(sh)
                    if title in seen:
                        self.assertEqual([c.text for c in sh.table.rows[1].cells],['#','المهمة','تاريخ الإنجاز المخطط','ملاحظات'])
                        seen[title].extend([[c.text for c in r.cells][1:] for r in list(sh.table.rows)[2:]])
                    elif title in workflow.splitter.TASK_TITLE.values():
                        self.assertNotIn('تاريخ الإنجاز المخطط',[c.text for c in sh.table.rows[1].cells])
            for status in ('على المخطط','متأخرة'):
                expected=[r for sl in master.slides for sh in workflow.weekly.builder.detail_tables(sl,'task') for r in list(sh.table.rows)[2:] if workflow.splitter.norm_status(r.cells[4].text)==status]
                self.assertTrue(expected)
                self.assertEqual(seen[workflow.splitter.TASK_TITLE[status]],[[r.cells[i].text for i in (1,3,5)] for r in expected])

    def test_sector_tables_omit_sector_and_on_plan_title_is_bold(self):
        found=False
        for deck in (self.root/'sectors').glob('*.pptx'):
            for sl in Presentation(deck).slides:
                for sh in workflow.weekly.builder.tables(sl):
                    self.assertNotIn('القطاع',[c.text for c in sh.table.rows[1].cells])
                    if workflow.weekly.builder.table_title(sh)=='المهام المعلقة':
                        self.assertIn('ملاحظات',[c.text for c in sh.table.rows[1].cells])
                        self.assertNotIn('التحديث',[c.text for c in sh.table.rows[1].cells])
                    if workflow.weekly.builder.table_title(sh)=='المهام على المخطط':
                        found=True
                        runs=[run for cell in sh.table.rows[0].cells
                              for p in cell.text_frame.paragraphs for run in p.runs if run.text.strip()]
                        self.assertTrue(runs)
                        self.assertTrue(all(run.font.bold for run in runs))
        self.assertTrue(found)

    def test_sector_date_corruption_rejected(self):
        for title in ('المهام على المخطط','المهام المتأخرة'):
            with self.subTest(title=title),tempfile.TemporaryDirectory() as td:
                out=Path(td)/'sectors';shutil.copytree(self.root/'sectors',out)
                deck=next(out.glob('*.pptx'));prs=Presentation(deck)
                sh=next(sh for sl in prs.slides for sh in workflow.weekly.builder.tables(sl) if workflow.weekly.builder.table_title(sh)==title)
                sh.table.cell(2,2).text='تاريخ خاطئ';prs.save(deck)
                import json
                manifest=json.loads((out/'manifest.json').read_text(encoding='utf-8'))
                for entry in manifest:entry['out']=str(deck)
                (out/'manifest.json').write_text(json.dumps(manifest),encoding='utf-8')
                with contextlib.redirect_stdout(io.StringIO()):
                    self.assertFalse(workflow.fixtures.sector_qa.check(out,self.root/'weekly.pptx',self.root/'transactions.xlsx'))

    def test_sector_task_widths_match_on_plan_with_full_width_tables(self):
        for deck in (self.root/'sectors').glob('*.pptx'):
            tables={workflow.weekly.builder.table_title(sh):sh
                    for sl in Presentation(deck).slides for sh in workflow.weekly.builder.tables(sl)}
            reference=tables['المهام على المخطط']
            widths={c.text:reference.table.columns[i].width for i,c in enumerate(reference.table.rows[1].cells)}
            for title in (*workflow.splitter.TASK_TITLE.values(),'طلبات الدعم'):
                sh=tables[title]
                actual={c.text:sh.table.columns[i].width for i,c in enumerate(sh.table.rows[1].cells)}
                with self.subTest(title=title):
                    self.assertEqual(sh.width,reference.width)
                    self.assertEqual(actual['#'],widths['#'])
                    self.assertEqual(actual['المهمة'],widths['المهمة'])
                    if 'ملاحظات' in actual:
                        extra=0 if 'تاريخ الإنجاز المخطط' in actual else widths['تاريخ الإنجاز المخطط']
                        self.assertAlmostEqual(actual['ملاحظات'],widths['ملاحظات']+extra,delta=2)

    def test_sector_qa_rejects_nonbold_title_and_restored_headers(self):
        deck=next((self.root/'sectors').glob('*.pptx'))
        for damage in ('title','sector_header','status_header','suspended_notes','task_width'):
            with self.subTest(damage=damage):
                prs=Presentation(deck)
                title='المهام المعلقة' if damage=='suspended_notes' else 'المهام على المخطط'
                table=next(sh.table for sl in prs.slides for sh in workflow.weekly.builder.tables(sl)
                           if workflow.weekly.builder.table_title(sh)==title)
                if damage=='title':
                    for cell in table.rows[0].cells:
                        for p in cell.text_frame.paragraphs:
                            for run in p.runs:run.font.bold=False
                    message='title must be bold'
                elif damage=='sector_header':
                    table.cell(1,1).text='القطاع'
                    message='Sector column must be omitted'
                elif damage=='status_header':
                    table.cell(1,3).text='الحالة'
                    message='without status'
                elif damage=='suspended_notes':
                    table.cell(1,2).text='التحديث'
                    message='notes header'
                else:
                    table.columns[1].width+=100000
                    table.columns[3].width-=100000
                    message='Task column widths must match'
                with self.assertRaisesRegex(ValueError,message):
                    workflow.fixtures.sector_qa.actual_records(prs)

    def test_sector_planned_dates_preserve_distinct_and_blank_values(self):
        with tempfile.TemporaryDirectory() as td:
            root=Path(td);master=root/'source.pptx';prs=Presentation(self.root/'weekly.pptx')
            for sl in prs.slides:
                for sh in workflow.weekly.builder.detail_tables(sl,'task'):
                    for row in list(sh.table.rows)[2:]:
                        status=workflow.splitter.norm_status(row.cells[4].text)
                        if status in ('متأخرة','على المخطط'):
                            row.cells[3].text='15/03/1448' if status=='متأخرة' else ''
            prs.save(master);before=master.read_bytes();out=root/'sectors'
            workflow.require(workflow.fixtures.run(workflow.fixtures.SECTOR/'split_report.py',master,'--transactions',self.root/'transactions.xlsx','--outdir',out,'--no-refine'))
            expected=workflow.fixtures.sector_qa.expected_records(master,self.root/'transactions.xlsx')
            for deck in out.glob('*.pptx'):
                actual=workflow.fixtures.sector_qa.actual_records(Presentation(deck))
                self.assertEqual(actual,expected['وكالة شؤون الحج'])
                self.assertEqual(actual['المهام المتأخرة'][0][1],'15/03/1448')
                self.assertEqual(actual['المهام على المخطط'][0][1],'')
            self.assertEqual(master.read_bytes(),before)

    def rich_case(self,folder):
        folder.mkdir(exist_ok=True)
        for file in self.root.glob('*.xlsx'):shutil.copy2(file,folder/file.name)
        wb=load_workbook(folder/'projects.xlsx');ws=wb['تفاصيل مشاريع سهيل']
        values=[
            'تاريخ التحديث 2 أكتوبر 2026 نص بعد التاريخ\n• تجهيز المتطلبات\n• التنسيق مع الجهات',
            'تاريخ التحديث: 02/10/2026\nفقرة عادية تمتد إلى سطر آخر\nمتابعة التنفيذ والتاريخ 03/10/2026',
            'تاريخ التحديث ٢ أكتوبر ٢٠٢٦\n- نقطة واحدة فقط',
            'تحديث بلا تاريخ\n1. متابعة التنفيذ\n2. معالجة الملاحظات',
        ]
        rows=[r for r in range(3,ws.max_row+1) if ws.cell(r,1).value]
        for row,value in zip(rows,values):ws.cell(row,6,value)
        wb.save(folder/'projects.xlsx')
        with contextlib.redirect_stdout(io.StringIO()):workflow.populate(folder)
        return values

    def test_update_date_and_actual_bullets_preserve_prose(self):
        with tempfile.TemporaryDirectory() as td:
            folder=Path(td);values=self.rich_case(folder)
            prs=Presentation(folder/'weekly.pptx')
            cells=[r.cells[5] for sl in prs.slides for sh in workflow.weekly.builder.detail_tables(sl,'project') for r in list(sh.table.rows)[2:]]
            self.assertEqual(len(cells),4)
            self.assertEqual(cells[0].text,'تاريخ التحديث 2 أكتوبر 2026 نص بعد التاريخ\nتجهيز المتطلبات\nالتنسيق مع الجهات')
            expected_bold=['تاريخ التحديث 2 أكتوبر 2026','تاريخ التحديث: 02/10/2026','تاريخ التحديث ٢ أكتوبر ٢٠٢٦','']
            for cell,bold,bullet_count in zip(cells,expected_bold,[2,2,0,2]):
                bold_text=''.join(r.text for p in cell.text_frame.paragraphs for r in p.runs if r.font.bold)
                self.assertEqual(bold_text,bold)
                bullets=[p for p in cell.text_frame.paragraphs if p._p.find('./'+qn('a:pPr')+'/'+qn('a:buChar')) is not None]
                self.assertEqual(len(bullets),bullet_count)
                self.assertTrue(all(p._p.get_or_add_pPr().get('rtl')=='1' for p in cell.text_frame.paragraphs))
            self.assertEqual(cells[1].text,values[1]);self.assertEqual(cells[2].text,values[2].replace('- نقطة','نقطة'))

    def test_rich_text_corruption_is_rejected(self):
        with tempfile.TemporaryDirectory() as td:
            folder=Path(td);self.rich_case(folder)
            deck=folder/'weekly.pptx';original=deck.read_bytes()
            for change in ('whole_paragraph_bold','missing_bullet','body_bullet'):
                with self.subTest(change=change):
                    deck.write_bytes(original);prs=Presentation(deck)
                    cell=next(r.cells[5] for sl in prs.slides for sh in workflow.weekly.builder.detail_tables(sl,'project') for r in list(sh.table.rows)[2:])
                    if change=='whole_paragraph_bold':
                        for r in cell.text_frame.paragraphs[0].runs:r.font.bold=True
                    elif change=='missing_bullet':
                        p=cell.text_frame.paragraphs[1];node=p._p.find('./'+qn('a:pPr')+'/'+qn('a:buChar'));node.getparent().remove(node)
                    else:
                        from pptx.oxml.xmlchemy import OxmlElement
                        cell.text_frame.paragraphs[0]._p.get_or_add_pPr().append(OxmlElement('a:buChar'))
                    prs.save(deck)
                    with self.assertRaises(ValueError):workflow.weekly_validation.validate(deck,folder/'tasks.xlsx',folder/'projects.xlsx',folder/'transactions.xlsx')

if __name__=='__main__':unittest.main()

# /// script
# requires-python = ">=3.12"
# dependencies = ["pypdf>=5", "python-pptx>=1.0.2", "pdfplumber>=0.11"]
# ///
"""Read-only PDF font, page-count and visible RTL-numbering checks."""
import argparse
import hashlib
import json
import re
from pathlib import Path
from pypdf import PdfReader
from pptx import Presentation
import pdfplumber


def check_numbering(prs, pdf, errors):
    """Check PDF coordinates, independently of the PPTX's declared RTL flag."""
    def tables(shapes):
        for shape in shapes:
            if shape.has_table:yield shape
            if hasattr(shape,'shapes'):yield from tables(shape.shapes)
    checked=0
    for number,(slide,page) in enumerate(zip(prs.slides,pdf.pages),1):
        words=page.extract_words()
        scale_x,scale_y=page.width/prs.slide_width,page.height/prs.slide_height
        for shape in tables(slide.shapes):
            table=shape.table
            headers=[(i,' '.join(c.text.split())) for i,row in enumerate(list(table.rows)[:2])
                     for c in row.cells if ' '.join(c.text.split()) in {'#','م','رقم المعاملة'}]
            if not headers:
                errors.append(f'Page {number}: cannot identify table numbering header');continue
            row,label=headers[0];marker=label.split()[0]
            left,right=shape.left*scale_x,(shape.left+shape.width)*scale_x
            top=(shape.top+sum(r.height for r in list(table.rows)[:row]))*scale_y
            tolerance=max(25,table.rows[row].height*scale_y+12)
            matches=[w for w in words if w['text'] in {marker,marker[::-1]}
                     and left<=w['x0']<=right and abs(w['top']-top)<=tolerance]
            if not matches:
                errors.append(f'Page {number}: table numbering header missing from PDF');continue
            word=min(matches,key=lambda w:abs(w['top']-top))
            if (word['x0']+word['x1'])/2<=(left+right)/2:
                errors.append(f'Page {number}: table numbering is visibly on the left')
            checked+=1
    return checked


def check(root):
    results = []
    for deck in sorted(root.rglob('*.pptx')):
        pdf = deck.with_suffix('.pdf')
        errors, fonts, arabic_runs = [], set(), 0
        if not pdf.exists():
            results.append({'file': str(deck), 'errors': ['Missing PDF']})
            continue
        reader = PdfReader(pdf)
        prs = Presentation(deck)
        if len(reader.pages) != len(prs.slides):
            errors.append('PDF page count differs from PowerPoint')
        for number, page in enumerate(reader.pages, 1):
            def visit(text, cm, tm, font, size):
                nonlocal arabic_runs
                if not text.strip():
                    return
                name = str(font.get('/BaseFont', '') if font else '')
                fonts.add(name)
                if re.search(r'[\u0600-\u06ff]', text):
                    arabic_runs += 1
                    if 'Abar' not in name:
                        errors.append(f'Page {number}: Arabic text uses {name}: {text[:40]!r}')
            page.extract_text(visitor_text=visit)
        if not arabic_runs:
            errors.append('No readable Arabic text; font acceptance cannot be established')
        with pdfplumber.open(pdf) as rendered:
            rtl_checked=check_numbering(prs,rendered,errors)
        results.append({
            'file': str(deck), 'pdf': str(pdf), 'pages': len(reader.pages),
            'arabic_runs': arabic_runs, 'fonts': sorted(fonts), 'errors': errors,
            'rtl_tables_checked': rtl_checked,
            'pptx_sha256': hashlib.sha256(deck.read_bytes()).hexdigest(),
            'pdf_sha256': hashlib.sha256(pdf.read_bytes()).hexdigest(),
        })
    if not results:
        raise ValueError('No presentations found')
    return results


if __name__ == '__main__':
    parser = argparse.ArgumentParser()
    parser.add_argument('root', type=Path)
    args = parser.parse_args()
    results = check(args.root)
    (args.root / 'pdf-font-acceptance.json').write_text(
        json.dumps(results, ensure_ascii=False, indent=2), encoding='utf-8')
    for row in results:
        print(('FAIL' if row['errors'] else 'PASS') + ' ' + row['file'])
        print('\n'.join(row['errors']))
    raise SystemExit(any(row['errors'] for row in results))

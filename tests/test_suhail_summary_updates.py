"""Exercise date emphasis and list formatting through both complete report builds."""
import contextlib, io, tempfile, unittest
from pathlib import Path
from openpyxl import load_workbook
from pptx import Presentation
from pptx.oxml.ns import qn
import run_weekly_workflow as workflow
from report_common.update_format import update_column, SUMMARY_UPDATE_TITLES

class SuhailSummaryUpdateTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.temp=tempfile.TemporaryDirectory();cls.root=Path(cls.temp.name)
        with contextlib.redirect_stdout(io.StringIO()):workflow.scenario(cls.root,'normal')
        wb=load_workbook(cls.root/'projects.xlsx')
        text='تاريخ التحديث 10 أكتوبر 2026\nاستكمال مراجعة المتطلبات\nالتنسيق مع الجهات'
        ws=wb['ملخص مشاريع سهيل']
        for row in ws:
            if any(str(row[0].value or '').endswith(title) for title in SUMMARY_UPDATE_TITLES):
                ws.cell(row[0].row+2,4,text if str(row[0].value).endswith('أبرز التحديثات') else
                        'تاريخ التحديث 10 أكتوبر 2026\nتستمر متابعة التنفيذ مع الجهات المعنية.')
                if str(row[0].value or '').endswith('أبرز التحديثات'):
                    for col,value in enumerate([1,'مشروع تجريبي 1','وكالة شؤون الحج'],1):
                        ws.cell(row[0].row+2,col,value)
        ws=wb['تفاصيل مشاريع سهيل']
        records=[r for r in range(1,ws.max_row+1) if str(ws.cell(r,1).value or '').isdigit()]
        for r,value in zip(records,[text,'تاريخ التحديث 10 أكتوبر 2026\nتستمر متابعة التنفيذ مع الجهات المعنية.',
                                    'تحديث عادي دون تاريخ.',text]):ws.cell(r,6,value)
        wb.save(cls.root/'projects.xlsx')
        with contextlib.redirect_stdout(io.StringIO()):workflow.populate(cls.root)
        workflow.require(workflow.run(workflow.fixtures.MONDAY/'build_report.py',
            '--topics',cls.root/'topics.xlsx','--tasks',cls.root/'tasks.xlsx',
            '--suhail',cls.root/'projects.xlsx','--output',cls.root/'monday.pptx'))

    @classmethod
    def tearDownClass(cls):cls.temp.cleanup()

    def check_report(self,name):
        deck=self.root/(name+'.pptx');original=deck.read_bytes();prs=Presentation(deck)
        cells=[];titles=[]
        for sl in prs.slides:
            for sh in sl.shapes:
                if not sh.has_table:continue
                j=update_column(sh.table)
                if j is None:continue
                titles.append(workflow.weekly.builder.table_title(sh))
                cells.extend(r.cells[j] for r in list(sh.table.rows)[2:])
        self.assertTrue(SUMMARY_UPDATE_TITLES.issubset(titles))
        self.assertGreater(len(cells),2)
        for cell in cells:
            bold=''.join(r.text for p in cell.text_frame.paragraphs for r in p.runs if r.font.bold)
            self.assertEqual(bold,'تاريخ التحديث 10 أكتوبر 2026' if cell.text.startswith('تاريخ التحديث') else '')
            count=sum(p._p.find('./'+qn('a:pPr')+'/'+qn('a:buChar')) is not None for p in cell.text_frame.paragraphs)
            self.assertEqual(count,2 if 'استكمال مراجعة' in cell.text else 0)
        for damage in ['bold','bullets']:
            prs=Presentation(deck)
            sh=next(sh for sl in prs.slides for sh in sl.shapes if sh.has_table and workflow.weekly.builder.table_title(sh)=='أبرز التحديثات')
            cell=sh.table.cell(2,update_column(sh.table))
            if damage=='bold':cell.text_frame.paragraphs[0].runs[0].font.bold=False
            else:
                props=cell.text_frame.paragraphs[1]._p.get_or_add_pPr();props.remove(props.find(qn('a:buChar')))
            prs.save(deck)
            try:
                if name=='weekly':
                    with self.assertRaises(ValueError):workflow.verify(self.root,deck)
                else:
                    result=workflow.run(workflow.fixtures.MONDAY/'report_gate_v2.py','--report',deck,
                        '--topics',self.root/'topics.xlsx','--tasks',self.root/'tasks.xlsx','--suhail',self.root/'projects.xlsx')
                    self.assertNotEqual(result.returncode,0,result.stdout)
            finally:deck.write_bytes(original)

    def test_weekly_summary_and_details(self):self.check_report('weekly')
    def test_monday_summary_and_details(self):self.check_report('monday')

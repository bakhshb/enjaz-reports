"""Digital weekly model keeps reviewed PDF values and rejects unsafe inputs."""
import hashlib
import json
import sys
import tempfile
import unittest
from pathlib import Path

from pptx import Presentation

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / 'skills/weekly-consolidated-report/scripts'))
import digital_maturity as digital


class DigitalMaturityTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.pdf = self.root / 'source.pdf'
        self.pdf.write_bytes(b'synthetic source identity')
        self.data = {
            'source_pdf_sha256': hashlib.sha256(self.pdf.read_bytes()).hexdigest(),
            'source_date_hijri': '1448/1/1',
            'reviewed': True,
            'metrics': {'experience': 15, 'satisfaction': 10, 'complaints': 20, 'tools': 65},
            'sectors': [
                {'name': 'القطاع الأول', 'metrics': {'experience': 15, 'satisfaction': 10, 'complaints': 20, 'tools': 65}, 'note': None},
                {'name': 'القطاع الثاني', 'metrics': {'experience': None, 'satisfaction': None, 'complaints': None, 'tools': None}, 'note': 'قيد الدراسة'},
                {'name': 'القطاع الثالث', 'metrics': {'experience': 25, 'satisfaction': None, 'complaints': None, 'tools': 70}, 'note': None},
                {'name': 'القطاع الرابع', 'metrics': {'experience': 5, 'satisfaction': 15, 'complaints': 30, 'tools': None}, 'note': None},
                {'name': 'القطاع الخامس', 'metrics': {'experience': 10, 'satisfaction': 20, 'complaints': 25, 'tools': 75}, 'note': None},
            ],
            'challenges': [{'perspective': 'عام', 'text': 'تحد أول'}, {'perspective': 'تجربة المستخدم', 'text': 'تحد ثان'}],
        }
        self.review = self.root / 'review.json'

    def load(self):
        self.review.write_text(json.dumps(self.data, ensure_ascii=False), encoding='utf-8')
        return digital.load_reviewed(self.review, self.pdf)

    def test_review_and_pdf_identity_are_required(self):
        self.data['reviewed'] = False
        with self.assertRaisesRegex(ValueError, 'Review'):
            self.load()
        self.data['reviewed'] = True
        self.load()
        self.pdf.write_bytes(b'a different PDF')
        with self.assertRaisesRegex(ValueError, 'does not match'):
            digital.load_reviewed(self.review, self.pdf)

    def test_capacity_and_conflicting_note_rejected(self):
        self.data['challenges'] *= 3
        with self.assertRaisesRegex(ValueError, 'exceed'):
            self.load()
        self.data['challenges'] = self.data['challenges'][:2]
        self.data['sectors'][1]['metrics']['experience'] = 5
        with self.assertRaisesRegex(ValueError, 'cannot have percentages'):
            self.load()

    def test_charts_and_challenges_match_reviewed_data(self):
        data = self.load()
        prs = Presentation(digital.MASTER)
        digital.populate(prs, data)
        report = self.root / 'report.pptx'
        prs.save(report)
        errors = []
        digital.check(Presentation(report), report, data, errors)
        self.assertEqual(errors, [])
        prs = Presentation(report)
        card = digital.helpers.shape_by_id(digital.digital_slide(prs), digital.CARDS['experience'][1])
        card.text = '99%'
        prs.save(report)
        errors = []
        digital.check(Presentation(report), report, data, errors)
        self.assertTrue(any('card label differs' in error for error in errors))
        card.text = '15%'
        sector = digital.helpers.shape_by_id(digital.digital_slide(prs), 93)
        sector.chart.series[0].data_labels.number_format = 'General'
        prs.save(report)
        errors = []
        digital.check(Presentation(report), report, data, errors)
        self.assertTrue(any('must display percentages' in error for error in errors))


if __name__ == '__main__':
    unittest.main()

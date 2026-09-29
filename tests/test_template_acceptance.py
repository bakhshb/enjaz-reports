"""Regressions for the approved template and zero-sector presentation rules."""
import tempfile
import unittest
from pathlib import Path
from pptx import Presentation
from pptx.oxml.ns import qn
import run_weekly_workflow as workflow

class TemplateAcceptanceTests(unittest.TestCase):
    def test_summary_continuations_remove_kpi_dividers(self):
        prs=Presentation(workflow.MASTER)
        seed=prs.slides[3]
        pattern=next(s for s in seed.shapes if s.has_table)
        original_top=pattern.top
        rows=[[str(i),'موضوع تجريبي','01/01/1448','جهة واردة','01/02/1448'] for i in range(40)]
        slides=workflow.flow(prs,seed,[(workflow.builder.table_title(pattern),rows,pattern)],summary=True)
        self.assertGreater(len(slides),1)
        total=0
        for index,slide in enumerate(slides):
            tables=[s for s in slide.shapes if s.has_table]
            total+=sum(len(s.table.rows)-2 for s in tables)
            if index:
                self.assertFalse(any(s.has_chart for s in slide.shapes))
                self.assertFalse(any(s.shape_type==9 for s in slide.shapes))
                self.assertLess(tables[0].top,original_top)
        self.assertEqual(total,len(rows))

    def test_charts_filter_empty_sectors_in_cache_and_workbook(self):
        statuses=['مكتملة','على المخطط','متأخر','معلق']
        for rows,expected in [
            ([['empty',None,0,None,0],['active',0,0,3,0]],[['active',0,0,3,0]]),
            ([['empty',0,0,0,0]],[]),
            ([['new',0,2,0,0],['completed',1,0,0,0]],[['new',0,2,0,0],['completed',1,0,0,0]])]:
            with self.subTest(rows=rows), tempfile.TemporaryDirectory() as td:
                prs=Presentation(workflow.MASTER)
                shape=next(s for s in prs.slides[1].shapes if s.has_chart)
                workflow.update_chart(shape,rows,statuses)
                path=Path(td)/'chart.pptx';prs.save(path)
                errors=[]
                workflow.gate.chart_metrics(path,'ppt/charts/chart1.xml',[[r[0],*[v or None for v in r[1:]]] for r in expected],statuses,errors)
                self.assertEqual(errors,[])
                monday=Path(td)/'monday.pptx'
                monday.write_bytes((workflow.fixtures.MONDAY.parent/'assets/monday-meeting-master.pptx').read_bytes())
                workflow.builder.patch_charts(monday,rows,rows)
                errors=[]
                workflow.gate.chart_metrics(monday,'ppt/charts/chart1.xml',[[r[0],*[v or None for v in r[1:]]] for r in expected],statuses,errors)
                self.assertEqual(errors,[])

    def test_current_transaction_and_notes_patterns_preserve_abar(self):
        prs=Presentation(workflow.MASTER)
        # Fill every existing body row and one cloned overflow row in both roles.
        for index in (3,5):
            for shape in prs.slides[index].shapes:
                if not shape.has_table:continue
                count=len(shape.table.rows)+1
                workflow.builder.resize_table(shape.table,count)
                for row in list(shape.table.rows)[2:]:
                    for cell in row.cells:
                        workflow.builder.set_cell_text(cell,'ملاحظات تجريبية 123')
                        run=cell.text_frame.paragraphs[0].runs[0]
                        for script in ('latin','cs'):
                            font=run._r.find(qn('a:rPr')).find(qn('a:'+script))
                            self.assertIsNotNone(font)
                            self.assertEqual(font.get('typeface'),'Abar Mid')

if __name__=='__main__':unittest.main()

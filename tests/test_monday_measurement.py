"""Native measurement failures must not discard or undersize report rows."""
import contextlib
import io
import json
from pathlib import Path
import subprocess
import sys
from types import SimpleNamespace
import unittest
from unittest.mock import patch
import warnings

from pptx import Presentation
from pptx.util import Pt
from test_deck_integrity import MONDAY

sys.path.insert(0, str(MONDAY))
import monday_layout as layout


class MondayMeasurementTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.template = MONDAY.parent / 'assets/monday-meeting-master.pptx'
        prs = Presentation(cls.template)
        shape = next(sh for sl in prs.slides for sh in sl.shapes
                     if sh.has_table and len(sh.table.columns) == 4)
        row = ['1', 'موضوع قصير', 'القطاع', 'تحديث قصير']
        cls.specs = [('agenda', shape, [row, row], 'جدول الأعمال', None),
                     ('followup', shape, [row], 'المتابعة', None)]
        cls.bullets = layout.bullet_example(prs)
        cls.estimates = cls.measure('estimate')

    @classmethod
    def measure(cls, engine):
        return layout.measure(cls.specs, cls.template, {}, cls.bullets, engine=engine)

    @staticmethod
    def native_result(payload):
        def run(args, **kwargs):
            if isinstance(payload, Exception):
                raise payload
            if payload == 'exit':
                return subprocess.CompletedProcess(args, 1, '', 'PowerPoint unavailable')
            if payload != 'missing':
                Path(args[-1]).with_suffix('.json').write_text(
                    payload if isinstance(payload, str) else json.dumps(payload), encoding='utf-8')
            return subprocess.CompletedProcess(args, 0, '', '')
        return run

    @staticmethod
    def failures():
        valid = {'agenda': [24, 24, 30, 31], 'followup': [24, 24, 32]}
        yield 'exit', 'exit'
        yield 'timeout', subprocess.TimeoutExpired('powershell.exe', 60)
        yield 'missing executable', FileNotFoundError('powershell.exe')
        yield 'missing JSON', 'missing'
        yield 'malformed JSON', '{'
        yield 'partial keys', {'agenda': valid['agenda']}
        yield 'short rows', dict(valid, agenda=[24, 24, 30])
        yield 'extra rows', dict(valid, agenda=[24, 24, 30, 31, 32])
        yield 'wrong container', []
        for value in (0, -1, '30', True, None, float('nan'), float('inf')):
            yield 'invalid height ' + repr(value), dict(valid, agenda=[24, 24, value, 31])

    def test_auto_windows_failure_warns_and_returns_complete_conservative_estimates(self):
        for reason, payload in self.failures():
            with self.subTest(reason=reason), \
                    patch.object(layout, 'os', SimpleNamespace(name='nt')), \
                    patch.object(layout.subprocess, 'run', side_effect=self.native_result(payload)), \
                    warnings.catch_warnings(record=True) as caught, \
                    contextlib.redirect_stderr(io.StringIO()) as stderr:
                warnings.simplefilter('always')
                result = self.measure('auto')
                self.assertEqual(set(result), set(self.estimates))
                for key, (headers, heights) in result.items():
                    expected_headers, expected_heights = self.estimates[key]
                    self.assertEqual(len(headers), len(expected_headers))
                    self.assertEqual(len(heights), len(expected_heights))
                    self.assertTrue(all(actual >= estimate > 0 for actual, estimate
                                        in zip(headers + heights, expected_headers + expected_heights)))
                warning = stderr.getvalue() + ' '.join(str(item.message) for item in caught)
                self.assertRegex(warning, '(?i)measurement|PowerPoint|قياس')

    def test_explicit_native_failure_has_meaningful_error(self):
        for reason, payload in self.failures():
            with self.subTest(reason=reason), \
                    patch.object(layout.subprocess, 'run', side_effect=self.native_result(payload)):
                with self.assertRaisesRegex(RuntimeError, '(?i)measurement|PowerPoint|قياس'):
                    self.measure('native')

    def test_complete_native_result_preserves_every_measured_row(self):
        actual = {'agenda': [24, 24, 30, 31], 'followup': [24, 24, 32]}
        with patch.object(layout.subprocess, 'run', side_effect=self.native_result(actual)):
            result = self.measure('native')
        self.assertEqual(set(result), set(actual))
        for key, values in actual.items():
            headers, heights = result[key]
            self.assertEqual(headers, [max(int(Pt(v)), baseline)
                                      for v, baseline in zip(values[:2], self.estimates[key][0])])
            self.assertEqual(heights, [int(Pt(v + 2)) for v in values[2:]])

    def test_probe_never_points_template_tags_at_a_slide_layout(self):
        # A raw copied table can retain rId1, which means "tags" in its source
        # but "slideLayout" in a fresh probe. PowerPoint refuses that package.
        def run(args, **kwargs):
            probe=Presentation(args[-1])
            for slide in probe.slides:
                for tag in slide._element.xpath('.//p:tags'):
                    rid=tag.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id')
                    self.assertTrue(slide.part.rels[rid].reltype.endswith('/tags'))
            return self.native_result({'agenda':[24,24,30,31],'followup':[24,24,32]})(args,**kwargs)
        with patch.object(layout.subprocess,'run',side_effect=run):
            self.measure('native')


if __name__ == '__main__':
    unittest.main()

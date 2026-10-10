"""Sector output must enforce typography even when the input uses old styling."""
from pathlib import Path
import tempfile,unittest
from pptx import Presentation
from pptx.util import Pt
from pptx.enum.text import PP_ALIGN
from test_deck_integrity import weekly_source_from_template,SECTOR,run,sector_qa
from report_common.table_typography import all_tables,check_table_typography

class SectorTypographyTests(unittest.TestCase):
    def test_old_and_current_inputs_and_corruption(self):
        with tempfile.TemporaryDirectory() as directory:
            root=Path(directory)
            for size in [10,11]:
                with self.subTest(source_size=size):
                    source=root/f'source-{size}.pptx';weekly_source_from_template(source)
                    prs=Presentation(source)
                    for table in all_tables(prs):
                        for row in table.rows:
                            for cell in row.cells:
                                for p in cell.text_frame.paragraphs:
                                    p.alignment=PP_ALIGN.RIGHT
                                    for r in p.runs:r.font.size=Pt(size)
                    prs.save(source);before=source.read_bytes();out=root/f'out-{size}'
                    result=run(SECTOR/'split_report.py',source,'--outdir',out,'--no-refine')
                    self.assertEqual(result.returncode,0,result.stdout+result.stderr)
                    result=run(SECTOR/'qa_report.py',out,'--master',source)
                    self.assertEqual(result.returncode,0,result.stdout+result.stderr)
                    self.assertEqual(before,source.read_bytes())
                    deck=next(out.glob('*.pptx'));prs=Presentation(deck)
                    errors=[];check_table_typography(prs,errors);self.assertFalse(errors)
                    original=deck.read_bytes()
                    for damage in ['size','alignment']:
                        deck.write_bytes(original);prs=Presentation(deck)
                        t=next(t for t in all_tables(prs) if any(c.text=='الحالة' for c in t.rows[1].cells))
                        j=next(i for i,c in enumerate(t.rows[1].cells) if c.text=='الحالة');cell=t.cell(2,j)
                        if damage=='size':next(r for p in cell.text_frame.paragraphs for r in p.runs if r.text).font.size=Pt(10)
                        else:cell.text_frame.paragraphs[0].alignment=PP_ALIGN.RIGHT
                        prs.save(deck)
                        result=run(SECTOR/'qa_report.py',out,'--master',source)
                        self.assertNotEqual(result.returncode,0)

if __name__=='__main__':unittest.main()

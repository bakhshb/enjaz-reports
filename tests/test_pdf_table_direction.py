"""A PDF must show numbered RTL columns on the right, not just declare RTL."""
import hashlib
from pathlib import Path
import shutil
import subprocess
import sys
import tempfile
import unittest

import pdfplumber
from pptx import Presentation
from pptx.util import Inches, Pt
import check_rendered_pdf as pdf_checks

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / 'skills'))


@unittest.skipUnless(shutil.which('soffice'), 'LibreOffice is required for PDF rendering')
class PDFTableDirectionTests(unittest.TestCase):
    def test_rtl_is_visible_in_every_rendered_table_and_source_is_unchanged(self):
        from report_common.render import export_pdf
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            source, output = root / 'report.pptx', root / 'report.pdf'
            prs = Presentation()
            for page in range(2):
                slide = prs.slides.add_slide(prs.slide_layouts[6])
                for index in range(2):
                    table = slide.shapes.add_table(3, 3, Inches(.3), Inches(.3 + index * 3),
                                                   Inches(9), Inches(2.4)).table
                    table._tbl.tblPr.set('rtl', '1')
                    table.cell(0, 0).merge(table.cell(0, 2))
                    table.cell(0, 0).text = f'TITLE-{page}-{index}'
                    for column, value in enumerate(('#', 'المهمة', 'ملاحظات')):
                        table.cell(1, column).text = value
                    for column, value in enumerate(('ID', 'TASK', 'NOTE')):
                        table.cell(2, column).text = f'{value}-{page}-{index}'
                    for row in table.rows:
                        for cell in row.cells:
                            for paragraph in cell.text_frame.paragraphs:
                                for run in paragraph.runs:
                                    run.font.name = 'Abar Mid'
                                    run.font.size = Pt(16)
            prs.save(source)
            before = hashlib.sha256(source.read_bytes()).hexdigest()
            export_pdf(source, output)
            self.assertEqual(hashlib.sha256(source.read_bytes()).hexdigest(), before)
            with pdfplumber.open(output) as pdf:
                self.assertEqual(len(pdf.pages), 2)
                errors = []
                self.assertEqual(pdf_checks.check_numbering(Presentation(source), pdf, errors), 4)
                self.assertEqual(errors, [])
                for number, page in enumerate(pdf.pages):
                    words = {word['text']: word for word in page.extract_words()}
                    for index in range(2):
                        self.assertIn(f'TITLE-{number}-{index}', words)
                        self.assertGreater(words[f'ID-{number}-{index}']['x0'],
                                           words[f'TASK-{number}-{index}']['x0'])
                        self.assertGreater(words[f'TASK-{number}-{index}']['x0'],
                                           words[f'NOTE-{number}-{index}']['x0'])
            # The unadapted exporter reproduces the user's screenshot despite rtl=1.
            raw = root / 'unadapted'
            raw.mkdir()
            subprocess.run([shutil.which('soffice'),
                            '-env:UserInstallation=' + (root / 'raw-profile').as_uri(),
                            '--headless', '--convert-to', 'pdf', '--outdir', str(raw), str(source)],
                           check=True, capture_output=True, timeout=90)
            with pdfplumber.open(raw / 'report.pdf') as pdf:
                errors = []
                self.assertEqual(pdf_checks.check_numbering(Presentation(source), pdf, errors), 4)
                self.assertTrue(any('visibly on the left' in error for error in errors))


if __name__ == '__main__':
    unittest.main()

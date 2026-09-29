# /// script
# requires-python = ">=3.12"
# dependencies = ["pandas>=2.2,<3", "openpyxl>=3.1,<4", "python-pptx>=1.0.2", "lxml>=5.3", "pillow"]
# ///
"""Exercise report gates with synthetic inputs and deliberately corrupted decks."""
import contextlib
from copy import deepcopy
import importlib.util
import io
import json
import shutil
import sys
import tempfile
import unittest
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

from openpyxl import Workbook, load_workbook
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches

from test_upstream_integrity import ROOT, SUHAIL, TASKS, run, suhail_source, task_source

MONDAY=ROOT/'skills/monday-meeting-report/scripts'
SECTOR=ROOT/'skills/weekly-report-by-sector/scripts'
sys.path.insert(0,str(SECTOR))
import qa_report as sector_qa


def load_module(name,path):
    spec=importlib.util.spec_from_file_location(name,path)
    module=importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


gate=load_module('monday_gate',MONDAY/'report_gate_v2.py')


def mutate_zip(path,part,change):
    with zipfile.ZipFile(path) as archive:
        entries={n:archive.read(n) for n in archive.namelist()}
    entries[part]=change(entries[part])
    with zipfile.ZipFile(path,'w',zipfile.ZIP_DEFLATED) as archive:
        for name,data in entries.items():archive.writestr(name,data)


def table(slide,title,headers,records,top=1):
    shape=slide.shapes.add_table(len(records)+2,len(headers),Inches(.3),Inches(top),Inches(8),Inches(.4*(len(records)+2)))
    t=shape.table
    t.cell(0,0).merge(t.cell(0,len(headers)-1));t.cell(0,0).text=title
    for i,h in enumerate(headers):t.cell(1,i).text=h
    for r,values in enumerate(records,2):
        for c,value in enumerate(values):t.cell(r,c).text=str(value)
    return shape


def label(slide,text):
    slide.shapes.add_textbox(Inches(.2),0,Inches(8),Inches(.4)).text=text


def synthetic_sector_files(root):
    sector='قطاع تجريبي'
    src=Presentation()
    cover=src.slides.add_slide(src.slide_layouts[6]);label(cover,'تقرير الإنجاز الأسبوعي')
    for title in ['إجمالي المهام','إجمالي مشاريع سهيل']:
        sl=src.slides.add_slide(src.slide_layouts[6]);label(sl,title)
        data=CategoryChartData();data.categories=[sector];data.add_series('مكتملة',[2])
        sl.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED,0,Inches(1),Inches(4),Inches(2),data)
    task_rows=[['1','مهمة أولى',sector,'01/02/1448','مكتملة','ملاحظة أولى'],['2','مهمة ثانية',sector,'02/02/1448','مكتملة','ملاحظة ثانية']]
    sl=src.slides.add_slide(src.slide_layouts[6])
    table(sl,'تفاصيل المهام (مصدر تجريبي)',gate.TASK_HEADERS,task_rows)
    sl=src.slides.add_slide(src.slide_layouts[6])
    table(sl,sector,gate.SUHAIL_HEADERS,[['1','مشروع أول','01/01/1448','01/02/1448','مكتملة','تحديث']])
    label(src.slides.add_slide(src.slide_layouts[6]),'شكرا')
    master=root/'master.pptx';src.save(master)
    out=root/'sectors';out.mkdir()
    prs=Presentation();sl=prs.slides.add_slide(prs.slide_layouts[6])
    table(sl,'المهام المكتملة',['#','المهمة','القطاع','ملاحظات'],[[r[0],r[1],sector,r[5]] for r in task_rows])
    label(sl,sector)
    sl=prs.slides.add_slide(prs.slide_layouts[6])
    table(sl,'المشاريع المكتملة',gate.SUHAIL_HEADERS,[['1','مشروع أول','01/01/1448','01/02/1448','مكتملة','تحديث']])
    deck=out/'sector.pptx';prs.save(deck)
    (out/'manifest.json').write_text(json.dumps([{'sector':sector,'out':str(deck),'n':[2,1,0,0,0,0]}],ensure_ascii=False),encoding='utf-8')
    return master,out,deck


def build_scenario(root,name,task_count,project_count,new_names=False):
    root.mkdir(parents=True,exist_ok=True)
    raw_tasks=root/'raw-tasks.xlsx';raw_projects=root/'raw-projects.xlsx'
    task_source(raw_tasks);wb=load_workbook(raw_tasks);ws=wb['المهام'];ws.delete_rows(2,ws.max_row)
    wb['طلبات الدعم'].delete_rows(2,wb['طلبات الدعم'].max_row)
    sector='قطاع جديد تجريبي' if new_names else 'وكالة شؤون الحج'
    source='مصدر جديد تجريبي' if new_names else 'اجتماع القيادات'
    for i in range(task_count):
        ws.append([f'مهمة تجريبية {i+1}',sector,source,'01/02/1448',['مكتمل','متأخر','معلق','على المخطط'][i%4],f'ملاحظة تجريبية {i+1}'])
    wb.save(raw_tasks)
    suhail_source(raw_projects);wb=load_workbook(raw_projects);ws=wb.active;ws.delete_rows(2,ws.max_row)
    for i in range(project_count):
        ws.append([sector,f'مشروع تجريبي {i+1}','01/01/1448','01/02/1448','مكتمل','قديم',['مكتمل','متأخر','لم يبدأ','على المخطط'][i%4],f'تاريخ التحديث 29 سبتمبر 2026\nتحديث تجريبي {i+1}','ملاحظة',''])
    wb.save(raw_projects)
    tasks=root/'tasks.xlsx';projects=root/'projects.xlsx'
    for script,args in [(TASKS,['--input',raw_tasks,'--output',tasks]),(SUHAIL,[raw_projects,projects])]:
        result=run(script,*args)
        if result.returncode:raise AssertionError(result.stderr+result.stdout)
    topics=root/'topics.xlsx';wb=Workbook();ws=wb.active;ws.title='Sheet1'
    ws.append([name]);ws.append(['الترتيب','الموضوع','المتحدث','الوقت المطلوب من المالك'])
    for i in range(18 if name=='overflow' else 1):ws.append([i+1,f'موضوع تجريبي {i+1}','مالك',5])
    ws.append([ws.max_row-1,'ملخص الاجتماع','مالك',5]);wb.save(topics)
    deck=root/'monday.pptx'
    result=run(MONDAY/'build_report.py','--topics',topics,'--tasks',tasks,'--suhail',projects,'--output',deck)
    if result.returncode:raise AssertionError(result.stderr+result.stdout)
    result=run(MONDAY/'report_gate_v2.py','--report',deck,'--topics',topics,'--tasks',tasks,'--suhail',projects)
    return result,deck


def weekly_source_from_template(path):
    """Replace the bundled master's sample table data with synthetic records."""
    prs=Presentation(ROOT/'skills/weekly-consolidated-report/assets/weekly-report-master.pptx')
    added_task=added_project=False
    for slide in prs.slides:
        for shape in list(slide.shapes):
            if not shape.has_table or len(shape.table.rows)<2:continue
            t=shape.table;headers=[c.text.strip() for c in t.rows[1].cells]
            body=deepcopy(t.rows[2]._tr if len(t.rows)>2 else t.rows[1]._tr)
            for row in list(t.rows)[2:]:t._tbl.remove(row._tr)
            records=[]
            if len(headers)==6 and 'المهمة' in headers and not added_task:
                added_task=True
                records=[['1','مهمة أولى','قطاع تجريبي','01/02/1448','مكتملة','ملاحظة'],['2','مهمة ثانية','قطاع تجريبي','02/02/1448','متأخر','ملاحظة أخرى']]
            elif len(headers)==6 and 'اسم المشروع' in headers and not added_project:
                added_project=True;t.cell(0,0).text='قطاع تجريبي'
                records=[['1','مشروع تجريبي','01/01/1448','01/02/1448','مكتملة','تحديث']]
            if not records:
                shape._element.getparent().remove(shape._element)
                continue
            for record in records:
                t._tbl.append(deepcopy(body))
                for c,value in enumerate(record):t.cell(len(t.rows)-1,c).text=value
    prs.save(path)


class SplitIntegrationTests(unittest.TestCase):
    def test_template_split_preserves_records_and_removes_unused_slide_parts(self):
        with tempfile.TemporaryDirectory() as td:
            root=Path(td);master=root/'weekly.pptx';weekly_source_from_template(master);out=root/'out'
            result=run(SECTOR/'split_report.py',master,'--outdir',out,'--no-refine')
            self.assertEqual(result.returncode,0,result.stderr+result.stdout)
            result=run(SECTOR/'qa_report.py',out,'--master',master)
            self.assertEqual(result.returncode,0,result.stderr+result.stdout)


class ScenarioTests(unittest.TestCase):
    def test_empty_sources(self):
        with tempfile.TemporaryDirectory() as td:
            result,_=build_scenario(Path(td),'empty',0,0)
            self.assertEqual(result.returncode,0,result.stderr+result.stdout)

    def test_new_source_and_sector(self):
        with tempfile.TemporaryDirectory() as td:
            result,deck=build_scenario(Path(td),'new',4,4,True)
            self.assertEqual(result.returncode,0,result.stderr+result.stdout)
            self.assertTrue(any('قطاع جديد تجريبي' in sh.table.cell(0,0).text for sl in Presentation(deck).slides for sh in sl.shapes if sh.has_table))

    def test_agenda_and_detail_overflow(self):
        with tempfile.TemporaryDirectory() as td:
            result,deck=build_scenario(Path(td),'overflow',32,24)
            self.assertEqual(result.returncode,0,result.stderr+result.stdout)
            self.assertGreater(len(Presentation(deck).slides),len(Presentation(MONDAY.parent/'assets/monday-meeting-master.pptx').slides))


class MondayGateTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.temp=tempfile.TemporaryDirectory();cls.root=Path(cls.temp.name)
        raw_tasks=cls.root/'raw-tasks.xlsx';raw_projects=cls.root/'raw-projects.xlsx'
        task_source(raw_tasks)
        wb=load_workbook(raw_tasks);wb['المهام'].append(['مهمة ثانية','وكالة شؤون الحج','اجتماع القيادات','02/02/1448','متأخر','ملاحظة ثانية']);wb.save(raw_tasks)
        suhail_source(raw_projects,'متأخر')
        cls.tasks=cls.root/'tasks.xlsx';cls.projects=cls.root/'projects.xlsx'
        for script,arguments in [(TASKS,['--input',raw_tasks,'--output',cls.tasks]),(SUHAIL,[raw_projects,cls.projects])]:
            result=run(script,*arguments)
            if result.returncode:raise AssertionError(result.stderr+result.stdout)
        wb=Workbook();ws=wb.active;ws.title='Sheet1';ws.append(['أسبوع تجريبي']);ws.append(['الترتيب','الموضوع','المتحدث','الوقت المطلوب من المالك']);ws.append([1,'مراجعة','مالك',10]);ws.append([2,'ملخص الاجتماع','مالك',5])
        cls.topics=cls.root/'topics.xlsx';wb.save(cls.topics)
        cls.baseline=cls.root/'baseline.pptx'
        result=run(MONDAY/'build_report.py','--topics',cls.topics,'--tasks',cls.tasks,'--suhail',cls.projects,'--output',cls.baseline)
        if result.returncode:raise AssertionError(result.stderr+result.stdout)

    @classmethod
    def tearDownClass(cls):cls.temp.cleanup()

    def candidate(self):
        path=self.root/'candidate.pptx';shutil.copy2(self.baseline,path);return path

    def check(self,path):
        return run(MONDAY/'report_gate_v2.py','--report',path,'--topics',self.topics,'--tasks',self.tasks,'--suhail',self.projects)

    def test_generated_deck_passes(self):
        result=self.check(self.candidate())
        self.assertEqual(result.returncode,0,result.stderr+result.stdout)

    def test_chart_cache_corruption_fails(self):
        path=self.candidate()
        def change(data):
            root=ET.fromstring(data);root.find('.//c:ser/c:val//c:v',gate.CHART_NS).text='999';return ET.tostring(root)
        mutate_zip(path,'ppt/charts/chart1.xml',change)
        result=self.check(path)
        self.assertNotEqual(result.returncode,0)
        self.assertIn('cache differs from source',result.stdout)

    def test_embedded_workbook_corruption_fails(self):
        path=self.candidate()
        def change(data):
            wb=load_workbook(io.BytesIO(data));wb.active['B2']=999;out=io.BytesIO();wb.save(out);return out.getvalue()
        mutate_zip(path,'ppt/embeddings/Microsoft_Excel_Worksheet.xlsx',change)
        result=self.check(path)
        self.assertNotEqual(result.returncode,0)
        self.assertIn('embedded workbook differs from source',result.stdout)

    def test_wrong_kpi_fails(self):
        path=self.candidate();prs=Presentation(path)
        slide=next(s for s in prs.slides if any(sh.has_chart for sh in s.shapes) and any(sh.has_text_frame and 'ملخص المهام' in sh.text for sh in s.shapes))
        next(sh for sh in slide.shapes if sh.shape_id==21).text='999';prs.save(path)
        result=self.check(path);self.assertNotEqual(result.returncode,0);self.assertIn('KPI shape',result.stdout)

    def test_completed_status_with_wrong_fill_fails(self):
        from pptx.dml.color import RGBColor
        path=self.candidate();prs=Presentation(path)
        cells=[row.cells[4] for sl in prs.slides for sh in sl.shapes if sh.has_table and len(sh.table.columns)==6 for row in list(sh.table.rows)[2:] if row.cells[4].text in ('مكتمل','مكتملة')]
        self.assertTrue(cells);cells[0].fill.solid();cells[0].fill.fore_color.rgb=RGBColor.from_string('EBF1DE');prs.save(path)
        result=self.check(path);self.assertNotEqual(result.returncode,0);self.assertIn('status fill',result.stdout)

    def test_order_change_fails_even_when_counts_match(self):
        errors=[];gate.cmp('tasks',[('1','A'),('2','B')],[('2','B'),('1','A')],errors)
        self.assertTrue(errors)

    def test_missing_relationship_target_fails(self):
        path=self.candidate()
        def change(data):
            root=ET.fromstring(data);root[0].set('Target','missing.xml');return ET.tostring(root)
        mutate_zip(path,'ppt/charts/_rels/chart1.xml.rels',change)
        errors=[];gate.package(path,errors)
        self.assertTrue(any('missing relationship target' in error for error in errors))

    def test_shape_relationship_without_definition_fails(self):
        path=self.candidate()
        def change(data):
            root=ET.fromstring(data)
            root.set('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id','rIdMissing')
            return ET.tostring(root)
        mutate_zip(path,'ppt/slides/slide1.xml',change)
        errors=[];gate.package(path,errors)
        self.assertTrue(any('missing relationship ID' in error for error in errors))

    def test_invalid_paragraph_property_order_fails(self):
        path=self.candidate()
        def change(data):
            root=ET.fromstring(data);ns='{http://schemas.openxmlformats.org/drawingml/2006/main}'
            props=next(el for el in root.iter(ns+'pPr'))
            for el in list(props):props.remove(el)
            ET.SubElement(props,ns+'buNone');ET.SubElement(props,ns+'lnSpc')
            return ET.tostring(root)
        mutate_zip(path,'ppt/slides/slide1.xml',change)
        errors=[];gate.package(path,errors)
        self.assertTrue(any('invalid bullet property order' in error for error in errors))


class SectorGateTests(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory();self.root=Path(self.temp.name)
        self.master,self.out,self.deck=synthetic_sector_files(self.root)

    def tearDown(self):self.temp.cleanup()

    def check(self):
        with contextlib.redirect_stdout(io.StringIO()):
            return sector_qa.check(self.out,self.master)

    def test_matching_records_pass(self):self.assertTrue(self.check())

    def test_duplicate_record_with_same_count_fails(self):
        prs=Presentation(self.deck);t=prs.slides[0].shapes[0].table
        for c in range(1,4):t.cell(3,c).text=t.cell(2,c).text
        prs.save(self.deck);self.assertFalse(self.check())

    def test_wrong_status_table_fails(self):
        prs=Presentation(self.deck);prs.slides[0].shapes[0].table.cell(0,0).text='المهام المتأخرة'
        prs.save(self.deck);self.assertFalse(self.check())

    def test_swapped_rows_fail(self):
        prs=Presentation(self.deck);t=prs.slides[0].shapes[0].table
        for c in range(1,4):
            a,b=t.cell(2,c).text,t.cell(3,c).text;t.cell(2,c).text=b;t.cell(3,c).text=a
        prs.save(self.deck);self.assertFalse(self.check())

    def test_missing_master_fails(self):
        with contextlib.redirect_stdout(io.StringIO()):self.assertFalse(sector_qa.check(self.out))

    def test_forged_manifest_count_does_not_hide_removed_record(self):
        prs=Presentation(self.deck);t=prs.slides[0].shapes[0].table;t._tbl.remove(t.rows[3]._tr);prs.slides[0].shapes[0].height=sum(r.height for r in t.rows);prs.save(self.deck)
        manifest=json.loads((self.out/'manifest.json').read_text(encoding='utf-8'));manifest[0]['n'][0]=1
        (self.out/'manifest.json').write_text(json.dumps(manifest),encoding='utf-8')
        self.assertFalse(self.check())


if __name__=='__main__':unittest.main()

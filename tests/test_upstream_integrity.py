# /// script
# requires-python = ">=3.12"
# dependencies = ["pandas>=2.2,<3", "openpyxl>=3.1,<4"]
# ///
"""Synthetic regression cases for the three source-workbook skills."""

import os
import importlib.util
import contextlib
import io
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook, load_workbook


ROOT = Path(os.environ.get('ENJAZ_REPO', Path(__file__).resolve().parents[1])).resolve()
SUHAIL = ROOT / "skills/suhail-project-report/scripts/build_report.py"
TASKS = ROOT / "skills/task-details/scripts/build_report.py"
TRANSACTIONS = ROOT / "skills/delayed-transactions-summary/scripts/build_report.py"


def run(script, *args):
    return subprocess.run(
        [sys.executable, str(script), *(str(arg) for arg in args)],
        text=True,
        capture_output=True,
        encoding="utf-8",
        errors="replace",
        env={**os.environ, "PYTHONIOENCODING": "utf-8"},
        check=False,
    )


def suhail_source(path, status="", headers=None):
    wb = Workbook()
    ws = wb.active
    ws.title = "مشاريع سهيل"
    headers = headers or [
        "القطاع", "اسم المشروع", "تاريخ البداية", "تاريخ النهاية",
        "الحالة الفعلية W37", "ما تم حتى تاريخه W37", "الحالة الفعلية W38",
        "ما تم حتى تاريخه W38", "ملاحظات", "التحدي",
    ]
    ws.append(headers)
    values = {
        "القطاع": "وكالة شؤون الحج", "اسم المشروع": "مشروع تجريبي",
        "تاريخ البداية": "01/01/1448", "تاريخ النهاية": "01/02/1448",
        "الحالة الفعلية W37": "مكتمل", "ما تم حتى تاريخه W37": "قديم",
        "الحالة الفعلية W38": status, "ما تم حتى تاريخه W38": "حديث",
        "ملاحظات": "ملاحظة", "التحدي": "",
    }
    ws.append([values.get(header) for header in headers])
    wb.save(path)


def task_source(path, support=True, support_headers=None):
    wb = Workbook()
    ws = wb.active
    ws.title = "المهام"
    ws.append(["المهمة", "القطاع", "المصدر", "تاريخ الإنجاز المخطط", "الحالة", "الملاحظات"])
    ws.append(["مهمة تجريبية", "وكالة شؤون الحج", "اجتماع القيادات", "01/02/1448", "مكتمل", "ملاحظة"])
    if support:
        sup = wb.create_sheet("طلبات الدعم")
        support_headers = support_headers or ["المهمة", "القطاع", "طلب الدعم"]
        sup.append(support_headers)
        sup.append([{"المهمة": "مهمة تجريبية", "القطاع": "وكالة شؤون الحج", "طلب الدعم": "طلب محدد"}.get(h) for h in support_headers])
    wb.save(path)


def transaction_source(path, *, header=True, sector="قطاع أ", status="متأخر", number="000123", empty=False):
    wb = Workbook()
    ws = wb.active
    ws.cell(2, 13, sector)
    if header:
        ws.cell(4, 31, "رقم المعاملة")
    if header and not empty:
        ws.cell(5, 20, status)
        ws.cell(5, 24, "01/02/1448")
        ws.cell(5, 27, "جهة واردة")
        ws.cell(5, 28, "01/01/1448")
        ws.cell(5, 29, "موضوع تجريبي")
        ws.cell(5, 31, number)
    wb.save(path)


class UpstreamIntegrityTests(unittest.TestCase):
    def test_suhail_uses_latest_week_independently_and_preserves_status(self):
        with tempfile.TemporaryDirectory() as td:
            source, output = Path(td) / "source.xlsx", Path(td) / "out.xlsx"
            suhail_source(source, "متأخر")
            wb = load_workbook(source)
            ws = wb.active
            ws.cell(1, 11, "ما تم حتى تاريخه W39")
            ws.cell(2, 11, "تحديث أحدث\nسطر ثان")
            wb.save(source)
            result = run(SUHAIL, source, output)
            self.assertEqual(result.returncode, 0, result.stderr + result.stdout)
            self.assertIn("W38 (status) / W39 (progress)", result.stdout)
            wb = load_workbook(output, data_only=True)
            self.assertEqual(wb["ملخص مشاريع سهيل"]["D3"].value, 1)
            detail = wb["تفاصيل مشاريع سهيل"]
            self.assertEqual(detail["E3"].value, "متأخر")
            self.assertEqual(detail["F3"].value, "تحديث أحدث\nسطر ثان")

    def test_suhail_defaults_only_blank_latest_status_and_reports_count(self):
        with tempfile.TemporaryDirectory() as td:
            source, output = Path(td) / "source.xlsx", Path(td) / "out.xlsx"
            suhail_source(source)
            result = run(SUHAIL, source, output)
            self.assertEqual(result.returncode, 0, result.stderr + result.stdout)
            self.assertIn("blank latest-week statuses defaulted: 1", result.stdout)
            wb = load_workbook(output, data_only=True)
            self.assertEqual(wb["ملخص مشاريع سهيل"]["C3"].value, 1)
            self.assertIn("حديث", [c.value for row in wb["تفاصيل مشاريع سهيل"] for c in row])

    def test_suhail_rejects_unknown_latest_status(self):
        with tempfile.TemporaryDirectory() as td:
            source, output = Path(td) / "source.xlsx", Path(td) / "out.xlsx"
            suhail_source(source, "قيد المراجعة")
            result = run(SUHAIL, source, output)
            self.assertNotEqual(result.returncode, 0)
            self.assertIn("قيد المراجعة", result.stderr + result.stdout)
            self.assertFalse(output.exists())

    def test_suhail_names_missing_required_column(self):
        with tempfile.TemporaryDirectory() as td:
            source, output = Path(td) / "source.xlsx", Path(td) / "out.xlsx"
            suhail_source(source, headers=["القطاع", "اسم المشروع", "تاريخ البداية", "الحالة الفعلية W38", "ما تم حتى تاريخه W38"])
            result = run(SUHAIL, source, output)
            self.assertNotEqual(result.returncode, 0)
            self.assertIn("تاريخ النهاية", result.stderr + result.stdout)
            self.assertFalse(output.exists())

    def assert_blank_manual_table(self, output, title, headers):
        wb = load_workbook(output)
        ws = wb["ملخص المهام" if title == "طلبات الدعم" else "ملخص مشاريع سهيل"]
        row = next(c.row for cells in ws for c in cells if c.value == title)
        self.assertEqual([ws.cell(row+1, i).value for i in range(1,5)], headers)
        self.assertEqual([ws.cell(row+2, i).value for i in range(1,5)], [None]*4)
        for i in range(1,5):
            cell=ws.cell(row+2,i)
            self.assertEqual(cell.font.name, "Abar Mid")
            self.assertEqual(cell.font.sz, 10)
            self.assertEqual(cell.border.bottom.style, "thin")
        self.assertTrue(ws.sheet_view.rightToLeft)
        return wb

    def test_tasks_manual_support_without_source_sheet(self):
        with tempfile.TemporaryDirectory() as td:
            source, output = Path(td)/"source.xlsx", Path(td)/"out.xlsx"
            task_source(source, support=False)
            result = run(TASKS, "--input", source, "--output", output)
            self.assertEqual(result.returncode, 0, result.stderr+result.stdout)
            self.assertNotIn("طلبات الدعم", result.stderr+result.stdout)
            wb=self.assert_blank_manual_table(output, "طلبات الدعم", ["#","المهمة","القطاع","طلب الدعم"])
            self.assertEqual(wb.sheetnames, ["ملخص المهام","تفاصيل المهام"])
            self.assertEqual(wb["ملخص المهام"]["B2"].value,1)

    def test_tasks_ignore_populated_or_malformed_support_sheets(self):
        for headers in (["المهمة","القطاع","طلب الدعم"],["المهمة","طلب الدعم"],["المهمة","القطاع"]):
            with self.subTest(headers=headers), tempfile.TemporaryDirectory() as td:
                source, output=Path(td)/"source.xlsx",Path(td)/"out.xlsx"
                task_source(source,support_headers=headers)
                result=run(TASKS,"--input",source,"--output",output)
                self.assertEqual(result.returncode,0,result.stderr+result.stdout)
                wb=self.assert_blank_manual_table(output,"طلبات الدعم",["#","المهمة","القطاع","طلب الدعم"])
                self.assertNotIn("طلب محدد",[c.value for ws in wb for row in ws for c in row])

    def test_suhail_challenges_always_manual_even_with_source_records(self):
        for kind in ("absent","column","sheet","malformed"):
            with self.subTest(kind=kind), tempfile.TemporaryDirectory() as td:
                source,output=Path(td)/"source.xlsx",Path(td)/"out.xlsx"
                suhail_source(source,"متأخر")
                wb=load_workbook(source)
                if kind=="column":wb.active.cell(2,10,"تحدي من المصدر")
                if kind in ("sheet","malformed"):
                    ws=wb.create_sheet("التحديات")
                    ws.append(["المشروع","القطاع","التحدي"] if kind=="sheet" else ["التحدي"])
                    ws.append(["مشروع تجريبي","وكالة شؤون الحج","تحدي من المصدر"])
                if kind=="absent":wb.active.delete_cols(10)
                wb.save(source)
                result=run(SUHAIL,source,output)
                self.assertEqual(result.returncode,0,result.stderr+result.stdout)
                self.assertNotIn("challenge",result.stderr.lower())
                wb=self.assert_blank_manual_table(output,"سادسا: التحديات",["#","المشروع","القطاع","التحدي"])
                self.assertNotIn("تحدي من المصدر",[c.value for ws in wb for row in ws for c in row])
                self.assertEqual(wb["تفاصيل مشاريع سهيل"]["E3"].value,"متأخر")

    def test_transactions_accept_valid_empty_sector(self):
        with tempfile.TemporaryDirectory() as td:
            input_dir, output = Path(td) / "input", Path(td) / "out.xlsx"
            input_dir.mkdir()
            transaction_source(input_dir / "empty.xlsx", empty=True)
            result = run(TRANSACTIONS, "--input-dir", input_dir, "--output", output, "--dedupe")
            self.assertEqual(result.returncode, 0, result.stderr + result.stdout)
            ws = load_workbook(output, data_only=True).active
            self.assertEqual([ws.cell(r, 2).value for r in range(2, 6)], [0, 0, 0, 0])
            self.assertEqual(ws["A9"].value, "قطاع أ")

    def test_transactions_count_all_keeps_cross_sector_records(self):
        with tempfile.TemporaryDirectory() as td:
            input_dir, output = Path(td) / "input", Path(td) / "out.xlsx"
            input_dir.mkdir()
            transaction_source(input_dir / "one.xlsx", sector="قطاع أ")
            transaction_source(input_dir / "two.xlsx", sector="قطاع ب")
            result = run(TRANSACTIONS, "--input-dir", input_dir, "--output", output)
            self.assertEqual(result.returncode, 0, result.stderr + result.stdout)
            ws = load_workbook(output, data_only=True).active
            self.assertEqual(ws["B2"].value, 2)
            self.assertEqual(sum(cell.value == "000123" for row in ws for cell in row), 2)

    def test_transactions_dedupe_identical_records(self):
        with tempfile.TemporaryDirectory() as td:
            input_dir, output = Path(td) / "input", Path(td) / "out.xlsx"
            input_dir.mkdir()
            transaction_source(input_dir / "one.xlsx")
            transaction_source(input_dir / "two.xlsx")
            result = run(TRANSACTIONS, "--input-dir", input_dir, "--output", output, "--dedupe")
            self.assertEqual(result.returncode, 0, result.stderr + result.stdout)
            ws = load_workbook(output, data_only=True).active
            self.assertEqual(ws["B2"].value, 1)
            self.assertEqual(sum(cell.value == "000123" for row in ws for cell in row), 1)

    def test_transactions_no_delayed_rows_is_valid(self):
        with tempfile.TemporaryDirectory() as td:
            input_dir, output = Path(td) / "input", Path(td) / "out.xlsx"
            input_dir.mkdir()
            transaction_source(input_dir / "one.xlsx", status="مكتمل")
            result = run(TRANSACTIONS, "--input-dir", input_dir, "--output", output)
            self.assertEqual(result.returncode, 0, result.stderr + result.stdout)
            ws = load_workbook(output, data_only=True).active
            self.assertEqual(ws["B3"].value, 1)
            self.assertEqual(ws["B5"].value, 0)

    def test_transactions_reject_unreadable_file(self):
        with tempfile.TemporaryDirectory() as td:
            input_dir, output = Path(td) / "input", Path(td) / "out.xlsx"
            input_dir.mkdir()
            transaction_source(input_dir / "valid.xlsx")
            transaction_source(input_dir / "missing-header.xlsx", header=False)
            result = run(TRANSACTIONS, "--input-dir", input_dir, "--output", output)
            self.assertNotEqual(result.returncode, 0)
            self.assertIn("missing-header.xlsx", result.stderr + result.stdout)
            self.assertFalse(output.exists())

    def test_transactions_reject_conflicting_duplicate_before_dedupe(self):
        with tempfile.TemporaryDirectory() as td:
            input_dir, output = Path(td) / "input", Path(td) / "out.xlsx"
            input_dir.mkdir()
            transaction_source(input_dir / "one.xlsx", status="متأخر")
            transaction_source(input_dir / "two.xlsx", status="مكتمل")
            result = run(TRANSACTIONS, "--input-dir", input_dir, "--output", output, "--dedupe")
            self.assertNotEqual(result.returncode, 0)
            self.assertIn("000123", result.stderr + result.stdout)
            self.assertFalse(output.exists())

    def test_transactions_reject_conflicting_sector_before_dedupe(self):
        with tempfile.TemporaryDirectory() as td:
            input_dir, output = Path(td) / "input", Path(td) / "out.xlsx"
            input_dir.mkdir()
            transaction_source(input_dir / "one.xlsx", sector="قطاع أ")
            transaction_source(input_dir / "two.xlsx", sector="قطاع ب")
            result = run(TRANSACTIONS, "--input-dir", input_dir, "--output", output, "--dedupe")
            self.assertNotEqual(result.returncode, 0)
            self.assertIn("000123", result.stderr + result.stdout)
            self.assertFalse(output.exists())

    def test_transactions_valid_report_keeps_identifier_as_text(self):
        with tempfile.TemporaryDirectory() as td:
            input_dir, output = Path(td) / "input", Path(td) / "out.xlsx"
            input_dir.mkdir()
            transaction_source(input_dir / "one.xlsx")
            result = run(TRANSACTIONS, "--input-dir", input_dir, "--output", output)
            self.assertEqual(result.returncode, 0, result.stderr + result.stdout)
            ws = load_workbook(output, data_only=True).active
            self.assertEqual(ws["B2"].value, 1)
            self.assertEqual(ws["B5"].value, 1)
            ids = [cell for row in ws for cell in row if cell.value == "000123"]
            self.assertEqual(len(ids), 1)
            self.assertEqual(ids[0].data_type, "s")

    def test_transactions_failed_verification_preserves_previous_output(self):
        with tempfile.TemporaryDirectory() as td:
            input_dir, output = Path(td) / "input", Path(td) / "out.xlsx"
            input_dir.mkdir()
            transaction_source(input_dir / "one.xlsx")
            output.write_bytes(b"previous approved report")
            spec = importlib.util.spec_from_file_location("transactions_under_test", TRANSACTIONS)
            module = importlib.util.module_from_spec(spec)
            spec.loader.exec_module(module)
            original_build = module.build
            def corrupt_candidate(*args, **kwargs):
                result = original_build(*args, **kwargs)
                wb = load_workbook(args[2])
                wb.active["B2"] = 999
                wb.save(args[2])
                wb.close()
                return result
            with contextlib.redirect_stdout(io.StringIO()), patch.object(module, "build", side_effect=corrupt_candidate), patch.object(
                sys, "argv", [str(TRANSACTIONS), "--input-dir", str(input_dir), "--output", str(output)]
            ):
                with self.assertRaises(SystemExit) as stopped:
                    module.main()
            self.assertIn("verification failed", str(stopped.exception))
            self.assertEqual(output.read_bytes(), b"previous approved report")
            self.assertEqual(list(Path(td).glob(".transactions-unverified-*")), [])


if __name__ == "__main__":
    unittest.main()

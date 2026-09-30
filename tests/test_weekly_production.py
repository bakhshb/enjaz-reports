"""Acceptance of the installed weekly entry point and deliberate failure cases."""
from pathlib import Path
import contextlib, io, json, shutil, tempfile, unittest
from unittest.mock import patch
from openpyxl import load_workbook
from pptx import Presentation
from pptx.dml.color import RGBColor
from xml.etree import ElementTree as ET
import run_weekly_workflow as workflow

weekly=workflow.weekly
gate=workflow.weekly_validation

class WeeklyProductionTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.temp=tempfile.TemporaryDirectory();cls.root=Path(cls.temp.name)/'baseline'
        with contextlib.redirect_stdout(io.StringIO()):workflow.scenario(cls.root,'normal')

    @classmethod
    def tearDownClass(cls):cls.temp.cleanup()

    def setUp(self):
        self.temp=tempfile.TemporaryDirectory();self.case=Path(self.temp.name)
        for path in self.root.glob('*.xlsx'):shutil.copy2(path,self.case/path.name)
        self.deck=self.case/'weekly.pptx';shutil.copy2(self.root/'weekly.pptx',self.deck)

    def tearDown(self):self.temp.cleanup()

    def args(self,output=None):
        return ['--tasks',str(self.case/'tasks.xlsx'),'--suhail',str(self.case/'projects.xlsx'),
                '--transactions',str(self.case/'transactions.xlsx'),'--output',str(output or self.deck)]

    def check(self):
        return gate.validate(self.deck,self.case/'tasks.xlsx',self.case/'projects.xlsx',self.case/'transactions.xlsx')

    def test_all_statuses_and_independent_summary_sections(self):
        self.assertEqual(self.check()['programmatic'],'passed')
        statuses={kind:set() for kind in ('task','project')}
        for sl in Presentation(self.deck).slides:
            for sh in weekly.builder.tables(sl,6):
                kind='task' if sh.table.cell(1,1).text=='المهمة' else 'project'
                statuses[kind].update(row.cells[4].text for row in list(sh.table.rows)[2:])
        self.assertEqual(statuses['task'],weekly.status_styles.TASK_STATUSES)
        self.assertEqual(statuses['project'],weekly.status_styles.PROJECT_STATUSES)

    def test_status_and_notes_font_corruption_rejected(self):
        original=self.deck.read_bytes()
        for column,change in [(4,'fill'),(4,'font'),(5,'font')]:
            with self.subTest(column=column,change=change):
                self.deck.write_bytes(original);prs=Presentation(self.deck)
                sh=next(sh for sl in prs.slides for sh in weekly.builder.tables(sl,6))
                cell=sh.table.cell(2,column)
                if change=='fill':cell.fill.solid();cell.fill.fore_color.rgb=RGBColor(255,0,255)
                else:cell.text_frame.paragraphs[0].runs[0].font.name='Calibri'
                prs.save(self.deck);before=self.deck.read_bytes()
                with self.assertRaises(ValueError):self.check()
                self.assertEqual(before,self.deck.read_bytes())

    def test_text_whitespace_and_record_order_corruption_rejected(self):
        original=self.deck.read_bytes()
        for kind in ('whitespace','order','group'):
            self.deck.write_bytes(original);prs=Presentation(self.deck)
            sh=next(sh for sl in prs.slides for sh in weekly.builder.tables(sl,6))
            if kind=='whitespace':weekly.set_cell_exact(sh.table.cell(2,1),sh.table.cell(2,1).text+' ')
            elif kind=='group':weekly.set_cell_exact(sh.table.cell(0,0),'تفاصيل المهام (مصدر خاطئ)')
            else:
                a=sh.table.cell(2,1).text;b=sh.table.cell(3,1).text
                weekly.set_cell_exact(sh.table.cell(2,1),b);weekly.set_cell_exact(sh.table.cell(3,1),a)
            prs.save(self.deck)
            with self.assertRaises(ValueError):self.check()

    def test_inconsistent_source_total_rejected_without_overwrite(self):
        wb=load_workbook(self.case/'tasks.xlsx');wb['ملخص المهام']['B2']=999;wb.save(self.case/'tasks.xlsx')
        before=self.deck.read_bytes()
        with self.assertRaises(ValueError):weekly.main(self.args())
        self.assertEqual(before,self.deck.read_bytes())

    def test_blank_notes_and_repeated_names_preserved(self):
        # Keep source identities consistent in summary/details; repeated names are records.
        wb=load_workbook(self.case/'tasks.xlsx')
        for ws in wb:
            for row in ws:
                for cell in row:
                    if isinstance(cell.value,str) and cell.value.startswith('مهمة تجريبية'):cell.value='اسم متكرر'
                    if isinstance(cell.value,str) and cell.value.startswith('ملاحظة تجريبية'):cell.value=None
        wb.save(self.case/'tasks.xlsx')
        with contextlib.redirect_stdout(io.StringIO()):weekly.main(self.args())
        self.check()
        notes=[row.cells[5].text for sl in Presentation(self.deck).slides for sh in weekly.builder.detail_tables(sl,'task') for row in list(sh.table.rows)[2:]]
        self.assertEqual(notes,['']*4)

    def test_transaction_leading_zero_preserved(self):
        records=[]
        for sl in Presentation(self.deck).slides:
            for sh in weekly.builder.tables(sl,5):records.extend(c.text for row in list(sh.table.rows)[2:] for c in row.cells)
        self.assertIn('000123',records)

    def test_chart_cache_corruption_rejected(self):
        def change(data):
            root=ET.fromstring(data)
            value=root.find('.//{http://schemas.openxmlformats.org/drawingml/2006/chart}numCache/{http://schemas.openxmlformats.org/drawingml/2006/chart}pt/{http://schemas.openxmlformats.org/drawingml/2006/chart}v')
            value.text='999';return ET.tostring(root)
        workflow.fixtures.mutate_zip(self.deck,'ppt/charts/chart1.xml',change)
        with self.assertRaisesRegex(ValueError,'cache differs'):self.check()

    def test_wrong_navigation_target_rejected(self):
        prs=Presentation(self.deck)
        slide=weekly.builder.find_slide_with_text(prs,'تفاصيل مشاريع سهيل')
        shape=next(sh for sh in slide.shapes if getattr(sh,'has_text_frame',False) and sh.text=='تفاصيل مشاريع سهيل')
        link=next(shape._element.iter(weekly.qn('a:hlinkClick')))
        slide.part.rels[link.get(weekly.qn('r:id'))]._target=prs.slides[0].part
        prs.save(self.deck)
        with self.assertRaisesRegex(ValueError,'Navigation target'):self.check()

    def test_summary_section_order_rejected(self):
        prs=Presentation(self.deck);lst=prs.slides._sldIdLst
        slide=lst[1];lst.remove(slide);lst.insert(2,slide);prs.save(self.deck)
        with self.assertRaisesRegex(ValueError,'order changed|outside its section'):self.check()

    def test_long_arabic_notes_preserve_line_breaks(self):
        note='متابعة تنفيذ الأعمال والتنسيق مع الجهات ذات العلاقة لضمان اكتمال المتطلبات\nملاحظة ثانية للتأكد من حفظ النص كما ورد دون اختصار'
        wb=load_workbook(self.case/'tasks.xlsx')
        for ws in wb:
            for row in ws:
                for cell in row:
                    if isinstance(cell.value,str) and cell.value.startswith('ملاحظة تجريبية'):cell.value=note
        wb.save(self.case/'tasks.xlsx')
        with contextlib.redirect_stdout(io.StringIO()):weekly.main(self.args())
        self.check()
        notes=[row.cells[5].text for sl in Presentation(self.deck).slides for sh in weekly.builder.detail_tables(sl,'task') for row in list(sh.table.rows)[2:]]
        self.assertEqual(notes,[note]*4)

    def test_failed_stages_preserve_output_and_cleanup(self):
        for stage in ('build','validation','publication'):
            for existing in (False,True):
                with self.subTest(stage=stage,existing=existing):
                    output=self.case/'protected.pptx'
                    if existing:output.write_bytes(b'previous accepted output')
                    else:output.unlink(missing_ok=True)
                    target=(weekly,'build_candidate') if stage=='build' else (gate,'validate') if stage=='validation' else (weekly.status_styles.os,'replace')
                    with patch.object(*target,side_effect=RuntimeError('injected failure')):
                        with self.assertRaises(RuntimeError):weekly.main(self.args(output))
                    self.assertEqual(output.read_bytes() if output.exists() else None,b'previous accepted output' if existing else None)
                    self.assertFalse(list(self.case.glob('.weekly-build-*')))
                    self.assertFalse(list(self.case.glob('.monday-publish-*')))

    def test_output_cannot_alias_sources_or_master(self):
        for path in (self.case/'tasks.xlsx',self.case/'projects.xlsx',self.case/'transactions.xlsx',weekly.MASTER):
            before=path.read_bytes()
            with self.assertRaises(ValueError):weekly.main(self.args(path))
            self.assertEqual(before,path.read_bytes())

    def test_unknown_status_rejected(self):
        wb=load_workbook(self.case/'projects.xlsx');ws=wb['تفاصيل مشاريع سهيل']
        row=next(row for row in ws if isinstance(row[0].value,int))
        row[4].value='حالة غير معتمدة';wb.save(self.case/'projects.xlsx')
        before=self.deck.read_bytes()
        with self.assertRaisesRegex(ValueError,'Unsupported'):weekly.main(self.args())
        self.assertEqual(before,self.deck.read_bytes())

    def test_measured_layout_preserves_records_and_rejects_stale_inputs(self):
        tables={}
        for sl in Presentation(self.deck).slides:
            for sh in weekly.builder.tables(sl):
                title=weekly.builder.table_title(sh)
                table=tables.setdefault(title,{'header':[22,22],'body':[]})
                table['body'].extend([95]*(len(sh.table.rows)-2))
        paths={'tasks':self.case/'tasks.xlsx','suhail':self.case/'projects.xlsx','transactions':self.case/'transactions.xlsx','template':weekly.MASTER}
        evidence={'hashes':{key:gate.validation.sha(path) for key,path in paths.items()},'tables':tables}
        measurements=self.case/'layout.json';measurements.write_text(json.dumps(evidence),encoding='utf-8')
        with contextlib.redirect_stdout(io.StringIO()):weekly.main(self.args()+['--measurements',str(measurements)])
        self.check()
        before=self.deck.read_bytes();evidence['hashes']['tasks']='stale'
        measurements.write_text(json.dumps(evidence),encoding='utf-8')
        with self.assertRaisesRegex(ValueError,'stale'):weekly.main(self.args()+['--measurements',str(measurements)])
        self.assertEqual(before,self.deck.read_bytes())

if __name__=='__main__':unittest.main()

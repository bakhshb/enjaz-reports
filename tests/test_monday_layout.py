"""Regression checks for wasted space, RTL conversion and update paragraphs."""
import copy
import sys
import unittest
from pptx import Presentation
from pptx.oxml.ns import qn
from test_deck_integrity import MONDAY, gate

sys.path.insert(0,str(MONDAY))
import monday_layout as layout
from report_common.update_format import update_parts, format_update


class MondayLayoutTests(unittest.TestCase):
    def test_headers_count_once_and_next_table_uses_free_space(self):
        pages=layout.pack([('a',[1,2],[10,10],[30,30]),('b',[3],[10,10],[20])],0,125,gap=5)
        self.assertEqual(len(pages),1)
        self.assertEqual([item[0] for item in pages[0]],['a','b'])
        self.assertEqual(pages[0][1][-1],85)

    def test_continuations_preserve_order_and_repeat_headers(self):
        pages=layout.pack([('a',list(range(5)),[10,10],[30]*5)],0,100,gap=5)
        self.assertEqual([len(p[0][1]) for p in pages],[2,2,1])
        self.assertEqual([r for p in pages for item in p for r in item[1]],list(range(5)))
        with self.assertRaisesRegex(ValueError,'too tall'):
            layout.pack([('a',[1],[10,10],[81])],0,100)

    def test_rtl_conversion_preserves_columns_merges_and_schema_order(self):
        prs=Presentation(MONDAY.parent/'assets/monday-meeting-master.pptx')
        for slide in prs.slides:
            for shape in slide.shapes:
                if not shape.has_table or len(shape.table.columns)!=4 or shape.table._tbl.tblPr.get('rtl')=='1':continue
                t=shape.table
                headers=[c.text for c in t.rows[1].cells]
                widths=[c.width for c in t.columns]
                layout.rtl_table(t)
                self.assertEqual([c.text for c in t.rows[1].cells],list(reversed(headers)))
                self.assertEqual([c.width for c in t.columns],list(reversed(widths)))
                before=t._tbl.xml;layout.rtl_table(t);self.assertEqual(before,t._tbl.xml)
                for row in t.rows:
                    tags=[c.tag for c in row._tr]
                    if qn('a:extLst') in tags:self.assertEqual(tags[-1],qn('a:extLst'))
                for row in t.rows:
                    for i,c in enumerate(row.cells):
                        if c.is_merge_origin:
                            self.assertTrue(all(row.cells[j].is_spanned for j in range(i+1,i+c.span_width)))

    def test_date_bold_multiple_points_and_single_sentence(self):
        prs=Presentation(MONDAY.parent/'assets/monday-meeting-master.pptx')
        style=layout.bullet_example(prs)
        shape=next(sh for sl in prs.slides for sh in sl.shapes if sh.has_table and len(sh.table.columns)==6 and sh.table.cell(1,1).text=='اسم المشروع')
        cell=shape.table.cell(2,5)
        for text,count in [('\nتاريخ التحديث 04 أكتوبر 2026\nنقطة أولى\nنقطة ثانية\n',2),('تاريخ التحديث 04 أكتوبر 2026\nجملة واحدة ممتدة دون تقسيم.',0),('تاريخ التحديث 04 أكتوبر 2026\n• نقطة واحدة',0)]:
            cell.text=text;format_update(cell,style,paragraph_items=True)
            self.assertEqual(sum(p._p.find('./'+qn('a:pPr')+'/'+qn('a:buChar')) is not None for p in cell.text_frame.paragraphs),count)
            expected=update_parts(text,paragraph_items=True)
            self.assertEqual([p.text for p in cell.text_frame.paragraphs],[part[0] for part in expected])
            self.assertTrue(all(r.font.bold for r in cell.text_frame.paragraphs[0].runs))
            self.assertTrue(all(not r.font.bold for p in cell.text_frame.paragraphs[1:] for r in p.runs))

    def test_empty_update_retains_one_empty_paragraph(self):
        self.assertEqual(update_parts('  \n',paragraph_items=True),[('',0,False)])

    def test_gate_rejects_lost_bold_and_bullets(self):
        prs=Presentation(MONDAY.parent/'assets/monday-meeting-master.pptx')
        style=layout.bullet_example(prs)
        shape=next(sh for sl in prs.slides for sh in sl.shapes if sh.has_table and len(sh.table.columns)==6 and sh.table.cell(1,1).text=='اسم المشروع')
        seed=copy.deepcopy(shape.element)
        for sl in list(prs.slides):layout.remove_slide(prs,sl)
        sl=prs.slides.add_slide(prs.slide_layouts[0]);sl.shapes._spTree.insert_element_before(seed,'p:extLst');sh=sl.shapes[-1]
        for row in list(sh.table.rows)[3:]:sh.table._tbl.remove(row._tr)
        text='تاريخ التحديث 04 أكتوبر 2026\nأول نقطة\nثاني نقطة'
        cell=sh.table.cell(2,5);cell.text=text;format_update(cell,style,paragraph_items=True)
        for p in cell.text_frame.paragraphs:p._p.get_or_add_pPr().set('rtl','1')
        errors=[];gate.verify_update_format(prs,[(text,)],errors);self.assertFalse(errors,errors)
        cell.text_frame.paragraphs[0].runs[0].font.bold=False
        p=cell.text_frame.paragraphs[1]._p.get_or_add_pPr();p.remove(p.find(qn('a:buChar')))
        errors=[];gate.verify_update_format(prs,[(text,)],errors)
        self.assertIn('Suhail update date/body bold differs',errors)
        self.assertIn('Suhail update bullets differ',errors)

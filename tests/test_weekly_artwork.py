"""Inherited artwork must survive continuation and fail closed when damaged."""
import contextlib
import io
import tempfile
import unittest
from pathlib import Path
from pptx import Presentation
from pptx.oxml.ns import qn
import run_weekly_workflow as workflow
import weekly_artwork as artwork

class WeeklyArtworkTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.temp = tempfile.TemporaryDirectory()
        cls.root = Path(cls.temp.name)
        with contextlib.redirect_stdout(io.StringIO()):
            workflow.scenario(cls.root, 'normal')
        cls.original = (cls.root/'weekly.pptx').read_bytes()

    @classmethod
    def tearDownClass(cls):
        cls.temp.cleanup()

    def test_artwork_corruption_is_rejected_without_repair(self):
        for damage in ('hidden_slide', 'hidden_layout', 'logo', 'footer', 'label',
                       'wrong_layout', 'wrong_link', 'missing_link', 'duplicate_id', 'moved_logo'):
            with self.subTest(damage=damage):
                deck = self.root/'damaged.pptx'; deck.write_bytes(self.original)
                prs = Presentation(deck)
                slide = next(sl for sl in prs.slides if artwork.detail_kind(sl) == 'project')
                layout = slide.slide_layout
                if damage == 'hidden_slide': slide._element.set('showMasterSp', '0')
                elif damage == 'hidden_layout': layout._element.set('showMasterSp', '0')
                elif damage == 'wrong_layout':
                    rel = next(r for r in slide.part.rels.values() if r.reltype.endswith('/slideLayout'))
                    rel._target = prs.slides[0].slide_layout.part
                elif damage in ('logo', 'footer', 'label', 'moved_logo'):
                    shape = next(sh for sh in layout.shapes if
                                 (sh.name == 'Picture 10' if damage in ('logo', 'moved_logo') else
                                  sh.name == 'Straight Connector 11' if damage == 'footer' else
                                  getattr(sh, 'text', '') == 'المعاملات'))
                    if damage == 'moved_logo': shape.top += 914400
                    else: shape.element.getparent().remove(shape.element)
                elif damage in ('wrong_link', 'missing_link'):
                    _, shapes = next(artwork.buttons(layout))
                    link = next(shapes[0].element.iter(qn('a:hlinkClick')))
                    if damage == 'missing_link': link.getparent().remove(link)
                    else: layout.part.rels[link.get(qn('r:id'))]._target = prs.slides[-1].part
                else:
                    group = next(sh for sh in slide.shapes if hasattr(sh, 'shapes'))
                    table = next(sh for sh in slide.shapes if sh.has_table)
                    table.element.find('.//'+qn('p:cNvPr')).set('id', str(group.shapes[0].shape_id))
                prs.save(deck); before = deck.read_bytes()
                with self.assertRaises(ValueError): workflow.verify(self.root, deck)
                self.assertEqual(before, deck.read_bytes())

    def test_continuation_keeps_artwork_links_and_unique_ids(self):
        prs = Presentation(workflow.weekly.MASTER)
        generated = []
        for kind in ('task', 'project'):
            seed = workflow.weekly.builder.detail_pattern_slides(prs, kind)[0]
            pattern = workflow.weekly.builder.detail_tables(seed, kind)[0]
            rows = [[str(i+1), 'سجل تجريبي', 'قطاع تجريبي' if kind == 'task' else '01/01/1448',
                     '01/02/1448', 'على المخطط', 'نص تجريبي'] for i in range(45)]
            seed._element.set('showMasterSp', '0')
            pages = workflow.weekly.flow(prs, seed, [('اختبار', rows, pattern)])
            self.assertGreater(len(pages), 1)
            generated.extend(pages)
        artwork.prepare_navigation(prs)
        deck = self.root/'continuations.pptx'; prs.save(deck); prs = Presentation(deck)
        errors = []; artwork.check_detail_artwork(prs, Presentation(workflow.weekly.MASTER), errors)
        self.assertFalse(errors, errors)
        self.assertGreater(len(generated), 4)
        for slide in prs.slides:
            if artwork.detail_kind(slide):
                ids = [sh.shape_id for sh in artwork.all_shapes(slide.shapes)]
                self.assertEqual(len(ids), len(set(ids)))

    def test_all_four_buttons_link_to_their_sections(self):
        prs = Presentation(self.root/'weekly.pptx')
        destinations = artwork.targets(prs)
        for slide in prs.slides:
            if not artwork.detail_kind(slide): continue
            layout = slide.slide_layout
            found = []
            for name, shapes in artwork.buttons(layout):
                found.append(name)
                for shape in shapes:
                    link = next(shape.element.iter(qn('a:hlinkClick')))
                    self.assertEqual(layout.part.rels[link.get(qn('r:id'))].target_part, destinations[name])
            self.assertEqual(set(found), set(artwork.NAVIGATION))
